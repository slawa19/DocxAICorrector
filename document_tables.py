import html

from docx.table import Table

from models import ImageAsset, RawTable


def build_raw_table(
    table: Table,
    image_assets: list[ImageAsset],
    *,
    raw_index: int,
    asset_id: str,
    build_paragraph_text_with_placeholders,
) -> RawTable | None:
    html_table = render_table_html(
        table,
        image_assets,
        build_paragraph_text_with_placeholders=build_paragraph_text_with_placeholders,
    )
    if not html_table.strip():
        return None
    return RawTable(raw_index=raw_index, html_text=html_table, asset_id=asset_id)


def render_table_html(
    table: Table,
    image_assets: list[ImageAsset],
    *,
    build_paragraph_text_with_placeholders,
) -> str:
    rows: list[list[str]] = []
    for row in table.rows:
        rendered_row = [
            render_table_cell(
                cell,
                image_assets,
                build_paragraph_text_with_placeholders=build_paragraph_text_with_placeholders,
            )
            for cell in row.cells
        ]
        rows.append(rendered_row)

    if not any(any(cell.strip() for cell in row) for row in rows):
        return ""

    has_header = len(rows) > 1 and all(cell.strip() for cell in rows[0])
    lines = ["<table>"]
    if has_header:
        lines.append("<thead>")
        lines.append(render_table_html_row(rows[0], cell_tag="th"))
        lines.append("</thead>")
        body_rows = rows[1:]
    else:
        body_rows = rows

    lines.append("<tbody>")
    for row in body_rows:
        lines.append(render_table_html_row(row, cell_tag="td"))
    lines.append("</tbody>")
    lines.append("</table>")
    return "\n".join(lines)


def render_table_cell(
    cell,
    image_assets: list[ImageAsset],
    *,
    build_paragraph_text_with_placeholders,
) -> str:
    cell_parts: list[str] = []
    for paragraph in cell.paragraphs:
        text = build_paragraph_text_with_placeholders(paragraph, image_assets).strip()
        if text:
            cell_parts.append(escape_html_preserving_breaks(text))
    return "<br/>".join(cell_parts)


def render_table_html_row(cells: list[str], *, cell_tag: str) -> str:
    rendered_cells = "".join(f"<{cell_tag}>{cell or '&nbsp;'}</{cell_tag}>" for cell in cells)
    return f"<tr>{rendered_cells}</tr>"


def escape_html_preserving_breaks(text: str) -> str:
    return "<br/>".join(html.escape(part, quote=False) for part in text.split("<br/>"))
