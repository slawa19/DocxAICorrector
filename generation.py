import logging
import re
import tempfile
import time
from collections.abc import Iterable, Mapping, Sequence, Sized
from functools import lru_cache
from pathlib import Path
from typing import TYPE_CHECKING, Any, cast

import lxml.etree as etree
import pypandoc
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from image_shared import is_retryable_error
from logger import log_event

if TYPE_CHECKING:
    from docx.document import Document as DocxDocument
    from openai import OpenAI


_SUPPORTED_RESPONSE_TEXT_TYPES = {"output_text", "text"}
_PARAGRAPH_MARKER_PATTERN = re.compile(r"\[\[DOCX_PARA_([A-Za-z0-9_]+)\]\]")
_IMAGE_ONLY_TARGET_PATTERN = re.compile(r"^(?:\s*\[\[DOCX_IMAGE_img_\d+\]\]\s*)+$")
_WORD_TOKEN_PATTERN = re.compile(r"\w+(?:[-']\w+)*", re.UNICODE)
_INCOMPLETE_RESPONSE_RETRY_MIN_OUTPUT_TOKENS = 1024
_INCOMPLETE_RESPONSE_RECOVERY_MIN_OUTPUT_TOKENS = 1536
_CONTEXT_LEAKAGE_RETRY_WARNING = (
    "ВАЖНО: Ваш предыдущий ответ содержал текст из контекста. "
    "Используйте ТОЛЬКО текст из [TARGET BLOCK]."
)


class ContextLeakageError(RuntimeError):
    pass


@lru_cache(maxsize=1)
def ensure_pandoc_available() -> None:
    try:
        pypandoc.get_pandoc_version()
    except OSError as exc:
        raise RuntimeError(
            "Pandoc не найден в текущем WSL runtime. Для штатного workflow установите его внутри WSL, "
            "например через: sudo apt-get install -y pandoc"
        ) from exc


def normalize_model_output(text: str) -> str:
    cleaned = text.strip()
    if cleaned.startswith("```"):
        first_newline = cleaned.find("\n")
        if first_newline != -1:
            cleaned = cleaned[first_newline + 1:].strip()
        else:
            cleaned = cleaned[3:].strip()
    if cleaned.endswith("```"):
        cleaned = cleaned[:-3].strip()
    return cleaned


def _normalize_context_text(text: str | None) -> str:
    if text is None:
        return "[контекст отсутствует]"
    cleaned = text.strip()
    return cleaned or "[контекст отсутствует]"


_CONTEXT_IMAGE_PLACEHOLDER_PATTERN = re.compile(r"\[\[DOCX_IMAGE_img_\d+\]\]")


def _strip_image_placeholders(text: str) -> str:
    """Remove DOCX image placeholder tokens from context strings.

    Image placeholders must not appear in context_before / context_after because
    the model consistently returns empty responses when it encounters them in the
    surrounding context (as opposed to the target block, where they must be
    preserved for later image reinsertion).
    """
    return _CONTEXT_IMAGE_PLACEHOLDER_PATTERN.sub("", text).strip()


def _strip_prompt_internal_tokens(text: str) -> str:
    without_images = _CONTEXT_IMAGE_PLACEHOLDER_PATTERN.sub("", text)
    without_markers = _PARAGRAPH_MARKER_PATTERN.sub("", without_images)
    return without_markers.strip()


def _should_passthrough_target(target_text: str) -> bool:
    stripped_target = target_text.strip()
    if not stripped_target:
        return True
    if _IMAGE_ONLY_TARGET_PATTERN.fullmatch(stripped_target):
        return True
    return not _strip_prompt_internal_tokens(target_text)


def _validate_prompt_inputs(target_text: str, context_before: str, context_after: str) -> list[str]:
    warnings: list[str] = []
    if not target_text.strip():
        warnings.append("empty_target_text")
    elif _IMAGE_ONLY_TARGET_PATTERN.fullmatch(target_text.strip()):
        warnings.append("image_only_target_text")
    elif not _strip_prompt_internal_tokens(target_text):
        warnings.append("placeholder_only_target_text")

    if not context_before.strip():
        warnings.append("empty_context_before")
    if not context_after.strip():
        warnings.append("empty_context_after")
    return warnings


def _build_standard_user_prompt(*, target_text: str, context_before: str, context_after: str) -> str:
    return (
        "Ниже передан целевой блок документа и соседний контекст.\n"
        "Используй соседний контекст только для понимания смысла, терминологии и связности.\n"
        "Редактируй только целевой блок и верни только его итоговый текст.\n\n"
        f"[CONTEXT BEFORE]\n{context_before}\n\n"
        f"[TARGET BLOCK]\n{target_text}\n\n"
        f"[CONTEXT AFTER]\n{context_after}"
    )


def _build_marker_preserving_user_prompt(*, target_text: str, context_before: str, context_after: str) -> str:
    return (
        "Ниже передан целевой блок документа с обязательными маркерами абзацев вида [[DOCX_PARA_...]].\n"
        "Сохрани каждый marker в точности, в том же количестве и порядке.\n"
        "Не удаляй, не дублируй и не переименовывай markers.\n"
        "Не объединяй абзацы между markers и не дели один marker на несколько абзацев.\n"
        "Редактируй только текст после каждого marker и верни весь блок целиком вместе с markers.\n"
        "Используй соседний контекст только для смысла и терминологии.\n\n"
        f"[CONTEXT BEFORE]\n{context_before}\n\n"
        f"[TARGET BLOCK WITH MARKERS]\n{target_text}\n\n"
        f"[CONTEXT AFTER]\n{context_after}"
    )


def _build_empty_response_recovery_user_prompt(*, target_text: str) -> str:
    return (
        "Предыдущая попытка вернула пустой ответ.\n"
        "Повтори обработку, но игнорируй любой внешний контекст и работай только с целевым блоком ниже.\n"
        "Сохрани весь смысл, структуру и факты блока.\n"
        "Верни только итоговый отредактированный текст блока без пояснений, без Markdown-обрамления и без пустого ответа.\n\n"
        f"[TARGET BLOCK ONLY]\n{target_text}"
    )


def _build_marker_recovery_user_prompt(*, target_text: str) -> str:
    return (
        "Предыдущая попытка нарушила контракт paragraph markers.\n"
        "Повтори обработку строго по правилам ниже.\n"
        "Сохрани каждый marker [[DOCX_PARA_...]] в исходном виде и порядке.\n"
        "Не удаляй markers, не добавляй новые, не меняй их местами.\n"
        "Каждому marker должен соответствовать ровно один абзац текста после него.\n"
        "Верни только итоговый блок целиком вместе с markers, без пояснений.\n\n"
        f"[TARGET BLOCK WITH MARKERS ONLY]\n{target_text}"
    )


def _split_marker_preserved_markdown(markdown: str, expected_paragraph_ids: Sequence[str]) -> list[str]:
    matches = list(_PARAGRAPH_MARKER_PATTERN.finditer(markdown))
    if not matches:
        raise RuntimeError("paragraph_marker_validation_failed:markers_missing")

    found_ids = [match.group(1) for match in matches]
    expected_ids = list(expected_paragraph_ids)
    if found_ids != expected_ids:
        raise RuntimeError("paragraph_marker_validation_failed:marker_order_or_identity")

    leading_text = markdown[: matches[0].start()].strip()
    if leading_text:
        raise RuntimeError("paragraph_marker_validation_failed:unexpected_prefix")

    paragraph_chunks: list[str] = []
    for index, match in enumerate(matches):
        content_end = matches[index + 1].start() if index + 1 < len(matches) else len(markdown)
        chunk = markdown[match.end() : content_end].strip()
        if not chunk:
            raise RuntimeError("paragraph_marker_validation_failed:empty_marker_chunk")
        if "\n\n" in chunk:
            raise RuntimeError("paragraph_marker_validation_failed:paragraph_split_detected")
        paragraph_chunks.append(chunk)
    return paragraph_chunks


def _strip_and_validate_paragraph_markers(markdown: str, expected_paragraph_ids: Sequence[str] | None, *, marker_mode: bool) -> str:
    if not marker_mode:
        return markdown
    if not expected_paragraph_ids:
        raise RuntimeError("paragraph_marker_validation_failed:missing_expected_ids")
    return "\n\n".join(_split_marker_preserved_markdown(markdown, expected_paragraph_ids))


def _normalize_leakage_comparison_text(text: str) -> str:
    return " ".join(match.group(0).lower() for match in _WORD_TOKEN_PATTERN.finditer(text))


def _detect_context_leakage(
    response_text: str,
    target_text: str,
    context_before: str,
    context_after: str,
    *,
    min_word_sequence: int = 6,
) -> str | None:
    response_tokens = list(_WORD_TOKEN_PATTERN.finditer(response_text))
    if len(response_tokens) < min_word_sequence:
        return None

    normalized_target = _normalize_leakage_comparison_text(target_text)
    normalized_contexts = [
        normalized_context
        for normalized_context in (
            _normalize_leakage_comparison_text(context_before),
            _normalize_leakage_comparison_text(context_after),
        )
        if normalized_context
    ]
    if not normalized_contexts:
        return None

    for start_index in range(0, len(response_tokens) - min_word_sequence + 1):
        end_index = start_index + min_word_sequence
        fragment = response_text[
            response_tokens[start_index].start() : response_tokens[end_index - 1].end()
        ]
        normalized_fragment = _normalize_leakage_comparison_text(fragment)
        if not normalized_fragment or normalized_fragment in normalized_target:
            continue
        if any(normalized_fragment in normalized_context for normalized_context in normalized_contexts):
            return fragment
    return None


def _trim_boundary_context_leakage(response_text: str, leaked_fragment: str) -> tuple[str, bool]:
    trimmed_response = response_text.strip()
    if not trimmed_response or leaked_fragment not in trimmed_response:
        return response_text, False

    matches = list(re.finditer(re.escape(leaked_fragment), trimmed_response))
    if not matches:
        return response_text, False
    if any(match.start() != 0 and match.end() != len(trimmed_response) for match in matches):
        return response_text, False

    updated_text = trimmed_response
    changed = False
    while updated_text.startswith(leaked_fragment):
        updated_text = updated_text[len(leaked_fragment) :].lstrip(" \t\r\n-–—,:;.!?")
        changed = True
    while updated_text.endswith(leaked_fragment):
        updated_text = updated_text[: -len(leaked_fragment)].rstrip(" \t\r\n-–—,:;.!?")
        changed = True

    if not changed or not updated_text:
        return response_text, False
    return updated_text, True


def _inject_context_leakage_retry_warning(request_kwargs: dict[str, object]) -> dict[str, object]:
    updated_request = dict(request_kwargs)
    payload = updated_request.get("input")
    if not isinstance(payload, list):
        return updated_request

    updated_payload: list[object] = []
    for index, message in enumerate(payload):
        if index != 1 or not isinstance(message, dict):
            updated_payload.append(message)
            continue

        updated_message = dict(message)
        content_items = list(updated_message.get("content", []))
        if content_items and isinstance(content_items[0], dict):
            updated_content = dict(content_items[0])
            text = updated_content.get("text")
            if isinstance(text, str) and _CONTEXT_LEAKAGE_RETRY_WARNING not in text:
                updated_content["text"] = f"{_CONTEXT_LEAKAGE_RETRY_WARNING}\n\n{text}"
                content_items[0] = updated_content
        updated_message["content"] = content_items
        updated_payload.append(updated_message)

    updated_request["input"] = updated_payload
    return updated_request


def _finalize_generated_markdown(
    markdown: str,
    *,
    target_text: str,
    context_before: str,
    context_after: str,
    expected_paragraph_ids: Sequence[str] | None,
    marker_mode: bool,
    allow_persistent_context_leakage: bool,
) -> str:
    cleaned_markdown = _strip_and_validate_paragraph_markers(
        markdown,
        expected_paragraph_ids,
        marker_mode=marker_mode,
    )
    leaked_fragment = _detect_context_leakage(
        cleaned_markdown,
        target_text,
        context_before,
        context_after,
    )
    if leaked_fragment is None:
        return cleaned_markdown

    trimmed_markdown, was_trimmed = _trim_boundary_context_leakage(cleaned_markdown, leaked_fragment)
    if was_trimmed:
        return trimmed_markdown

    if allow_persistent_context_leakage:
        log_event(
            logging.WARNING,
            "context_leakage_persisted",
            "После последней попытки генерации сохранилась verbatim-протечка текста из соседнего контекста; возвращаю fail-open результат.",
            leaked_fragment=leaked_fragment,
            target_chars=len(target_text),
            marker_mode=marker_mode,
        )
        return cleaned_markdown

    raise ContextLeakageError(f"context_leakage_detected:{leaked_fragment}")


def _read_response_field(value: object, field_name: str) -> object:
    if isinstance(value, Mapping):
        return value.get(field_name)
    return getattr(value, field_name, None)


def _coerce_response_text_value(value: object) -> str | None:
    if isinstance(value, str):
        return value
    nested_value = _read_response_field(value, "value")
    if isinstance(nested_value, str):
        return nested_value
    return None


def _extract_text_from_content_item(content_item: object) -> tuple[str | None, bool]:
    item_type = _read_response_field(content_item, "type")
    if item_type not in _SUPPORTED_RESPONSE_TEXT_TYPES:
        return None, False
    text_value = _coerce_response_text_value(_read_response_field(content_item, "text"))
    if text_value is None:
        raise RuntimeError("Модель вернула ответ в неподдерживаемом формате (unsupported_response_shape).")
    return text_value, True


def _extract_response_output_text(response: object) -> str:
    output_text = getattr(response, "output_text", None)
    raw_output_text: str | None = None
    if output_text is not None:
        if not isinstance(output_text, str):
            raise RuntimeError("Модель вернула ответ в неподдерживаемом формате (unsupported_response_shape).")
        if output_text.strip():
            return output_text
        raw_output_text = output_text

    output_items = _read_response_field(response, "output")
    if output_items is None:
        return raw_output_text or ""
    if isinstance(output_items, (str, bytes)) or not isinstance(output_items, Iterable):
        raise RuntimeError("Модель вернула ответ в неподдерживаемом формате (unsupported_response_shape).")

    collected_texts: list[str] = []
    saw_output_items = False
    saw_supported_text_shape = False
    saw_empty_content_container = False

    for output_item in output_items:
        saw_output_items = True

        direct_text, direct_supported = _extract_text_from_content_item(output_item)
        if direct_supported:
            saw_supported_text_shape = True
            if direct_text:
                collected_texts.append(direct_text)
            continue

        content_items = _read_response_field(output_item, "content")
        if content_items is None:
            continue
        if isinstance(content_items, (str, bytes)) or not isinstance(content_items, Iterable):
            raise RuntimeError("Модель вернула ответ в неподдерживаемом формате (unsupported_response_shape).")

        content_list = list(content_items)
        if not content_list:
            saw_empty_content_container = True
            continue

        for content_item in content_list:
            extracted_text, supported = _extract_text_from_content_item(content_item)
            if not supported:
                continue
            saw_supported_text_shape = True
            if extracted_text:
                collected_texts.append(extracted_text)

    if collected_texts:
        return "\n".join(collected_texts)
    if raw_output_text is not None:
        return raw_output_text
    if not saw_output_items or saw_supported_text_shape or saw_empty_content_container:
        return ""
    raise RuntimeError("Модель вернула ответ в неподдерживаемом формате (unsupported_response_shape).")


def _log_empty_response_shape(response: object, raw_output_text: str, *, error_code: str) -> None:
    output_items = _read_response_field(response, "output")
    output_items_len = len(output_items) if isinstance(output_items, Sized) else None

    first_item_summary: dict[str, object] | None = None
    if isinstance(output_items, Iterable) and not isinstance(output_items, (str, bytes)):
        for item in output_items:
            item_type = _read_response_field(item, "type")
            refusal = _read_response_field(item, "refusal")
            status = _read_response_field(item, "status")
            content_items = _read_response_field(item, "content")
            content_types: list[str] = []
            if isinstance(content_items, Iterable) and not isinstance(content_items, (str, bytes)):
                content_types = [
                    str(_read_response_field(c, "type") or type(c).__name__)
                    for c in content_items
                ]
            first_item_summary = {
                "type": item_type,
                "refusal": refusal,
                "status": status,
                "content_types": content_types,
            }
            break

    log_event(
        logging.WARNING,
        "model_empty_response_shape",
        "Модель вернула пустой или схлопнувшийся текстовый ответ",
        error_code=error_code,
        has_output_text_attr=getattr(response, "output_text", None) is not None,
        raw_output_len=len(raw_output_text),
        output_items_type=type(output_items).__name__ if output_items is not None else "None",
        output_items_len=output_items_len,
        response_status=_read_response_field(response, "status"),
        first_output_item=first_item_summary,
    )


def _extract_normalized_markdown(response: object) -> str:
    response_status = _read_response_field(response, "status")
    if response_status == "incomplete":
        _log_empty_response_shape(response, "", error_code="incomplete_response")
        raise RuntimeError("Модель не завершила генерацию (incomplete_response).")
    if isinstance(response_status, str) and response_status != "completed":
        _log_empty_response_shape(response, "", error_code="non_completed_response")
        raise RuntimeError(f"Модель вернула неожиданный статус ответа: {response_status} (non_completed_response).")

    raw_output_text = _extract_response_output_text(response)
    markdown = normalize_model_output(raw_output_text)
    if markdown:
        return markdown
    error_code = "collapsed_output" if raw_output_text else "empty_response"
    _log_empty_response_shape(response, raw_output_text, error_code=error_code)
    if raw_output_text:
        raise RuntimeError("Модель вернула ответ, который схлопнулся после нормализации (collapsed_output).")
    raise RuntimeError("Модель вернула пустой ответ (empty_response).")


def _call_responses_create(client: "OpenAI", request_kwargs: dict[str, Any]) -> object:
    return cast(Any, client.responses).create(**request_kwargs)


def _estimate_max_output_tokens(target_text: str) -> int:
    estimated_output_tokens = max((len(target_text) // 3) * 4, 512)
    return min(estimated_output_tokens, 16384)


def _build_request_kwargs(*, model: str, system_prompt: str, user_prompt: str, target_text: str) -> dict[str, object]:
    payload: Any = [
        {
            "role": "system",
            "content": [{"type": "input_text", "text": system_prompt}],
        },
        {
            "role": "user",
            "content": [{"type": "input_text", "text": user_prompt}],
        },
    ]
    return {
        "model": model,
        "input": payload,
        "max_output_tokens": _estimate_max_output_tokens(target_text),
    }


def _boost_request_output_budget(
    request_kwargs: dict[str, object],
    *,
    minimum_tokens: int,
) -> dict[str, object]:
    boosted_request = dict(request_kwargs)
    current_value = boosted_request.get("max_output_tokens")
    if isinstance(current_value, int) and not isinstance(current_value, bool):
        boosted_request["max_output_tokens"] = min(max(current_value * 2, minimum_tokens), 16384)
        return boosted_request
    boosted_request["max_output_tokens"] = min(max(minimum_tokens, 512), 16384)
    return boosted_request


def _call_markdown_request_with_sdk_fallback(client: "OpenAI", request_kwargs: dict[str, object]) -> tuple[str, bool]:
    max_output_tokens_removed = False
    try:
        response = _call_responses_create(client, cast(dict[str, Any], request_kwargs))
        return _extract_normalized_markdown(response), max_output_tokens_removed
    except TypeError as exc:
        if "max_output_tokens" not in str(exc) or "max_output_tokens" not in request_kwargs:
            raise
        request_kwargs = dict(request_kwargs)
        request_kwargs.pop("max_output_tokens", None)
        max_output_tokens_removed = True
        response = _call_responses_create(client, cast(dict[str, Any], request_kwargs))
        return _extract_normalized_markdown(response), max_output_tokens_removed


def _recover_from_persistent_empty_response(
    *,
    client: "OpenAI",
    model: str,
    system_prompt: str,
    target_text: str,
    expected_paragraph_ids: Sequence[str] | None = None,
    marker_mode: bool = False,
    minimum_output_tokens: int | None = None,
) -> str:
    log_event(
        logging.WARNING,
        "markdown_empty_response_recovery_started",
        "Обычные retry исчерпаны; запускаю recovery-вызов без соседнего контекста.",
        model=model,
        target_chars=len(target_text),
    )
    request_kwargs = _build_request_kwargs(
        model=model,
        system_prompt=system_prompt,
        user_prompt=(
            _build_marker_recovery_user_prompt(target_text=target_text)
            if marker_mode
            else _build_empty_response_recovery_user_prompt(target_text=target_text)
        ),
        target_text=target_text,
    )
    if minimum_output_tokens is not None:
        request_kwargs = _boost_request_output_budget(
            request_kwargs,
            minimum_tokens=minimum_output_tokens,
        )
    markdown = _call_markdown_request_with_sdk_fallback(client, request_kwargs)[0]
    return _strip_and_validate_paragraph_markers(
        markdown,
        expected_paragraph_ids,
        marker_mode=marker_mode,
    )


def _is_incomplete_response_error(exc: Exception) -> bool:
    return isinstance(exc, RuntimeError) and "incomplete_response" in str(exc)


def _can_fallback_to_source_text_after_incomplete_response(target_text: str) -> bool:
    return bool(target_text.strip())


def _is_retryable_empty_generation_error(exc: Exception) -> bool:
    return isinstance(exc, RuntimeError) and (
        "empty_response" in str(exc) or "collapsed_output" in str(exc) or "incomplete_response" in str(exc)
    )


def _is_retryable_marker_validation_error(exc: Exception) -> bool:
    return isinstance(exc, RuntimeError) and "paragraph_marker_validation_failed" in str(exc)


def _is_retryable_context_leakage_error(exc: Exception) -> bool:
    return isinstance(exc, ContextLeakageError)


def generate_markdown_block(
    client: "OpenAI",
    model: str,
    system_prompt: str,
    target_text: str,
    context_before: str,
    context_after: str,
    max_retries: int,
    expected_paragraph_ids: Sequence[str] | None = None,
    marker_mode: bool = False,
) -> str:
    if isinstance(max_retries, bool) or not isinstance(max_retries, int):
        raise TypeError("max_retries должен быть целым числом.")
    if max_retries < 1:
        raise ValueError("max_retries должен быть не меньше 1.")

    if _should_passthrough_target(target_text):
        log_event(
            logging.WARNING,
            "image_only_target_passthrough",
            "Целевой блок не содержит редактируемого текста; возвращаю его без вызова модели.",
            target_chars=len(target_text),
            marker_mode=marker_mode,
        )
        return target_text

    context_before_text = _normalize_context_text(_strip_image_placeholders(context_before))
    context_after_text = _normalize_context_text(_strip_image_placeholders(context_after))
    prompt_warnings = _validate_prompt_inputs(target_text, context_before_text, context_after_text)
    if prompt_warnings:
        log_event(
            logging.WARNING,
            "prompt_quality_warning",
            "Входные данные prompt содержат потенциально проблемный shape.",
            warnings=prompt_warnings,
            target_chars=len(target_text),
            context_before_chars=len(context_before_text),
            context_after_chars=len(context_after_text),
            marker_mode=marker_mode,
        )

    request_kwargs = _build_request_kwargs(
        model=model,
        system_prompt=system_prompt,
        user_prompt=(
            _build_marker_preserving_user_prompt(
                target_text=target_text,
                context_before=context_before_text,
                context_after=context_after_text,
            )
            if marker_mode
            else _build_standard_user_prompt(
                target_text=target_text,
                context_before=context_before_text,
                context_after=context_after_text,
            )
        ),
        target_text=target_text,
    )
    target_text_for_leakage = _strip_and_validate_paragraph_markers(
        target_text,
        expected_paragraph_ids,
        marker_mode=marker_mode,
    )
    last_exception: Exception | None = None

    for attempt in range(1, max_retries + 1):
        try:
            markdown = _call_markdown_request_with_sdk_fallback(client, request_kwargs)[0]
            return _finalize_generated_markdown(
                markdown,
                target_text=target_text_for_leakage,
                context_before=context_before_text,
                context_after=context_after_text,
                expected_paragraph_ids=expected_paragraph_ids,
                marker_mode=marker_mode,
                allow_persistent_context_leakage=attempt >= max_retries,
            )
        except Exception as exc:
            last_exception = exc
            should_retry = attempt < max_retries and (
                is_retryable_error(exc)
                or _is_retryable_empty_generation_error(exc)
                or _is_retryable_marker_validation_error(exc)
                or _is_retryable_context_leakage_error(exc)
            )
            if not should_retry:
                break
            if _is_incomplete_response_error(exc):
                request_kwargs = _boost_request_output_budget(
                    request_kwargs,
                    minimum_tokens=_INCOMPLETE_RESPONSE_RETRY_MIN_OUTPUT_TOKENS,
                )
            if _is_retryable_context_leakage_error(exc):
                request_kwargs = _inject_context_leakage_retry_warning(request_kwargs)
            time.sleep(min(2 ** (attempt - 1), 8))

    if last_exception is not None and (
        _is_retryable_empty_generation_error(last_exception)
        or _is_retryable_marker_validation_error(last_exception)
    ):
        try:
            return _recover_from_persistent_empty_response(
                client=client,
                model=model,
                system_prompt=system_prompt,
                target_text=target_text,
                expected_paragraph_ids=expected_paragraph_ids,
                marker_mode=marker_mode,
                minimum_output_tokens=(
                    _INCOMPLETE_RESPONSE_RECOVERY_MIN_OUTPUT_TOKENS
                    if _is_incomplete_response_error(last_exception)
                    else None
                ),
            )
        except Exception as recovery_exc:
            if _is_incomplete_response_error(recovery_exc) and _can_fallback_to_source_text_after_incomplete_response(target_text):
                log_event(
                    logging.WARNING,
                    "markdown_incomplete_response_source_fallback",
                    "Recovery для блока снова завершился incomplete_response; сохраняю исходный текст блока как controlled fallback.",
                    model=model,
                    target_chars=len(target_text),
                    marker_mode=marker_mode,
                )
                return target_text
            if _is_retryable_empty_generation_error(recovery_exc) or _is_retryable_marker_validation_error(recovery_exc):
                raise recovery_exc
            raise recovery_exc

    if last_exception is not None:
        raise last_exception

    raise RuntimeError("Не удалось получить ответ модели.")


_THEME_RELATIONSHIP_TYPE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme"
)
_DRAWINGML_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def convert_markdown_to_docx_bytes(
    markdown_text: str,
    *,
    body_font: str | None = None,
    heading_font: str | None = None,
) -> bytes:
    """Convert *markdown_text* to DOCX bytes via Pandoc.

    *body_font* and *heading_font* are optional overrides for the reference
    document. Body-facing styles are updated directly because python-docx writes
    them as explicit ``w:rFonts`` values; heading styles additionally require a
    theme patch because Word gives ``w:asciiTheme=majorHAnsi`` precedence over
    the direct font name. When both are ``None`` (the default) the python-docx
    built-in theme is left unchanged.
    """
    ensure_pandoc_available()
    try:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            markdown_path = temp_path / "result.md"
            docx_path = temp_path / "result.docx"
            reference_docx_path = temp_path / "reference.docx"
            markdown_path.write_text(markdown_text, encoding="utf-8")
            _build_reference_docx(reference_docx_path, body_font=body_font, heading_font=heading_font)
            pypandoc.convert_file(
                str(markdown_path),
                to="docx",
                format="md",
                outputfile=str(docx_path),
                extra_args=[f"--reference-doc={reference_docx_path}"],
            )
            return docx_path.read_bytes()
    except Exception as exc:
        raise RuntimeError(f"Ошибка при сборке DOCX: {exc}") from exc


def _patch_reference_theme_fonts(
    reference_document: "DocxDocument",
    *,
    body_font: str | None,
    heading_font: str | None,
) -> None:
    """Overwrite the major/minor font slots in the reference document's theme.

    Word resolves ``w:asciiTheme="majorHAnsi"`` (used by all built-in Heading
    styles) and ``w:asciiTheme="minorHAnsi"`` (body/list/caption) by looking
    up the document's embedded ``theme1.xml``.  python-docx's default template
    maps those slots to Calibri (major) and Cambria (minor).

    Patching the theme here means every ``w:asciiTheme`` reference in every
    style automatically picks up the configured font **without** touching
    individual style ``w:rFonts`` elements — this is the OOXML-idiomatic
    approach and the only reliable way to override heading fonts given that
    python-docx's ``Style.font.name`` setter leaves ``w:asciiTheme`` intact.

    Called only when at least one font is configured; both arguments may be
    ``None`` to skip their respective slot.
    """
    try:
        theme_part = reference_document.part.part_related_by(_THEME_RELATIONSHIP_TYPE)
    except KeyError:
        return  # Template has no theme part — nothing to patch.

    root = etree.fromstring(theme_part.blob)

    if heading_font is not None:
        for el in root.findall(f".//{{{_DRAWINGML_NS}}}majorFont/{{{_DRAWINGML_NS}}}latin"):
            el.set("typeface", heading_font)

    if body_font is not None:
        for el in root.findall(f".//{{{_DRAWINGML_NS}}}minorFont/{{{_DRAWINGML_NS}}}latin"):
            el.set("typeface", body_font)

    theme_part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _build_reference_docx(
    reference_docx_path: Path,
    *,
    body_font: str | None = None,
    heading_font: str | None = None,
) -> None:
    reference_document = Document()
    styles = reference_document.styles
    effective_body_font = body_font or "Aptos"
    effective_heading_font = heading_font or "Aptos Display"

    body_baseline = {
        "font_name": effective_body_font,
        "font_size": 11,
        "space_after": 8,
        "line_spacing": 1.15,
    }

    _configure_paragraph_style(styles["Normal"], **body_baseline)

    if "Body Text" in styles:
        _configure_paragraph_style(styles["Body Text"], **body_baseline)

    heading_specs = (
        ("Heading 1", 18, 18, 8),
        ("Heading 2", 16, 16, 7),
        ("Heading 3", 14, 14, 6),
        ("Heading 4", 13, 12, 5),
        ("Heading 5", 12, 10, 4),
        ("Heading 6", 11, 8, 3),
    )
    for style_name, font_size, space_before, space_after in heading_specs:
        if style_name not in styles:
            continue
        _configure_paragraph_style(
            styles[style_name],
            font_name=effective_heading_font,
            font_size=font_size,
            bold=True,
            space_before=space_before,
            space_after=space_after,
            line_spacing=1.1,
            keep_with_next=True,
        )

    if "Caption" in styles:
        _configure_paragraph_style(
            styles["Caption"],
            font_name=effective_body_font,
            font_size=10,
            italic=True,
            space_before=4,
            space_after=10,
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
        )

    if "List Paragraph" in styles:
        _configure_paragraph_style(
            styles["List Paragraph"],
            font_name=effective_body_font,
            font_size=11,
            space_before=0,
            space_after=4,
            line_spacing=1.1,
        )

    if "Table Grid" in styles:
        table_grid_style = cast(Any, styles["Table Grid"])
        table_grid_style.font.name = effective_body_font
        table_grid_style.font.size = Pt(10)

    _ensure_reference_numbering_definitions(reference_document)

    # Patch theme fonts only when explicitly configured. This is required for
    # heading styles, whose built-in w:asciiTheme bindings outrank the direct
    # font name, and keeps theme-bound fallback slots aligned with explicit
    # style overrides.
    if body_font is not None or heading_font is not None:
        _patch_reference_theme_fonts(reference_document, body_font=body_font, heading_font=heading_font)

    reference_document.save(str(reference_docx_path))


def _ensure_reference_numbering_definitions(reference_document: "DocxDocument") -> None:
    numbering_part = reference_document.part.numbering_part
    numbering = numbering_part.element
    baseline_specs = (
        {
            "num_fmt": "decimal",
            "level_text_patterns": ("%1.", "%1.%2.", "%1.%2.%3."),
        },
        {
            "num_fmt": "bullet",
            "level_text_patterns": (chr(0x2022), chr(0x25E6), chr(0x25AA)),
        },
    )

    for spec in baseline_specs:
        baseline_abstract_num = _find_reference_baseline_abstract_num(
            numbering,
            num_fmt=spec["num_fmt"],
            level_text_patterns=spec["level_text_patterns"],
        )
        if baseline_abstract_num is None:
            abstract_num_id = _next_numbering_id(numbering, "w:abstractNum", "abstractNumId")
            _append_multilevel_numbering_definition(
                numbering,
                abstract_num_id=abstract_num_id,
                num_fmt=spec["num_fmt"],
                level_text_patterns=spec["level_text_patterns"],
            )
            baseline_abstract_num = _find_abstract_num_by_id(numbering, abstract_num_id)
            if baseline_abstract_num is None:
                raise RuntimeError("Не удалось создать baseline numbering definition.")

        abstract_num_id = int(baseline_abstract_num.get(qn("w:abstractNumId")))
        if not _num_instance_exists(numbering, abstract_num_id=abstract_num_id):
            num_id = _next_numbering_id(numbering, "w:num", "numId")
            _append_num_instance(numbering, num_id=num_id, abstract_num_id=abstract_num_id)


def _find_reference_baseline_abstract_num(numbering, *, num_fmt: str, level_text_patterns: tuple[str, ...]):
    for abstract_num in numbering.xpath('./*[local-name()="abstractNum"]'):
        if _abstract_num_matches_reference_baseline(
            abstract_num,
            num_fmt=num_fmt,
            level_text_patterns=level_text_patterns,
        ):
            return abstract_num
    return None


def _find_abstract_num_by_id(numbering, abstract_num_id: int):
    matches = numbering.xpath(
        f'./*[local-name()="abstractNum" and @*[local-name()="abstractNumId"]="{abstract_num_id}"]'
    )
    return matches[0] if matches else None


def _iter_num_instances(numbering):
    return numbering.xpath('./*[local-name()="num"]')


def _num_instance_abstract_num_id(num_instance) -> str | None:
    abstract_num_id_values = num_instance.xpath(
        './*[local-name()="abstractNumId"]/@*[local-name()="val"]'
    )
    if not abstract_num_id_values:
        return None
    return str(abstract_num_id_values[0])


def _abstract_num_matches_reference_baseline(abstract_num, *, num_fmt: str, level_text_patterns: tuple[str, ...]) -> bool:
    levels = abstract_num.xpath('./*[local-name()="lvl"]')
    if len(levels) != len(level_text_patterns):
        return False

    for ilvl, level_text in enumerate(level_text_patterns):
        level_matches = [level for level in levels if level.get(qn("w:ilvl")) == str(ilvl)]
        if len(level_matches) != 1:
            return False
        if not _level_matches_reference_baseline(level_matches[0], num_fmt=num_fmt, level_text=level_text, ilvl=ilvl):
            return False
    return True


def _level_matches_reference_baseline(level, *, num_fmt: str, level_text: str, ilvl: int) -> bool:
    num_fmt_values = level.xpath('./*[local-name()="numFmt"]/@*[local-name()="val"]')
    level_text_values = level.xpath('./*[local-name()="lvlText"]/@*[local-name()="val"]')
    left_values = level.xpath('./*[local-name()="pPr"]/*[local-name()="ind"]/@*[local-name()="left"]')
    hanging_values = level.xpath('./*[local-name()="pPr"]/*[local-name()="ind"]/@*[local-name()="hanging"]')
    after_values = level.xpath('./*[local-name()="pPr"]/*[local-name()="spacing"]/@*[local-name()="after"]')
    line_values = level.xpath('./*[local-name()="pPr"]/*[local-name()="spacing"]/@*[local-name()="line"]')
    line_rule_values = level.xpath('./*[local-name()="pPr"]/*[local-name()="spacing"]/@*[local-name()="lineRule"]')
    ascii_fonts = level.xpath('./*[local-name()="rPr"]/*[local-name()="rFonts"]/@*[local-name()="ascii"]')
    hansi_fonts = level.xpath('./*[local-name()="rPr"]/*[local-name()="rFonts"]/@*[local-name()="hAnsi"]')
    cs_fonts = level.xpath('./*[local-name()="rPr"]/*[local-name()="rFonts"]/@*[local-name()="cs"]')

    return (
        num_fmt_values == [num_fmt]
        and level_text_values == [level_text]
        and left_values == [str(720 + (ilvl * 360))]
        and hanging_values == ["360"]
        and after_values == ["80"]
        and line_values == ["264"]
        and line_rule_values == ["auto"]
        and ascii_fonts == ["Aptos"]
        and hansi_fonts == ["Aptos"]
        and cs_fonts == ["Aptos"]
    )


def _num_instance_exists(numbering, *, abstract_num_id: int) -> bool:
    expected_abstract_num_id = str(abstract_num_id)
    return any(
        _num_instance_abstract_num_id(num_instance) == expected_abstract_num_id
        for num_instance in _iter_num_instances(numbering)
    )


def _next_numbering_id(numbering, element_name: str, attr_name: str) -> int:
    existing_ids = []
    local_name = element_name.split(":", 1)[1]
    for element in numbering.xpath(f'./*[local-name()="{local_name}"]'):
        value = element.get(qn(f"w:{attr_name}"))
        if value is not None:
            existing_ids.append(int(value))
    return (max(existing_ids) + 1) if existing_ids else 0


def _append_multilevel_numbering_definition(numbering, *, abstract_num_id: int, num_fmt: str, level_text_patterns: tuple[str, str, str]) -> None:
    abstract_num = OxmlElement("w:abstractNum")
    abstract_num.set(qn("w:abstractNumId"), str(abstract_num_id))

    nsid = OxmlElement("w:nsid")
    nsid.set(qn("w:val"), f"{abstract_num_id + 1:08X}")
    abstract_num.append(nsid)

    multi_level_type = OxmlElement("w:multiLevelType")
    multi_level_type.set(qn("w:val"), "multilevel")
    abstract_num.append(multi_level_type)

    template_code = OxmlElement("w:tmpl")
    template_code.set(qn("w:val"), f"{abstract_num_id + 257:08X}")
    abstract_num.append(template_code)

    for ilvl, level_text in enumerate(level_text_patterns):
        abstract_num.append(_build_numbering_level(ilvl=ilvl, num_fmt=num_fmt, level_text=level_text))

    numbering.append(abstract_num)


def _build_numbering_level(*, ilvl: int, num_fmt: str, level_text: str):
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), str(ilvl))

    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)

    num_fmt_element = OxmlElement("w:numFmt")
    num_fmt_element.set(qn("w:val"), num_fmt)
    level.append(num_fmt_element)

    level_text_element = OxmlElement("w:lvlText")
    level_text_element.set(qn("w:val"), level_text)
    level.append(level_text_element)

    level_jc = OxmlElement("w:lvlJc")
    level_jc.set(qn("w:val"), "left")
    level.append(level_jc)

    paragraph_properties = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(720 + (ilvl * 360)))
    ind.set(qn("w:hanging"), "360")
    paragraph_properties.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "264")
    spacing.set(qn("w:lineRule"), "auto")
    paragraph_properties.append(spacing)
    level.append(paragraph_properties)

    run_properties = OxmlElement("w:rPr")
    run_fonts = OxmlElement("w:rFonts")
    run_fonts.set(qn("w:ascii"), "Aptos")
    run_fonts.set(qn("w:hAnsi"), "Aptos")
    run_fonts.set(qn("w:cs"), "Aptos")
    run_properties.append(run_fonts)
    level.append(run_properties)

    return level


def _append_num_instance(numbering, *, num_id: int, abstract_num_id: int) -> None:
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id_element = OxmlElement("w:abstractNumId")
    abstract_num_id_element.set(qn("w:val"), str(abstract_num_id))
    num.append(abstract_num_id_element)
    numbering.append(num)


def _configure_paragraph_style(
    style,
    *,
    font_name: str,
    font_size: int,
    bold: bool | None = None,
    italic: bool | None = None,
    space_before: int | None = None,
    space_after: int | None = None,
    line_spacing: float | None = None,
    keep_with_next: bool | None = None,
    alignment=None,
) -> None:
    style.font.name = font_name
    style.font.size = Pt(font_size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic

    paragraph_format = style.paragraph_format
    if space_before is not None:
        paragraph_format.space_before = Pt(space_before)
    if space_after is not None:
        paragraph_format.space_after = Pt(space_after)
    if line_spacing is not None:
        paragraph_format.line_spacing = line_spacing
    if keep_with_next is not None:
        paragraph_format.keep_with_next = keep_with_next
    if alignment is not None:
        paragraph_format.alignment = alignment


def build_output_filename(filename: str) -> str:
    return f"{Path(filename).stem}_edited.docx"


def build_markdown_filename(filename: str) -> str:
    return f"{Path(filename).stem}_edited.md"
