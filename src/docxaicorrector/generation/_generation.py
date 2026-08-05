import logging
import re
import tempfile
import time
from collections.abc import Iterable, Sequence, Sized
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from typing import TYPE_CHECKING, Any, cast

import lxml.etree as etree
import pypandoc
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from docxaicorrector.core.logger import log_event
from docxaicorrector.core.model_accounting import (
    PARAGRAPH_DISPOSITION_STATUSES,
    STAGE_TEXT_GENERATION,
    record_model_call_usage,
    record_model_output_discarded,
    record_paragraph_disposition,
    record_retry_attempt,
)
from docxaicorrector.image.shared import call_responses_create_with_retry, extract_unsupported_parameter_name, is_retryable_error
from docxaicorrector.generation.marker_attempt_capture import capture_rejected_marker_attempt
from docxaicorrector.generation.openai_response_utils import collect_response_text_traversal, read_response_field

if TYPE_CHECKING:
    from docx.document import Document as DocxDocument
    from openai import OpenAI


_PARAGRAPH_MARKER_PATTERN = re.compile(r"\[\[DOCX_PARA_([A-Za-z0-9_]+)\]\]")
_IMAGE_ONLY_TARGET_PATTERN = re.compile(r"^(?:\s*\[\[DOCX_IMAGE_img_\d+\]\]\s*)+$")
_WORD_TOKEN_PATTERN = re.compile(r"\w+(?:[-']\w+)*", re.UNICODE)
_INLINE_HTML_SUP_PATTERN = re.compile(r"<sup>(.*?)</sup>", re.IGNORECASE | re.DOTALL)
_INLINE_HTML_SUB_PATTERN = re.compile(r"<sub>(.*?)</sub>", re.IGNORECASE | re.DOTALL)
_INLINE_HTML_BREAK_PATTERN = re.compile(r"<br\s*/?>", re.IGNORECASE)
_INLINE_HTML_UNDERLINE_PATTERN = re.compile(
    r"(?P<prefix>\\*!?)<u>(?P<content>.*?)</u>", re.IGNORECASE | re.DOTALL
)
_PANDOC_SPAN_SPECIAL_PATTERN = re.compile(r"[\\\[\]]")
_NARRATION_INTERNAL_PLACEHOLDER_PATTERN = re.compile(r"\[\[DOCX_[A-Za-z0-9_]+\]\]")
_NARRATION_MARKDOWN_LINK_PATTERN = re.compile(r"\[([^\]\n]+)\]\(([^)\n]+)\)")
_NARRATION_HEADING_PATTERN = re.compile(r"^\s{0,3}#{1,6}\s*(.*)$")
_NARRATION_TAGGED_HEADING_PATTERN = re.compile(r"^(\s*(?:\[[^\]\n]+\]\s*)+)#{1,6}\s*(.*)$")
_NARRATION_BLOCKQUOTE_PATTERN = re.compile(r"^\s{0,3}>\s?(.*)$")
# A list marker is not speech. Markdown's own markers were already dropped; the printed
# bullet GLYPH was not, and the audiobook prompt's rule against lists does not bind the
# model — measured on the 2026-08-04 Money & Sustainability narration artifact, 116 of them
# reached the text a TTS engine would read aloud, 113 at the start of their own line and 3
# behind an ElevenLabs tag. The glyph set is the repository's existing bullet lexicon
# (``output_validation._BULLET_GLYPH_PATTERN``, which produced that same count of 116), not
# a new one: of its four glyphs only ``•`` occurs in the artifact, and the other three are
# carried because they are the same form, not because a book showed them. This is not a
# guess about structure — a bullet glyph is simply not speech, and the item's own text is
# kept.
# A separator after the glyph is required, so a glyph welded inside a token ("4●5") is left
# alone, exactly as ``_WELDED_BULLET_GLYPH_PATTERN`` leaves it alone in the quality gate.
_NARRATION_LIST_MARKER = r"(?:[-*+●•◦‣]\s+|\d+\.\s+)"
_NARRATION_LIST_PATTERN = re.compile(r"^\s{0,3}" + _NARRATION_LIST_MARKER + r"(.*)$")
# The same marker behind an ElevenLabs tag prefix, which the narration keeps: "[serious] • …"
# is the tagged twin of "• …" and has to lose the glyph the same way, or the tagged lines
# would be the only ones a listener hears a bullet in.
_NARRATION_TAGGED_LIST_PATTERN = re.compile(
    r"^(\s*(?:\[[^\]\n]+\]\s*)+)" + _NARRATION_LIST_MARKER + r"(.*)$"
)
_NARRATION_STRONG_PATTERN = re.compile(r"(\*\*|__)(?=\S)(.+?)(?<=\S)\1")
_NARRATION_EMPHASIS_PATTERN = re.compile(r"(\*|_)(?=\S)(.+?)(?<=\S)\1")
_NARRATION_RAW_URL_PATTERN = re.compile(r"(?:https?://\S+|www\.\S+)", re.IGNORECASE)
_NARRATION_INTERNAL_WHITESPACE_PATTERN = re.compile(r"[\t ]+")
# A returned marker chunk that keeps this little of its own source paragraph is a stub,
# not an edit or a translation. Calibrated on the 1761 recorded pairs of the 2026-08-03
# literary-edit run: the 7 "(Пусто)" stubs sit at 0.008–0.135, the next real edit above
# them at 0.207, and every one of the 609 genuine edits stays far above the floor.
_COLLAPSED_MARKER_CHUNK_RATIO = 0.15
# Below this the source paragraph is itself a fragment ("г.", ".htm.") where a length
# ratio carries no signal.
_COLLAPSED_MARKER_CHUNK_MIN_SOURCE_CHARS = 40
# A neighbour — on EITHER side — counts as the one that swallowed the collapsed paragraph
# when it grew by at least this share of what the collapsed paragraph lost.
_ABSORBING_NEIGHBOUR_RATIO = 0.5
_INCOMPLETE_RESPONSE_RETRY_MIN_OUTPUT_TOKENS = 1024
_INCOMPLETE_RESPONSE_RECOVERY_MIN_OUTPUT_TOKENS = 1536
_CONTEXT_LEAKAGE_RETRY_WARNING = (
    "IMPORTANT: Your previous answer included text from the surrounding context. "
    "Use ONLY the text that belongs to [TARGET BLOCK]."
)


class ContextLeakageError(RuntimeError):
    pass


# --- spec 056 E: what happened to ONE paragraph, said out loud ----------------------
# A block used to be all-or-nothing, so one paragraph the model could not answer for
# discarded every good translation next to it. These four statuses are the whole
# vocabulary; anything a paragraph can end up as is one of them, and the counts are
# published per run (``model_accounting.paragraph_disposition_counts``).
PARAGRAPH_STATUS_ACCEPTED = "accepted"
"""The model's own text for this paragraph is what ships."""
PARAGRAPH_STATUS_OMITTED = "omitted"
"""The model returned nothing for it: the source stands in the DOCX, the narration skips it."""
PARAGRAPH_STATUS_SOURCE_RESTORED = "source_restored"
"""The model's text was discarded (a visible merge) and this paragraph's source re-instated."""
PARAGRAPH_STATUS_RETRY_REQUIRED = "retry_required"
"""Unresolved: there is retry budget left, so the block is asked again. Never ships."""
# A status assigned here but absent from the report's vocabulary would be counted into a
# bucket nobody documented, so the two lists are held together at import time rather than
# by a comment. ``tests/test_generation.py`` asserts the same equality from the outside.
assert sorted(
    (
        PARAGRAPH_STATUS_ACCEPTED,
        PARAGRAPH_STATUS_OMITTED,
        PARAGRAPH_STATUS_SOURCE_RESTORED,
        PARAGRAPH_STATUS_RETRY_REQUIRED,
    )
) == sorted(PARAGRAPH_DISPOSITION_STATUSES)


@dataclass(frozen=True)
class ParagraphDisposition:
    """One paragraph of a marker-preserved block, and what became of it."""

    paragraph_id: str
    text: str
    status: str


class MarkerPreservedBlockText(str):
    """The block's markdown, which also remembers what happened to each paragraph.

    A ``str`` subclass rather than a new return type on purpose. ``generate_markdown_block``
    is reached through the ``MarkdownGenerator`` protocol and its result flows through
    assembly, narration, classification and the UI as plain text; changing the type would
    touch every one of those and every test double that returns a string.

    **The subclass is a transport, never a place to compute.** Any ``str`` operation on it
    — ``.strip()``, a slice, an ``f``-string — returns a plain ``str`` and the record is
    gone, and the loss is invisible because ``build_processed_paragraph_registry_entries``
    re-splits on ``\\n\\n`` and finds the same paragraph COUNT. That is how a review found
    English being read aloud with every check green (``_trim_boundary_context_leakage``
    built its result with ``.strip()`` and a slice). So the block's text is derived from the
    dispositions and never the other way round: ``_marker_preserved_block_text`` is the only
    constructor used inside the generator, and every terminal return in marker mode goes
    through ``_deliver_marker_preserved_block``, which RAISES if the record did not survive.
    """

    paragraph_dispositions: tuple[ParagraphDisposition, ...]

    def __new__(
        cls,
        text: str,
        dispositions: Sequence[ParagraphDisposition] = (),
    ) -> "MarkerPreservedBlockText":
        instance = super().__new__(cls, text)
        instance.paragraph_dispositions = tuple(dispositions)
        return instance

    def __reduce__(self) -> tuple[object, ...]:
        # Without this, ``copy.deepcopy`` and ``pickle`` reconstruct through
        # ``cls.__new__(cls, text)`` and lose the record — silently, and only in whichever
        # code path happens to copy a processed block.
        return (self.__class__, (str(self), self.paragraph_dispositions))


class MarkerParagraphRecordLost(RuntimeError):
    """The per-paragraph record was expected on this value and is not there.

    Not a ``MarkerValidationError``: the model did nothing wrong, the code did. It is
    therefore NOT retryable — resending the same request cannot put the record back — and
    it must not be confused with a rejected answer in the retry accounting. Raised so that
    a transport degradation is a stopped block with a named cause instead of source-language
    text delivered quietly under a green classification.
    """

    def __init__(self, stage: str) -> None:
        super().__init__(f"marker_paragraph_record_lost:{stage}")
        self.stage = stage


def marker_paragraph_dispositions(value: object) -> tuple[ParagraphDisposition, ...] | None:
    """The per-paragraph record carried by a processed block, or ``None`` if it has none."""

    dispositions = getattr(value, "paragraph_dispositions", None)
    if not isinstance(dispositions, tuple):
        return None
    if not all(isinstance(item, ParagraphDisposition) for item in dispositions):
        return None
    return dispositions


def _marker_preserved_block_text(
    dispositions: Sequence[ParagraphDisposition],
) -> MarkerPreservedBlockText:
    """The ONLY way the generator builds a marker-preserved block.

    The text is derived from the record, so the two cannot disagree and no code path can
    produce the text without the record.
    """

    return MarkerPreservedBlockText(
        "\n\n".join(disposition.text for disposition in dispositions),
        dispositions,
    )


def _deliver_marker_preserved_block(value: str, *, marker_mode: bool, stage: str) -> str:
    """Every terminal return of a marker-mode ANSWER passes through here.

    Two jobs, both of which have to happen exactly once per block and exactly at the end:

    * the record must still be attached — a plain ``str`` arriving here means some string
      operation upstream dropped it, and that is raised, not tolerated;
    * the per-paragraph counters are recorded HERE rather than where the statuses are
      resolved, because a status resolved on an attempt that is then rejected (context
      leakage, an outer retry) never ships. Counting it made the run report count attempts:
      a two-paragraph block that succeeded once after two rejected attempts reported
      ``accepted: 6``.
    """

    if not marker_mode:
        return value
    dispositions = marker_paragraph_dispositions(value)
    if dispositions is None:
        raise MarkerParagraphRecordLost(stage)
    _record_paragraph_dispositions(dispositions)
    return value


class MarkerValidationError(RuntimeError):
    def __init__(
        self,
        error_code: str,
        *,
        raw_markdown: str,
        expected_paragraph_ids: Sequence[str],
        found_paragraph_ids: Sequence[str] | None = None,
        leading_text: str | None = None,
    ) -> None:
        super().__init__(f"paragraph_marker_validation_failed:{error_code}")
        self.error_code = error_code
        # The FULL answer, kept alongside the preview. The preview feeds prompts and log
        # payloads, which must stay small; the full text feeds the attempt-capture artifact
        # (spec 056 D'), whose only purpose is that a rejected answer can be replayed
        # offline instead of being re-bought. One block's answer, one exception.
        self.raw_markdown = raw_markdown
        self.raw_markdown_preview = raw_markdown[:1000]
        self.expected_paragraph_ids = tuple(expected_paragraph_ids)
        self.found_paragraph_ids = tuple(found_paragraph_ids or ())
        self.leading_text = leading_text or ""
        self.leading_text_preview = (leading_text or "")[:400]


@lru_cache(maxsize=1)
def ensure_pandoc_available() -> None:
    try:
        pypandoc.get_pandoc_version()
    except OSError as exc:
        raise RuntimeError(
            "Pandoc не найден в текущем WSL runtime. Для штатного workflow установите его внутри WSL, "
            "например через: sudo apt-get install -y pandoc"
        ) from exc


def normalize_model_output(text: str) -> str:
    cleaned = text.strip()
    if cleaned.startswith("```"):
        first_newline = cleaned.find("\n")
        if first_newline != -1:
            cleaned = cleaned[first_newline + 1 :].strip()
        else:
            cleaned = cleaned[3:].strip()
    if cleaned.endswith("```"):
        cleaned = cleaned[:-3].strip()
    return cleaned


def strip_markdown_for_narration(text: str) -> str:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    normalized = _NARRATION_INTERNAL_PLACEHOLDER_PATTERN.sub("", normalized)
    normalized = _NARRATION_MARKDOWN_LINK_PATTERN.sub(r"\1", normalized)

    collapsed_lines: list[str] = []
    previous_blank = True
    for raw_line in normalized.split("\n"):
        line_source = raw_line
        is_heading = False

        tagged_heading_match = _NARRATION_TAGGED_HEADING_PATTERN.match(line_source)
        if tagged_heading_match is not None:
            is_heading = True
            tag_prefix = _NARRATION_INTERNAL_WHITESPACE_PATTERN.sub(" ", tagged_heading_match.group(1)).strip()
            heading_text = tagged_heading_match.group(2).strip()
            line_source = f"{tag_prefix} {heading_text}".strip()
        else:
            heading_match = _NARRATION_HEADING_PATTERN.match(line_source)
            if heading_match is not None:
                is_heading = True
                line_source = heading_match.group(1)
            else:
                blockquote_match = _NARRATION_BLOCKQUOTE_PATTERN.match(line_source)
                if blockquote_match is not None:
                    line_source = blockquote_match.group(1)
                else:
                    list_match = _NARRATION_LIST_PATTERN.match(line_source)
                    if list_match is not None:
                        line_source = list_match.group(1)
                    else:
                        tagged_list_match = _NARRATION_TAGGED_LIST_PATTERN.match(line_source)
                        if tagged_list_match is not None:
                            tag_prefix = _NARRATION_INTERNAL_WHITESPACE_PATTERN.sub(
                                " ", tagged_list_match.group(1)
                            ).strip()
                            item_text = tagged_list_match.group(2).strip()
                            line_source = f"{tag_prefix} {item_text}".strip()

        line = line_source.replace("`", "")
        line = _strip_narration_inline_emphasis(line)
        line = _NARRATION_RAW_URL_PATTERN.sub("", line)
        line = _NARRATION_INTERNAL_WHITESPACE_PATTERN.sub(" ", line).strip()
        if not line:
            if not previous_blank:
                collapsed_lines.append("")
            previous_blank = True
            continue
        collapsed_lines.append(line)
        previous_blank = False
        if is_heading:
            collapsed_lines.append("")
            previous_blank = True
    return "\n".join(collapsed_lines).strip()


def _strip_narration_inline_emphasis(text: str) -> str:
    stripped = text
    while True:
        updated = _NARRATION_STRONG_PATTERN.sub(r"\2", stripped)
        updated = _NARRATION_EMPHASIS_PATTERN.sub(r"\2", updated)
        if updated == stripped:
            return updated
        stripped = updated


def _normalize_context_text(text: str | None) -> str:
    if text is None:
        return "[no context]"
    cleaned = text.strip()
    return cleaned or "[no context]"


_CONTEXT_IMAGE_PLACEHOLDER_PATTERN = re.compile(r"\[\[DOCX_IMAGE_img_\d+\]\]")


def _strip_image_placeholders(text: str) -> str:
    """Remove DOCX image placeholder tokens from context strings.

    Image placeholders must not appear in context_before / context_after because
    the model consistently returns empty responses when it encounters them in the
    surrounding context (as opposed to the target block, where they must be
    preserved for later image reinsertion).
    """
    return _CONTEXT_IMAGE_PLACEHOLDER_PATTERN.sub("", text).strip()


def _strip_prompt_internal_tokens(text: str) -> str:
    without_images = _CONTEXT_IMAGE_PLACEHOLDER_PATTERN.sub("", text)
    without_markers = _PARAGRAPH_MARKER_PATTERN.sub("", without_images)
    return without_markers.strip()


def _should_passthrough_target(target_text: str) -> bool:
    stripped_target = target_text.strip()
    if not stripped_target:
        return True
    if _IMAGE_ONLY_TARGET_PATTERN.fullmatch(stripped_target):
        return True
    return not _strip_prompt_internal_tokens(target_text)


def _validate_prompt_inputs(target_text: str, context_before: str, context_after: str) -> list[str]:
    warnings: list[str] = []
    if not target_text.strip():
        warnings.append("empty_target_text")
    elif _IMAGE_ONLY_TARGET_PATTERN.fullmatch(target_text.strip()):
        warnings.append("image_only_target_text")
    elif not _strip_prompt_internal_tokens(target_text):
        warnings.append("placeholder_only_target_text")

    if not context_before.strip():
        warnings.append("empty_context_before")
    if not context_after.strip():
        warnings.append("empty_context_after")
    return warnings


def _build_standard_user_prompt(*, target_text: str, context_before: str, context_after: str) -> str:
    return (
        "Below is a target document block and surrounding context.\n"
        "Use the surrounding context only to understand meaning, terminology, and continuity.\n"
        "Process only the target block according to the system instructions and return only its final text.\n\n"
        f"[CONTEXT BEFORE]\n{context_before}\n\n"
        f"[TARGET BLOCK]\n{target_text}\n\n"
        f"[CONTEXT AFTER]\n{context_after}"
    )


def _build_marker_preserving_user_prompt(*, target_text: str, context_before: str, context_after: str) -> str:
    return (
        "Below is a target document block with required paragraph markers of the form [[DOCX_PARA_...]].\n"
        "Preserve every marker exactly, in the same quantity and order.\n"
        "Do not delete, duplicate, rename, or reorder markers.\n"
        "Do not merge paragraphs across markers and do not split one marker into multiple paragraphs.\n"
        "Every marker must be followed by the processed text of ITS OWN paragraph. If you believe a "
        "paragraph continues the previous one, still return that paragraph's own text under its own "
        "marker, unchanged. Never answer with a placeholder, a stub, a dash, or a note about what you "
        "did (for example \"(Пусто)\", \"(Empty)\", \"(see above)\").\n"
        "Process only the text after each marker according to the system instructions and return the whole block together with the markers.\n"
        "Use the surrounding context only for meaning and terminology.\n\n"
        f"[CONTEXT BEFORE]\n{context_before}\n\n"
        f"[TARGET BLOCK WITH MARKERS]\n{target_text}\n\n"
        f"[CONTEXT AFTER]\n{context_after}"
    )


def _build_empty_response_recovery_user_prompt(*, target_text: str) -> str:
    return (
        "The previous attempt returned an empty answer.\n"
        "Repeat the processing, ignore all external context, and work only with the target block below.\n"
        "Preserve the full meaning, structure, and facts of the block.\n"
        "Return only the final processed block text with no explanation, no Markdown fence, and no empty answer.\n\n"
        f"[TARGET BLOCK ONLY]\n{target_text}"
    )
def _build_marker_recovery_user_prompt(
    *,
    target_text: str,
    expected_paragraph_ids: Sequence[str] | None = None,
    last_error: Exception | None = None,
) -> str:
    required_marker_lines = ""
    if expected_paragraph_ids:
        rendered_markers = "\n".join(f"[[DOCX_PARA_{paragraph_id}]]" for paragraph_id in expected_paragraph_ids)
        required_marker_lines = f"Required marker sequence:\n{rendered_markers}\n\n"

    previous_output_lines = ""
    if isinstance(last_error, MarkerValidationError):
        previous_output_lines = (
            f"Previous marker validation error: {last_error.error_code}.\n"
            f"Expected marker ids: {', '.join(last_error.expected_paragraph_ids) or '[none]'}.\n"
            f"Found marker ids: {', '.join(last_error.found_paragraph_ids) or '[none]'}.\n"
        )
        if last_error.leading_text_preview:
            previous_output_lines += f"Unexpected prefix preview:\n{last_error.leading_text_preview}\n\n"
        if last_error.raw_markdown_preview:
            previous_output_lines += f"Previous invalid output preview:\n{last_error.raw_markdown_preview}\n\n"

    return (
        "The previous attempt violated the paragraph marker contract.\n"
        "Repeat the processing strictly according to the rules below.\n"
        "Preserve every marker [[DOCX_PARA_...]] exactly as it appears and in the same order.\n"
        "Do not delete markers, add new ones, or reorder them.\n"
        "Each marker must correspond to exactly one paragraph of text after it.\n"
        "A marker whose paragraph you would rather merge into its neighbour must still carry that "
        "paragraph's own text, unchanged — never a placeholder, a stub, or a note.\n"
        "The response must begin with the first required marker and contain no explanation.\n\n"
        f"{required_marker_lines}"
        f"{previous_output_lines}"
        f"[TARGET BLOCK WITH MARKERS ONLY]\n{target_text}"
    )


def split_marker_preserved_paragraph_dispositions(
    markdown: str,
    expected_paragraph_ids: Sequence[str],
) -> list[ParagraphDisposition]:
    """Split a marker-preserved answer into ONE typed record per paragraph.

    Spec 056 E. The predecessor returned a list of strings and raised on the first bad
    paragraph, which made a block all-or-nothing: on the 2026-08-04 audiobook run block 274
    held ten paragraphs and every one of them was discarded — replaced by the block's own
    English source in a Russian narration — because paragraph ``p1336``, whose entire text
    is ``14``, came back empty. Nine good translations were thrown away with it.

    What stays block-fatal, unchanged, because these are the checks that detect REAL loss:

    - ``markers_missing`` — no marker at all;
    - ``marker_order_or_identity`` — a marker missing, duplicated or reordered. On the same
      run this caught the model deleting a paragraph TOGETHER with its marker, losing the
      heading ``## NGO Initiative s :`` outright;
    - ``unexpected_prefix`` — text before the first marker, which has no owner.

    What changes: an EMPTY chunk under an exact marker sequence is no longer a block
    failure. It becomes ``retry_required`` — the caller decides whether there is budget
    left to ask again (``resolve_marker_paragraph_dispositions``). Emptiness is what the
    audiobook prompt ORDERS for a paragraph that is pure reference apparatus, while the
    user prompt forbids a stub in its place, so the contract left the model no legal move
    and the block paid for it.

    ``paragraph_split_detected`` — a blank line INSIDE one chunk — is relaxed only where
    the relaxation is provable, and that is the single-marker block; see
    ``_collapse_single_marker_paragraph_break``.
    """

    matches = list(_PARAGRAPH_MARKER_PATTERN.finditer(markdown))
    if not matches:
        raise MarkerValidationError(
            "markers_missing",
            raw_markdown=markdown,
            expected_paragraph_ids=expected_paragraph_ids,
        )

    found_ids = [match.group(1) for match in matches]
    expected_ids = list(expected_paragraph_ids)
    if found_ids != expected_ids:
        raise MarkerValidationError(
            "marker_order_or_identity",
            raw_markdown=markdown,
            expected_paragraph_ids=expected_paragraph_ids,
            found_paragraph_ids=found_ids,
        )

    leading_text = markdown[: matches[0].start()].strip()
    if leading_text:
        raise MarkerValidationError(
            "unexpected_prefix",
            raw_markdown=markdown,
            expected_paragraph_ids=expected_paragraph_ids,
            found_paragraph_ids=found_ids,
            leading_text=leading_text,
        )

    dispositions: list[ParagraphDisposition] = []
    for index, match in enumerate(matches):
        content_end = matches[index + 1].start() if index + 1 < len(matches) else len(markdown)
        chunk = markdown[match.end() : content_end].strip()
        if "\n\n" in chunk:
            if len(matches) != 1 or _has_structural_output_line(chunk):
                raise MarkerValidationError(
                    "paragraph_split_detected",
                    raw_markdown=markdown,
                    expected_paragraph_ids=expected_paragraph_ids,
                    found_paragraph_ids=found_ids,
                )
            chunk = _collapse_single_marker_paragraph_break(chunk, paragraph_id=found_ids[index])
        dispositions.append(
            ParagraphDisposition(
                paragraph_id=found_ids[index],
                text=chunk,
                status=(PARAGRAPH_STATUS_ACCEPTED if chunk else PARAGRAPH_STATUS_RETRY_REQUIRED),
            )
        )
    return dispositions


_STRUCTURAL_OUTPUT_LINE_PATTERN = re.compile(
    r"^\s{0,3}(?:#{1,6}(?:\s|$)|[-*+]\s|\d{1,9}[.)]\s|```|~~~)"
)


def _has_structural_output_line(chunk: str) -> bool:
    """Does the MODEL'S OWN answer for this paragraph contain a structural line?

    A markdown heading, a list item or a code fence. Joining the pieces of a broken chunk
    with a space is only safe for prose: ``## Заголовок`` welded to the sentence after it
    stops being a heading, and ``- пункт`` welded to the text before it stops being a list
    item — the collapse would then repair a split by destroying a structural role, which is
    Constitution VII rule 7 exactly (content survived, its role did not).

    This inspects the shape of the model's OUTPUT, not the shape of the source document, so
    it is not the structure-guessing VII forbids: nothing is inferred about what the
    paragraph "really" is, only whether the answer is written in a form that a space would
    destroy. And the fallback is today's behaviour — refusing to collapse leaves the break
    ``paragraph_split_detected``, exactly as before spec 056 — so the guard is bounded above
    by "no worse than the previous release".

    Measured on the 504 recorded model answers of the 2026-08-04 audiobook run and the
    2026-08-03 literary-edit run: every one of the single-marker breaks the collapse rescues
    is prose, and none of them contains such a line, so the guard costs nothing that was
    being gained.
    """

    return any(
        _STRUCTURAL_OUTPUT_LINE_PATTERN.match(line) is not None
        for line in chunk.splitlines()
    )


def _collapse_single_marker_paragraph_break(chunk: str, *, paragraph_id: str) -> str:
    """Join a paragraph the model broke in two — ONLY when the block holds one marker.

    With a single marker there is no neighbour to steal from and no neighbour to steal for:
    ``unexpected_prefix`` already forbids a fragment before the marker, and every character
    after it belongs to this paragraph because the block contains no other. The collapse is
    provable, not plausible.

    With two or more markers it is NOT, and no rule was found that makes it so. The model
    can place text BEFORE a marker:

    .. code-block:: text

        [[DOCX_PARA_p1]]
        Перевод первого абзаца.

        ## Безопасность
        [[DOCX_PARA_p2]]
        [short pause]

    Marker identity and order are exact, both chunks are non-empty, and ``p2``'s source is
    under ``_COLLAPSED_MARKER_CHUNK_MIN_SOURCE_CHARS`` so the collapse-restore is skipped.
    Collapsing ``p1``'s chunk would hand ``p2``'s heading to ``p1`` with every remaining
    check passing — trading a DETECTED failure for a silent corruption. Attributing the
    trailing fragment would mean comparing it against a source in another language, or
    keying on how it looks; both are the per-document heuristics Constitution VII forbids.
    Restricting the collapse to the last chunk does not help either: the mirror case, where
    an earlier paragraph's text was pushed down past its own marker, has exactly the same
    shape. So for two or more markers the break stays block-fatal, and that is a recorded
    accepted outcome rather than a rule that was not found yet.

    Measured on the 2026-08-04 run: of the three ``paragraph_split_detected`` blocks, 118
    is a single-marker block (one 4 095-character quotation the importer welded together)
    and is rescued here; 164 and 185 hold two and three markers and stay fatal.

    The pieces are joined with a space rather than kept apart, because the contract is one
    marker — one paragraph, and the registry maps one paragraph id to one paragraph.
    """

    parts = [part.strip() for part in re.split(r"\n\s*\n", chunk) if part.strip()]
    collapsed = " ".join(parts)
    log_event(
        logging.INFO,
        "marker_chunk_paragraph_break_collapsed",
        "Модель разбила единственный абзац блока на несколько; части склеены обратно в один абзац.",
        paragraph_id=paragraph_id,
        part_count=len(parts),
        collapsed_chars=len(collapsed),
    )
    return collapsed


def resolve_marker_paragraph_dispositions(
    dispositions: Sequence[ParagraphDisposition],
    *,
    source_paragraph_chunks: Sequence[str] | None = None,
    allow_unresolved_paragraphs: bool,
    raw_markdown: str | None = None,
) -> list[ParagraphDisposition]:
    """Turn ``retry_required`` paragraphs into a final status, or ask for another attempt.

    While there is retry budget left an unresolved paragraph still raises
    ``empty_marker_chunk`` — the same error code, the same retry accounting and the same
    resend as before. That is deliberate: of the seven bare-number paragraphs recorded on
    the 2026-08-04 run, two WERE rescued by a resend, and giving that up to save a call
    would trade real text for money.

    On the last attempt the block is no longer thrown away. Each paragraph the model
    emptied becomes ``omitted``: its own SOURCE text stands in the document, so nothing is
    lost from the DOCX and the paragraph-per-marker mapping the restorer depends on still
    holds, and it is excluded from the narration, where source-language text is exactly
    what spec 054 says must not reach a listener. Every other paragraph keeps its
    translation.

    The collapse-restore runs FIRST, on the raw texts, so an emptied paragraph whose text
    was visibly absorbed by a neighbour is still repaired as a pair (``source_restored``)
    instead of being re-instated on its own and shipped twice.
    """

    texts = [disposition.text for disposition in dispositions]
    statuses = list(disposition.status for disposition in dispositions)

    if source_paragraph_chunks is not None:
        restored_texts = list(
            restore_collapsed_marker_paragraphs(
                list(texts),
                source_paragraph_chunks,
                expected_paragraph_ids=[disposition.paragraph_id for disposition in dispositions],
            )
        )
        if len(restored_texts) == len(texts):
            for index, (before, after) in enumerate(zip(texts, restored_texts)):
                if before != after:
                    statuses[index] = PARAGRAPH_STATUS_SOURCE_RESTORED
            texts = restored_texts

    unresolved_indexes = [
        index
        for index, text in enumerate(texts)
        if not text and statuses[index] != PARAGRAPH_STATUS_SOURCE_RESTORED
    ]
    if unresolved_indexes and not allow_unresolved_paragraphs:
        raise MarkerValidationError(
            "empty_marker_chunk",
            # The model's ACTUAL answer when the caller has it (spec 056 D' captures this
            # exception's ``raw_markdown`` verbatim, and a reconstruction would not replay).
            raw_markdown=(
                raw_markdown
                if raw_markdown is not None
                else "\n\n".join(
                    f"[[DOCX_PARA_{disposition.paragraph_id}]]\n{disposition.text}"
                    for disposition in dispositions
                )
            ),
            expected_paragraph_ids=[disposition.paragraph_id for disposition in dispositions],
            found_paragraph_ids=[disposition.paragraph_id for disposition in dispositions],
        )

    for index in unresolved_indexes:
        statuses[index] = PARAGRAPH_STATUS_OMITTED
        if source_paragraph_chunks is not None and index < len(source_paragraph_chunks):
            texts[index] = source_paragraph_chunks[index]

    resolved = [
        ParagraphDisposition(paragraph_id=disposition.paragraph_id, text=text, status=status)
        for disposition, text, status in zip(dispositions, texts, statuses)
    ]
    if unresolved_indexes:
        # The run report must be able to say WHICH paragraphs the model produced nothing
        # for, and HOW MUCH text stops being spoken because of it. The owner's metric for
        # spec 054 is a share of CHARACTERS, so a count of paragraphs alone cannot be
        # compared against it: a WARNING reading "1 paragraph" hid 3 000 characters of
        # prose on the measurement that prompted this.
        log_event(
            logging.WARNING,
            "marker_paragraph_omitted",
            "Модель не вернула текст для отдельных абзацев блока; в документе остаётся исходный текст, в озвучку они не попадают.",
            omitted_paragraph_ids=[resolved[index].paragraph_id for index in unresolved_indexes],
            omitted_paragraph_count=len(unresolved_indexes),
            omitted_source_chars=sum(len(resolved[index].text) for index in unresolved_indexes),
            paragraph_count=len(resolved),
        )
    return resolved


def _record_paragraph_dispositions(dispositions: Sequence[ParagraphDisposition]) -> None:
    counts: dict[str, int] = {}
    for disposition in dispositions:
        counts[disposition.status] = counts.get(disposition.status, 0) + 1
    for status, count in counts.items():
        record_paragraph_disposition(status=status, count=count)


def _split_marker_preserved_markdown(markdown: str, expected_paragraph_ids: Sequence[str]) -> list[str]:
    """Strict split, used on the SOURCE text where every chunk is a real paragraph.

    The source is built by the pipeline, not by the model, so an empty or broken chunk here
    is a pipeline defect and not a contract the model could not satisfy. It keeps raising.
    """

    dispositions = split_marker_preserved_paragraph_dispositions(markdown, expected_paragraph_ids)
    for disposition in dispositions:
        if not disposition.text:
            raise MarkerValidationError(
                "empty_marker_chunk",
                raw_markdown=markdown,
                expected_paragraph_ids=expected_paragraph_ids,
                found_paragraph_ids=[item.paragraph_id for item in dispositions],
            )
    return [disposition.text for disposition in dispositions]


def _visible_marker_chunk_length(chunk: str) -> int:
    return len(_strip_prompt_internal_tokens(chunk))


def restore_collapsed_marker_paragraphs(
    paragraph_chunks: list[str],
    source_paragraph_chunks: Sequence[str],
    *,
    expected_paragraph_ids: Sequence[str] | None = None,
) -> list[str]:
    """Re-instate the SOURCE text of a paragraph the model MERGED into a neighbour.

    Contract: one marker — one paragraph, and a paragraph's content stays in its own
    paragraph. When the model merges a paragraph into its neighbour it is forbidden to
    delete the emptied marker, so it fills it with an invented stub ("(Пусто)") that ships
    to the reader as literal garbage (7 occurrences on the 2026-08-03 literary-edit run).

    A shrunken chunk on its own is NOT evidence of a merge, and restoring on shrinkage
    alone silently reverts work the model was ASKED to do: a legitimately tightened
    sentence, or — in audiobook mode, where the prompt orders reference and bibliography
    paragraphs to be dropped — exactly the paragraphs that must not reach the narration.
    So a restore requires the merge to be VISIBLE: a neighbour that grew by at least half
    of what this paragraph holds, i.e. an identified place the text went to. No identified
    absorber, no restore.

    The absorber is searched on BOTH sides. The model merges forward as readily as back
    (``[stub for A, "text A text B"]``), and looking only at ``index - 1`` restored A while
    leaving A's text inside the next chunk — shipping A twice.

    A collapse is length-only and therefore operation-agnostic (``edit`` AND ``translate``):
    no translation turns an 861-character paragraph into 7 characters.

    Replayed over the 1761 recorded pairs of that run: all 7 stubs have an absorbing
    neighbour and are still repaired; the single collapse WITHOUT one (a 655-character
    quote) turned out not to be a merge at all — the model had shifted a whole endnote
    region, and that quote was already shipping five markers earlier, so re-instating it
    delivered the quote twice.
    """

    if not source_paragraph_chunks or len(source_paragraph_chunks) != len(paragraph_chunks):
        return paragraph_chunks

    source_lengths = [_visible_marker_chunk_length(chunk) for chunk in source_paragraph_chunks]
    returned_lengths = [_visible_marker_chunk_length(chunk) for chunk in paragraph_chunks]

    collapsed_indexes: list[int] = []
    for index in range(len(paragraph_chunks)):
        source_length = source_lengths[index]
        if returned_lengths[index] == 0:
            # An EMPTY answer is an omission, not a merge, and it has a status of its own
            # (``omitted``). Until spec 056 E an empty chunk raised ``empty_marker_chunk``
            # before this function was ever reached, so this is the input domain the
            # restorer was calibrated on — nothing that legitimately restored before stops
            # restoring. Letting empties in was measurably wrong: the merge evidence is "a
            # neighbour grew by at least half of what this paragraph holds", and a Russian
            # translation of an English paragraph is routinely 1.5x its source, so an
            # ordinary neighbour looked like an absorber. The block then reverted the
            # emptied paragraph AND its perfectly good neighbour to English, marked both
            # ``source_restored`` — a status the narration filter keeps — and read 227
            # characters of English aloud under a ``valid`` classification.
            continue
        if source_length < _COLLAPSED_MARKER_CHUNK_MIN_SOURCE_CHARS:
            continue
        if returned_lengths[index] > source_length * _COLLAPSED_MARKER_CHUNK_RATIO:
            continue
        collapsed_indexes.append(index)

    if not collapsed_indexes:
        return paragraph_chunks

    restored = list(paragraph_chunks)
    restored_indexes: list[int] = []
    merged_indexes: list[int] = []
    kept_indexes: list[int] = []
    for index in collapsed_indexes:
        required_growth = source_lengths[index] * _ABSORBING_NEIGHBOUR_RATIO
        absorbing_indexes = [
            neighbour_index
            for neighbour_index in (index - 1, index + 1)
            if 0 <= neighbour_index < len(paragraph_chunks)
            and returned_lengths[neighbour_index] - source_lengths[neighbour_index] >= required_growth
        ]
        if not absorbing_indexes:
            # No neighbour grew to hold this text: nothing shows a merge happened, so the
            # answer is a short answer and the model's output stands.
            kept_indexes.append(index)
            continue
        merged_indexes.append(index)
        for restore_index in (index, *absorbing_indexes):
            restored[restore_index] = source_paragraph_chunks[restore_index]
            restored_indexes.append(restore_index)

    if kept_indexes:
        log_event(
            logging.INFO,
            "marker_chunk_shrunk_without_absorber_kept",
            "Возвращённый абзац сильно короче исходного, но поглотивший его сосед не найден; ответ модели сохранён.",
            shrunken_paragraph_ids=[
                _marker_paragraph_id(expected_paragraph_ids, index) for index in kept_indexes
            ],
            paragraph_count=len(paragraph_chunks),
        )

    if not merged_indexes:
        return paragraph_chunks

    # The run report must show these paragraphs: the block still logs ``OK``, so without a
    # counter the reader believes a paragraph was edited when in fact its edit was thrown
    # away and the source was shipped instead.
    record_model_output_discarded(
        reason="marker_chunk_collapse",
        paragraph_count=len(set(restored_indexes)),
    )
    log_event(
        logging.WARNING,
        "marker_chunk_collapse_source_restored",
        "Возвращённый абзац схлопнулся в соседний; восстановлен исходный текст обоих абзацев.",
        collapsed_paragraph_ids=[
            _marker_paragraph_id(expected_paragraph_ids, index) for index in merged_indexes
        ],
        restored_paragraph_count=len(set(restored_indexes)),
        paragraph_count=len(paragraph_chunks),
    )
    return restored


def _marker_paragraph_id(expected_paragraph_ids: Sequence[str] | None, index: int) -> str:
    if expected_paragraph_ids and index < len(expected_paragraph_ids):
        return expected_paragraph_ids[index]
    return str(index)


def _strip_paragraph_markers_from_source(
    target_text: str,
    expected_paragraph_ids: Sequence[str] | None,
    *,
    marker_mode: bool,
) -> str:
    """Drop the markers from the block's OWN source text.

    Separate from the model-output path on purpose. The source is not an answer: it has no
    disposition, it must not be counted in the run's per-paragraph statuses (counting it
    would report every source paragraph as ``accepted``), and an empty chunk here is a
    pipeline defect rather than a contract the model could not satisfy — so it keeps
    raising, exactly as before.
    """

    if not marker_mode:
        return target_text
    if not expected_paragraph_ids:
        raise RuntimeError("paragraph_marker_validation_failed:missing_expected_ids")
    return "\n\n".join(_split_marker_preserved_markdown(target_text, expected_paragraph_ids))


def _strip_and_validate_paragraph_markers(
    markdown: str,
    expected_paragraph_ids: Sequence[str] | None,
    *,
    marker_mode: bool,
    source_paragraph_chunks: Sequence[str] | None = None,
    allow_unresolved_paragraphs: bool = False,
) -> str:
    if not marker_mode:
        return markdown
    if not expected_paragraph_ids:
        raise RuntimeError("paragraph_marker_validation_failed:missing_expected_ids")
    dispositions = resolve_marker_paragraph_dispositions(
        split_marker_preserved_paragraph_dispositions(markdown, expected_paragraph_ids),
        source_paragraph_chunks=source_paragraph_chunks,
        allow_unresolved_paragraphs=allow_unresolved_paragraphs,
        raw_markdown=markdown,
    )
    # The joined string still carries exactly one paragraph per marker — an ``omitted``
    # paragraph contributes its own source text — so everything downstream that counts
    # paragraphs keeps counting the same number. What rides along is WHY each one is
    # there, which is what the registry needs and could never recover from the join.
    return _marker_preserved_block_text(dispositions)


def _normalize_leakage_comparison_text(text: str) -> str:
    return " ".join(match.group(0).lower() for match in _WORD_TOKEN_PATTERN.finditer(text))


def _detect_context_leakage(
    response_text: str,
    target_text: str,
    context_before: str,
    context_after: str,
    *,
    min_word_sequence: int = 6,
) -> str | None:
    response_tokens = list(_WORD_TOKEN_PATTERN.finditer(response_text))
    if len(response_tokens) < min_word_sequence:
        return None

    normalized_target = _normalize_leakage_comparison_text(target_text)
    normalized_contexts = [
        normalized_context
        for normalized_context in (
            _normalize_leakage_comparison_text(context_before),
            _normalize_leakage_comparison_text(context_after),
        )
        if normalized_context
    ]
    if not normalized_contexts:
        return None

    for start_index in range(0, len(response_tokens) - min_word_sequence + 1):
        end_index = start_index + min_word_sequence
        fragment = response_text[
            response_tokens[start_index].start() : response_tokens[end_index - 1].end()
        ]
        normalized_fragment = _normalize_leakage_comparison_text(fragment)
        if not normalized_fragment or normalized_fragment in normalized_target:
            continue
        if any(normalized_fragment in normalized_context for normalized_context in normalized_contexts):
            return fragment
    return None


_LEAKAGE_TRIM_STRIP_CHARS = " \t\r\n-–—,:;.!?"


def _trim_boundary_context_leakage(response_text: str, leaked_fragment: str) -> tuple[str, bool]:
    trimmed_response = response_text.strip()
    if not trimmed_response or leaked_fragment not in trimmed_response:
        return response_text, False

    matches = list(re.finditer(re.escape(leaked_fragment), trimmed_response))
    if not matches:
        return response_text, False
    if any(match.start() != 0 and match.end() != len(trimmed_response) for match in matches):
        return response_text, False

    updated_text = trimmed_response
    changed = False
    while updated_text.startswith(leaked_fragment):
        updated_text = updated_text[len(leaked_fragment) :].lstrip(_LEAKAGE_TRIM_STRIP_CHARS)
        changed = True
    while updated_text.endswith(leaked_fragment):
        updated_text = updated_text[: -len(leaked_fragment)].rstrip(_LEAKAGE_TRIM_STRIP_CHARS)
        changed = True

    if not changed or not updated_text:
        return response_text, False
    return updated_text, True


def _trim_marker_preserved_boundary_leakage(
    value: str,
    leaked_fragment: str,
) -> tuple[str, bool]:
    """The boundary trim, applied to the RECORD instead of to the joined string.

    ``_trim_boundary_context_leakage`` builds its answer with ``.strip()`` and slices, which
    return a plain ``str``. Running it on a ``MarkerPreservedBlockText`` therefore threw the
    per-paragraph record away inside the generator, before the value ever reached a caller —
    the one degradation ``__reduce__`` cannot protect against — and the paragraph COUNT still
    matched afterwards, so nothing downstream noticed and an ``omitted`` paragraph's English
    source went into the narration.

    The leak is at a boundary of the whole answer by construction (the string-level function
    refuses to trim anything else), so it belongs to the first paragraph that has text or to
    the last one. Both are trimmed with the same rule, and the result is re-derived from the
    record. If the two do not agree character for character the trim could not be attributed
    to a paragraph — the answer is then rejected as a marker failure and retried, rather than
    shipped with a record that no longer describes it.
    """

    dispositions = marker_paragraph_dispositions(value)
    if dispositions is None:
        return _trim_boundary_context_leakage(value, leaked_fragment)

    trimmed_text, was_trimmed = _trim_boundary_context_leakage(str(value), leaked_fragment)
    if not was_trimmed:
        return value, False

    texts = [disposition.text for disposition in dispositions]
    populated_indexes = [index for index, text in enumerate(texts) if text]
    if populated_indexes:
        first_index = populated_indexes[0]
        while texts[first_index].startswith(leaked_fragment):
            texts[first_index] = texts[first_index][len(leaked_fragment) :].lstrip(
                _LEAKAGE_TRIM_STRIP_CHARS
            )
        last_index = populated_indexes[-1]
        while texts[last_index].endswith(leaked_fragment):
            texts[last_index] = texts[last_index][: -len(leaked_fragment)].rstrip(
                _LEAKAGE_TRIM_STRIP_CHARS
            )

    rebuilt = _marker_preserved_block_text(
        [
            ParagraphDisposition(
                paragraph_id=disposition.paragraph_id,
                text=text,
                status=disposition.status,
            )
            for disposition, text in zip(dispositions, texts)
        ]
    )
    if str(rebuilt) != trimmed_text:
        raise MarkerValidationError(
            "context_leakage_trim_unattributable",
            raw_markdown=str(value),
            expected_paragraph_ids=[disposition.paragraph_id for disposition in dispositions],
            found_paragraph_ids=[disposition.paragraph_id for disposition in dispositions],
        )
    return rebuilt, True


def _inject_context_leakage_retry_warning(request_kwargs: dict[str, object]) -> dict[str, object]:
    updated_request = dict(request_kwargs)
    payload = updated_request.get("input")
    if not isinstance(payload, list):
        return updated_request

    updated_payload: list[object] = []
    for index, message in enumerate(payload):
        if index != 1 or not isinstance(message, dict):
            updated_payload.append(message)
            continue

        updated_message = dict(message)
        content_items = list(updated_message.get("content", []))
        if content_items and isinstance(content_items[0], dict):
            updated_content = dict(content_items[0])
            text = updated_content.get("text")
            if isinstance(text, str) and _CONTEXT_LEAKAGE_RETRY_WARNING not in text:
                updated_content["text"] = f"{_CONTEXT_LEAKAGE_RETRY_WARNING}\n\n{text}"
                content_items[0] = updated_content
        updated_message["content"] = content_items
        updated_payload.append(updated_message)

    updated_request["input"] = updated_payload
    return updated_request


def _finalize_generated_markdown(
    markdown: str,
    *,
    target_text: str,
    context_before: str,
    context_after: str,
    expected_paragraph_ids: Sequence[str] | None,
    marker_mode: bool,
    allow_persistent_context_leakage: bool,
    source_paragraph_chunks: Sequence[str] | None = None,
    allow_unresolved_paragraphs: bool = False,
) -> str:
    cleaned_markdown = _strip_and_validate_paragraph_markers(
        markdown,
        expected_paragraph_ids,
        marker_mode=marker_mode,
        source_paragraph_chunks=source_paragraph_chunks,
        allow_unresolved_paragraphs=allow_unresolved_paragraphs,
    )
    leaked_fragment = _detect_context_leakage(
        cleaned_markdown,
        target_text,
        context_before,
        context_after,
    )
    if leaked_fragment is None:
        return _deliver_marker_preserved_block(
            cleaned_markdown, marker_mode=marker_mode, stage="finalize_clean"
        )

    trimmed_markdown, was_trimmed = _trim_marker_preserved_boundary_leakage(
        cleaned_markdown, leaked_fragment
    )
    if was_trimmed:
        return _deliver_marker_preserved_block(
            trimmed_markdown, marker_mode=marker_mode, stage="finalize_leakage_trimmed"
        )

    if allow_persistent_context_leakage:
        log_event(
            logging.WARNING,
            "context_leakage_persisted",
            "После последней попытки генерации сохранилась verbatim-протечка текста из соседнего контекста; возвращаю fail-open результат.",
            leaked_fragment=leaked_fragment,
            target_chars=len(target_text),
            marker_mode=marker_mode,
        )
        return _deliver_marker_preserved_block(
            cleaned_markdown, marker_mode=marker_mode, stage="finalize_leakage_persisted"
        )

    raise ContextLeakageError(f"context_leakage_detected:{leaked_fragment}")


def _extract_response_output_text(response: object) -> str:
    traversal = collect_response_text_traversal(
        response,
        unsupported_message="Модель вернула ответ в неподдерживаемом формате (unsupported_response_shape).",
    )
    if traversal.collected_texts:
        return "\n".join(traversal.collected_texts)
    if traversal.raw_output_text is not None:
        return traversal.raw_output_text
    if not traversal.saw_output_items or traversal.saw_supported_text_shape or traversal.saw_empty_content_container:
        return ""
    raise RuntimeError("Модель вернула ответ в неподдерживаемом формате (unsupported_response_shape).")


def _log_empty_response_shape(response: object, raw_output_text: str, *, error_code: str) -> None:
    output_items = read_response_field(response, "output")
    output_items_len = len(output_items) if isinstance(output_items, Sized) else None

    first_item_summary: dict[str, object] | None = None
    if isinstance(output_items, Iterable) and not isinstance(output_items, (str, bytes)):
        for item in output_items:
            item_type = read_response_field(item, "type")
            refusal = read_response_field(item, "refusal")
            status = read_response_field(item, "status")
            content_items = read_response_field(item, "content")
            content_types: list[str] = []
            if isinstance(content_items, Iterable) and not isinstance(content_items, (str, bytes)):
                content_types = [
                    str(read_response_field(c, "type") or type(c).__name__)
                    for c in content_items
                ]
            first_item_summary = {
                "type": item_type,
                "refusal": refusal,
                "status": status,
                "content_types": content_types,
            }
            break

    log_event(
        logging.WARNING,
        "model_empty_response_shape",
        "Модель вернула пустой или схлопнувшийся текстовый ответ",
        error_code=error_code,
        has_output_text_attr=getattr(response, "output_text", None) is not None,
        raw_output_len=len(raw_output_text),
        output_items_type=type(output_items).__name__ if output_items is not None else "None",
        output_items_len=output_items_len,
        response_status=read_response_field(response, "status"),
        first_output_item=first_item_summary,
    )


def _extract_normalized_markdown(response: object) -> str:
    response_status = read_response_field(response, "status")
    if response_status == "incomplete":
        _log_empty_response_shape(response, "", error_code="incomplete_response")
        raise RuntimeError("Модель не завершила генерацию (incomplete_response).")
    if isinstance(response_status, str) and response_status != "completed":
        _log_empty_response_shape(response, "", error_code="non_completed_response")
        raise RuntimeError(f"Модель вернула неожиданный статус ответа: {response_status} (non_completed_response).")

    raw_output_text = _extract_response_output_text(response)
    markdown = normalize_model_output(raw_output_text)
    if markdown:
        return markdown
    error_code = "collapsed_output" if raw_output_text else "empty_response"
    _log_empty_response_shape(response, raw_output_text, error_code=error_code)
    if raw_output_text:
        raise RuntimeError("Модель вернула ответ, который схлопнулся после нормализации (collapsed_output).")
    raise RuntimeError("Модель вернула пустой ответ (empty_response).")


def _call_responses_create(client: "OpenAI", request_kwargs: dict[str, Any]) -> object:
    return call_responses_create_with_retry(
        client,
        request_kwargs,
        max_retries=1,
        retryable_error_predicate=lambda exc: False,
        retryable_optional_params={"temperature", "max_output_tokens"},
        usage_stage=STAGE_TEXT_GENERATION,
    )


def _is_openrouter_client(client: object) -> bool:
    base_url = getattr(client, "base_url", None)
    if base_url is None:
        return False
    return "openrouter" in str(base_url).lower()


def _is_anthropic_client(client: object) -> bool:
    return callable(getattr(getattr(client, "messages", None), "create", None)) and not callable(
        getattr(getattr(client, "responses", None), "create", None)
    )


def _is_openrouter_responses_compatibility_error(exc: Exception) -> bool:
    status_code = getattr(exc, "status_code", None)
    if isinstance(status_code, int) and status_code not in {400, 404, 422}:
        return False
    error_text = str(exc).lower()
    return any(
        marker in error_text
        for marker in (
            "/responses",
            "responses api unsupported",
            "responses endpoint",
            "chat.completions",
            "unsupported parameter",
            "unsupported field",
            "unknown parameter",
            "unknown field",
        )
    )


def _canonicalize_model_selector_for_client(*, client: object, model: str) -> str:
    stripped_model = str(model).strip()
    if not stripped_model:
        return stripped_model
    if ":" in stripped_model:
        provider_name, _, model_id = stripped_model.partition(":")
        normalized_provider = provider_name.strip().lower()
        if normalized_provider in {"openai", "openrouter", "anthropic"} and model_id.strip():
            return f"{normalized_provider}:{model_id.strip()}"
    provider_name = "anthropic" if _is_anthropic_client(client) else "openrouter" if _is_openrouter_client(client) else "openai"
    return f"{provider_name}:{stripped_model}"


def _normalize_anthropic_model_id(model: str) -> str:
    model_id = str(model).strip()
    if model_id.startswith("anthropic:"):
        model_id = model_id.split(":", 1)[1].strip()
    aliases = {
        "claude-sonnet-4.6": "claude-sonnet-4-6",
        "anthropic/claude-sonnet-4.6": "claude-sonnet-4-6",
    }
    return aliases.get(model_id, model_id)


def _extract_chat_messages_from_request(request_kwargs: dict[str, object]) -> list[dict[str, str]]:
    raw_input = request_kwargs.get("input")
    if not isinstance(raw_input, list):
        raise RuntimeError("OpenRouter Chat Completions fallback не может собрать messages из request input.")

    messages: list[dict[str, str]] = []
    for item in raw_input:
        if not isinstance(item, dict):
            continue
        role = str(item.get("role") or "user").strip() or "user"
        raw_content = item.get("content")
        if isinstance(raw_content, list):
            text_parts = [
                str(content_item.get("text") or "")
                for content_item in raw_content
                if isinstance(content_item, dict) and content_item.get("type") == "input_text"
            ]
            content_text = "\n".join(part for part in text_parts if part)
        else:
            content_text = str(raw_content or "")
        messages.append({"role": role, "content": content_text})

    if not messages:
        raise RuntimeError("OpenRouter Chat Completions fallback не может собрать ни одного message.")
    return messages


def _call_chat_completions_create(client: "OpenAI", request_kwargs: dict[str, object]) -> object:
    chat_completions = getattr(getattr(client, "chat", None), "completions", None)
    create = getattr(chat_completions, "create", None)
    if not callable(create):
        raise RuntimeError("Provider 'openrouter' не поддерживает required text API surface для selector '<runtime>'.")

    payload: dict[str, object] = {
        "model": request_kwargs["model"],
        "messages": _extract_chat_messages_from_request(request_kwargs),
        "temperature": request_kwargs.get("temperature"),
    }
    max_output_tokens = request_kwargs.get("max_output_tokens")
    if isinstance(max_output_tokens, int) and not isinstance(max_output_tokens, bool):
        payload["max_tokens"] = max_output_tokens

    removable_optional_params = {"temperature", "max_tokens"}
    while True:
        try:
            response = create(**payload)
        except TypeError as exc:
            unsupported_param = extract_unsupported_parameter_name(str(exc))
            if unsupported_param in removable_optional_params and unsupported_param in payload:
                payload.pop(unsupported_param, None)
                continue
            raise
        except Exception as exc:
            unsupported_param = extract_unsupported_parameter_name(str(exc))
            if unsupported_param in removable_optional_params and unsupported_param in payload:
                payload.pop(unsupported_param, None)
                continue
            raise
        else:
            record_model_call_usage(stage=STAGE_TEXT_GENERATION, response=response)
            return response


def _extract_chat_completion_markdown(response: object) -> str:
    choices = getattr(response, "choices", None)
    if not isinstance(choices, list) or not choices:
        raise RuntimeError("Provider 'openrouter' не поддерживает required text API surface для selector '<runtime>'.")
    first_choice = choices[0]
    message = getattr(first_choice, "message", None)
    content = getattr(message, "content", None)
    if isinstance(content, str):
        markdown = normalize_model_output(content)
        if markdown:
            return markdown
        raise RuntimeError("Модель вернула ответ, который схлопнулся после нормализации (collapsed_output).")
    if isinstance(content, list):
        text_parts = [
            str(getattr(item, "text", "") or item.get("text") or "")
            for item in content
            if isinstance(item, dict) or hasattr(item, "text")
        ]
        markdown = normalize_model_output("\n".join(part for part in text_parts if part))
        if markdown:
            return markdown
    raise RuntimeError("Модель вернула пустой ответ (empty_response).")


def _extract_anthropic_messages_payload(request_kwargs: dict[str, object]) -> tuple[str, list[dict[str, str]]]:
    raw_input = request_kwargs.get("input")
    if not isinstance(raw_input, list):
        raise RuntimeError("Provider 'anthropic' не может собрать messages из request input.")

    system_parts: list[str] = []
    messages: list[dict[str, str]] = []
    for item in raw_input:
        if not isinstance(item, dict):
            continue
        role = str(item.get("role") or "user").strip() or "user"
        raw_content = item.get("content")
        if isinstance(raw_content, list):
            text = "\n".join(
                str(content_item.get("text") or "")
                for content_item in raw_content
                if isinstance(content_item, dict) and content_item.get("type") == "input_text"
            )
        else:
            text = str(raw_content or "")
        if not text.strip():
            continue
        if role == "system":
            system_parts.append(text)
        else:
            messages.append({"role": "assistant" if role == "assistant" else "user", "content": text})

    if not messages:
        raise RuntimeError("Provider 'anthropic' не может собрать ни одного message.")
    return "\n\n".join(system_parts), messages


def _call_anthropic_messages_create(client: object, request_kwargs: dict[str, object]) -> object:
    create = getattr(getattr(client, "messages", None), "create", None)
    if not callable(create):
        raise RuntimeError("Provider 'anthropic' не поддерживает required text API surface для selector '<runtime>'.")

    system_prompt, messages = _extract_anthropic_messages_payload(request_kwargs)
    payload: dict[str, object] = {
        "model": _normalize_anthropic_model_id(str(request_kwargs["model"])),
        "messages": messages,
        "max_tokens": request_kwargs.get("max_output_tokens"),
    }
    if system_prompt:
        payload["system"] = system_prompt
    temperature = request_kwargs.get("temperature")
    if temperature is not None:
        payload["temperature"] = temperature

    removable_optional_params = {"temperature"}
    while True:
        try:
            response = create(**payload)
        except TypeError as exc:
            unsupported_param = extract_unsupported_parameter_name(str(exc))
            if unsupported_param in removable_optional_params and unsupported_param in payload:
                payload.pop(unsupported_param, None)
                continue
            raise
        except Exception as exc:
            unsupported_param = extract_unsupported_parameter_name(str(exc))
            if unsupported_param in removable_optional_params and unsupported_param in payload:
                payload.pop(unsupported_param, None)
                continue
            raise
        else:
            # Anthropic reports ``usage.input_tokens``/``output_tokens`` but no cost, so
            # this path deliberately leaves cost unknown instead of applying a price list.
            record_model_call_usage(stage=STAGE_TEXT_GENERATION, response=response)
            return response


def _extract_anthropic_message_markdown(response: object) -> str:
    content = getattr(response, "content", None)
    if not isinstance(content, list) or not content:
        raise RuntimeError("Модель вернула пустой ответ (empty_response).")
    text_parts = [
        str(getattr(item, "text", "") or item.get("text") or "")
        for item in content
        if isinstance(item, dict) or hasattr(item, "text")
    ]
    markdown = normalize_model_output("\n".join(part for part in text_parts if part))
    if markdown:
        return markdown
    raise RuntimeError("Модель вернула пустой ответ (empty_response).")


def _estimate_max_output_tokens(target_text: str) -> int:
    estimated_output_tokens = max((len(target_text) // 3) * 4, 512)
    return min(estimated_output_tokens, 16384)


def _build_request_kwargs(*, model: str, system_prompt: str, user_prompt: str, target_text: str) -> dict[str, object]:
    payload: Any = [
        {
            "role": "system",
            "content": [{"type": "input_text", "text": system_prompt}],
        },
        {
            "role": "user",
            "content": [{"type": "input_text", "text": user_prompt}],
        },
    ]
    return {
        "model": model,
        "input": payload,
        "temperature": 0.4,
        "max_output_tokens": _estimate_max_output_tokens(target_text),
    }


def _boost_request_output_budget(
    request_kwargs: dict[str, object],
    *,
    minimum_tokens: int,
) -> dict[str, object]:
    boosted_request = dict(request_kwargs)
    current_value = boosted_request.get("max_output_tokens")
    if isinstance(current_value, int) and not isinstance(current_value, bool):
        boosted_request["max_output_tokens"] = min(max(current_value * 2, minimum_tokens), 16384)
        return boosted_request
    boosted_request["max_output_tokens"] = min(max(minimum_tokens, 512), 16384)
    return boosted_request


def _call_markdown_request_with_sdk_fallback(client: "OpenAI", request_kwargs: dict[str, object]) -> tuple[str, bool]:
    if _is_anthropic_client(client):
        response = _call_anthropic_messages_create(client, request_kwargs)
        return _extract_anthropic_message_markdown(response), False
    try:
        response = _call_responses_create(client, cast(dict[str, Any], request_kwargs))
        return _extract_normalized_markdown(response), False
    except Exception as exc:
        if not _is_openrouter_client(client) or not _is_openrouter_responses_compatibility_error(exc):
            raise
        log_event(
            logging.WARNING,
            "provider_text_api_fallback_engaged",
            "Responses API для text path отклонён provider-совместимостью; переключаюсь на Chat Completions fallback.",
            provider="openrouter",
            model=str(request_kwargs.get("model") or ""),
            model_selector=str(request_kwargs.get("model") or ""),
            canonical_model_selector=_canonicalize_model_selector_for_client(
                client=client,
                model=str(request_kwargs.get("model") or ""),
            ),
            api_surface="chat.completions",
            fallback_reason=str(exc),
        )
        chat_response = _call_chat_completions_create(client, request_kwargs)
        return _extract_chat_completion_markdown(chat_response), True


def _recover_from_persistent_empty_response(
    *,
    client: "OpenAI",
    model: str,
    system_prompt: str,
    target_text: str,
    expected_paragraph_ids: Sequence[str] | None = None,
    marker_mode: bool = False,
    minimum_output_tokens: int | None = None,
    last_exception: Exception | None = None,
) -> str:
    log_event(
        logging.WARNING,
        "markdown_empty_response_recovery_started",
        "Обычные retry исчерпаны; запускаю recovery-вызов без соседнего контекста.",
        model=model,
        target_chars=len(target_text),
    )
    request_kwargs = _build_request_kwargs(
        model=model,
        system_prompt=system_prompt,
        user_prompt=(
            _build_marker_recovery_user_prompt(
                target_text=target_text,
                expected_paragraph_ids=expected_paragraph_ids,
                last_error=last_exception,
            )
            if marker_mode
            else _build_empty_response_recovery_user_prompt(target_text=target_text)
        ),
        target_text=target_text,
    )
    if minimum_output_tokens is not None:
        request_kwargs = _boost_request_output_budget(
            request_kwargs,
            minimum_tokens=minimum_output_tokens,
        )
    markdown = _call_markdown_request_with_sdk_fallback(client, request_kwargs)[0]
    cleaned_markdown = _strip_and_validate_paragraph_markers(
        markdown,
        expected_paragraph_ids,
        marker_mode=marker_mode,
        source_paragraph_chunks=(
            _split_marker_preserved_markdown(target_text, expected_paragraph_ids)
            if marker_mode and expected_paragraph_ids
            else None
        ),
        # The recovery call is the last chance there is; there is nothing left to retry for.
        allow_unresolved_paragraphs=True,
    )
    if not cleaned_markdown.strip():
        raise RuntimeError("Модель вернула пустой ответ (empty_response).")
    return _deliver_marker_preserved_block(
        cleaned_markdown, marker_mode=marker_mode, stage="recovery"
    )


def _block_source_fallback_result(
    target_text_for_leakage: str,
    *,
    expected_paragraph_ids: Sequence[str] | None,
    source_paragraph_chunks: Sequence[str] | None,
    marker_mode: bool,
) -> str:
    """The block-level fallback — every paragraph reverted to its own source — WITH a record.

    The four block-level fallbacks used to return a bare string, so a marker-mode block could
    leave the generator with no per-paragraph record for a perfectly ordinary reason. That
    made "no record" ambiguous downstream, and an ambiguous invariant cannot be enforced.
    Now every marker-mode exit carries one, which is what lets ``generate_markdown_block``
    refuse to hand back a block whose record went missing.

    The TEXT is unchanged: ``_strip_paragraph_markers_from_source`` builds
    ``target_text_for_leakage`` by joining exactly these chunks with a blank line.
    ``source_restored`` is also the truthful status — the model's answer was discarded and
    the source re-instated — and the counter for it was already being recorded by hand on
    one of the four paths.
    """

    if not marker_mode or not expected_paragraph_ids or source_paragraph_chunks is None:
        return target_text_for_leakage
    if len(source_paragraph_chunks) != len(expected_paragraph_ids):
        return target_text_for_leakage
    return _deliver_marker_preserved_block(
        _marker_preserved_block_text(
            [
                ParagraphDisposition(
                    paragraph_id=paragraph_id,
                    text=chunk,
                    status=PARAGRAPH_STATUS_SOURCE_RESTORED,
                )
                for paragraph_id, chunk in zip(expected_paragraph_ids, source_paragraph_chunks)
            ]
        ),
        marker_mode=marker_mode,
        stage="block_source_fallback",
    )


def _is_incomplete_response_error(exc: Exception) -> bool:
    return isinstance(exc, RuntimeError) and "incomplete_response" in str(exc)


def _can_fallback_to_source_text_after_incomplete_response(target_text: str) -> bool:
    return bool(target_text.strip())


def _can_fallback_to_source_text_after_marker_validation_failure(target_text: str, *, marker_mode: bool) -> bool:
    return marker_mode and bool(target_text.strip())


def _is_empty_response_error(exc: Exception) -> bool:
    return isinstance(exc, RuntimeError) and "empty_response" in str(exc)


def _can_fallback_to_source_text_after_empty_response(target_text: str) -> bool:
    return bool(target_text.strip())


def _is_non_completed_response_error(exc: Exception) -> bool:
    return isinstance(exc, RuntimeError) and "non_completed_response" in str(exc)


def _can_fallback_to_source_text_after_non_completed_response(target_text: str) -> bool:
    return bool(target_text.strip())


def _is_retryable_empty_generation_error(exc: Exception) -> bool:
    return isinstance(exc, RuntimeError) and (
        "empty_response" in str(exc) or "collapsed_output" in str(exc) or "incomplete_response" in str(exc)
    )


def _is_retryable_marker_validation_error(exc: Exception) -> bool:
    return isinstance(exc, RuntimeError) and "paragraph_marker_validation_failed" in str(exc)


def _is_retryable_context_leakage_error(exc: Exception) -> bool:
    return isinstance(exc, ContextLeakageError)


def _classify_retry_reason(exc: Exception) -> str:
    """Name WHY a block was retried, so the run report is diagnosable, not just a count."""

    if _is_retryable_context_leakage_error(exc):
        return "context_leakage"
    if _is_retryable_marker_validation_error(exc):
        return "marker_validation"
    if _is_incomplete_response_error(exc):
        return "incomplete_response"
    if _is_non_completed_response_error(exc):
        return "non_completed_response"
    if _is_retryable_empty_generation_error(exc):
        return "empty_generation"
    if is_retryable_error(exc):
        return "transient_api_error"
    return "other"


def _capture_marker_attempt_failure(
    exc: Exception,
    *,
    block_index: int | None,
    attempt: int,
    max_attempts: int,
    stage: str,
    target_chars: int,
) -> str | None:
    """Persist a rejected marker answer, if this failure carries one. Never raises.

    Only ``MarkerValidationError`` is captured: it is the only failure that HOLDS the
    model's answer together with the expected/found marker ids. A transient API error has
    no answer to keep, and an empty response has nothing to replay.
    """

    if not isinstance(exc, MarkerValidationError):
        return None
    try:
        return capture_rejected_marker_attempt(
            block_index=block_index,
            attempt=attempt,
            max_attempts=max_attempts,
            stage=stage,
            error_code=exc.error_code,
            expected_paragraph_ids=exc.expected_paragraph_ids,
            found_paragraph_ids=exc.found_paragraph_ids,
            raw_response=exc.raw_markdown,
            leading_text=exc.leading_text,
            target_chars=target_chars,
        )
    except Exception:
        return None


def generate_markdown_block(
    client: "OpenAI",
    model: str,
    system_prompt: str,
    target_text: str,
    context_before: str,
    context_after: str,
    max_retries: int,
    expected_paragraph_ids: Sequence[str] | None = None,
    marker_mode: bool = False,
    block_index: int | None = None,
) -> str:
    if isinstance(max_retries, bool) or not isinstance(max_retries, int):
        raise TypeError("max_retries должен быть целым числом.")
    if max_retries < 1:
        raise ValueError("max_retries должен быть не меньше 1.")

    if _should_passthrough_target(target_text):
        log_event(
            logging.WARNING,
            "image_only_target_passthrough",
            "Целевой блок не содержит редактируемого текста; возвращаю его без вызова модели.",
            target_chars=len(target_text),
            marker_mode=marker_mode,
        )
        return target_text

    context_before_text = _normalize_context_text(_strip_image_placeholders(context_before))
    context_after_text = _normalize_context_text(_strip_image_placeholders(context_after))
    prompt_warnings = _validate_prompt_inputs(target_text, context_before_text, context_after_text)
    if prompt_warnings:
        log_event(
            logging.WARNING,
            "prompt_quality_warning",
            "Входные данные prompt содержат потенциально проблемный shape.",
            warnings=prompt_warnings,
            target_chars=len(target_text),
            context_before_chars=len(context_before_text),
            context_after_chars=len(context_after_text),
            marker_mode=marker_mode,
        )

    request_kwargs = _build_request_kwargs(
        model=model,
        system_prompt=system_prompt,
        user_prompt=(
            _build_marker_preserving_user_prompt(
                target_text=target_text,
                context_before=context_before_text,
                context_after=context_after_text,
            )
            if marker_mode
            else _build_standard_user_prompt(
                target_text=target_text,
                context_before=context_before_text,
                context_after=context_after_text,
            )
        ),
        target_text=target_text,
    )
    target_text_for_leakage = _strip_paragraph_markers_from_source(
        target_text,
        expected_paragraph_ids,
        marker_mode=marker_mode,
    )
    source_paragraph_chunks = (
        _split_marker_preserved_markdown(target_text, expected_paragraph_ids)
        if marker_mode and expected_paragraph_ids
        else None
    )
    last_exception: Exception | None = None
    # A block that needed a second attempt is invisible today: every retry is silent and
    # the block still reports ``OK``. Count the attempts, the blocks and — in marker mode,
    # where the mapping is exact — the paragraphs those blocks carried.
    block_paragraph_count = len(expected_paragraph_ids) if marker_mode and expected_paragraph_ids else 0
    block_was_retried = False

    for attempt in range(1, max_retries + 1):
        try:
            markdown = _call_markdown_request_with_sdk_fallback(client, request_kwargs)[0]
            return _finalize_generated_markdown(
                markdown,
                target_text=target_text_for_leakage,
                context_before=context_before_text,
                context_after=context_after_text,
                expected_paragraph_ids=expected_paragraph_ids,
                marker_mode=marker_mode,
                allow_persistent_context_leakage=attempt >= max_retries,
                source_paragraph_chunks=source_paragraph_chunks,
                # Same shape as the leakage fail-open one line above: while budget remains
                # an unanswered paragraph is worth asking about again (2 of 7 bare-number
                # paragraphs were rescued that way on the 2026-08-04 run), and on the last
                # attempt the answer stands with a per-paragraph status instead of the
                # whole block being replaced by its own source.
                allow_unresolved_paragraphs=attempt >= max_retries,
            )
        except Exception as exc:
            last_exception = exc
            # Spec 056 D': write the rejected answer down HERE, before anything decides
            # whether to retry or to fall back. Every other place is too late — a controlled
            # fallback returns a plain string and the call site never sees this exception.
            _capture_marker_attempt_failure(
                exc,
                block_index=block_index,
                attempt=attempt,
                max_attempts=max_retries,
                stage="attempt",
                target_chars=len(target_text),
            )
            should_retry = attempt < max_retries and (
                is_retryable_error(exc)
                or _is_retryable_empty_generation_error(exc)
                or _is_retryable_marker_validation_error(exc)
                or _is_retryable_context_leakage_error(exc)
                or _is_non_completed_response_error(exc)
            )
            if not should_retry:
                break
            record_retry_attempt(
                reason=_classify_retry_reason(exc),
                paragraph_count=block_paragraph_count,
                first_retry_for_block=not block_was_retried,
            )
            block_was_retried = True
            if _is_incomplete_response_error(exc):
                request_kwargs = _boost_request_output_budget(
                    request_kwargs,
                    minimum_tokens=_INCOMPLETE_RESPONSE_RETRY_MIN_OUTPUT_TOKENS,
                )
            if _is_retryable_context_leakage_error(exc):
                request_kwargs = _inject_context_leakage_retry_warning(request_kwargs)
            time.sleep(min(2 ** (attempt - 1), 8))

    if last_exception is not None and (
        _is_retryable_empty_generation_error(last_exception)
        or _is_retryable_marker_validation_error(last_exception)
    ):
        record_retry_attempt(
            reason="recovery_after_exhausted_retries",
            paragraph_count=block_paragraph_count,
            first_retry_for_block=not block_was_retried,
        )
        block_was_retried = True
        try:
            return _recover_from_persistent_empty_response(
                client=client,
                model=model,
                system_prompt=system_prompt,
                target_text=target_text,
                expected_paragraph_ids=expected_paragraph_ids,
                marker_mode=marker_mode,
                minimum_output_tokens=(
                    _INCOMPLETE_RESPONSE_RECOVERY_MIN_OUTPUT_TOKENS
                    if _is_incomplete_response_error(last_exception)
                    else None
                ),
                last_exception=last_exception,
            )
        except Exception as recovery_exc:
            # The recovery call is the LAST answer the model gives for this block, and the
            # one whose rejection sends the block's own English source into the artifact.
            # It is the single most valuable attempt to have on disk.
            _capture_marker_attempt_failure(
                recovery_exc,
                block_index=block_index,
                attempt=max_retries + 1,
                max_attempts=max_retries + 1,
                stage="recovery",
                target_chars=len(target_text),
            )
            if _is_incomplete_response_error(recovery_exc) and _can_fallback_to_source_text_after_incomplete_response(target_text):
                record_model_output_discarded(reason="incomplete_response_source_fallback", block_count=1)
                log_event(
                    logging.WARNING,
                    "markdown_incomplete_response_source_fallback",
                    "Recovery для блока снова завершился incomplete_response; сохраняю исходный текст блока как controlled fallback.",
                    model=model,
                    target_chars=len(target_text_for_leakage),
                    marker_mode=marker_mode,
                )
                return _block_source_fallback_result(
                    target_text_for_leakage,
                    expected_paragraph_ids=expected_paragraph_ids,
                    source_paragraph_chunks=source_paragraph_chunks,
                    marker_mode=marker_mode,
                )
            if _is_retryable_marker_validation_error(recovery_exc) and _can_fallback_to_source_text_after_marker_validation_failure(
                target_text_for_leakage,
                marker_mode=marker_mode,
            ):
                record_model_output_discarded(reason="marker_validation_source_fallback", block_count=1)
                log_event(
                    logging.WARNING,
                    "markdown_marker_validation_source_fallback",
                    "Recovery для блока снова завершился marker validation error; сохраняю исходный текст блока как controlled fallback.",
                    model=model,
                    target_chars=len(target_text_for_leakage),
                    marker_mode=marker_mode,
                    marker_error=str(recovery_exc),
                )
                return _block_source_fallback_result(
                    target_text_for_leakage,
                    expected_paragraph_ids=expected_paragraph_ids,
                    source_paragraph_chunks=source_paragraph_chunks,
                    marker_mode=marker_mode,
                )
            if _is_empty_response_error(recovery_exc) and _can_fallback_to_source_text_after_empty_response(
                target_text
            ):
                record_model_output_discarded(reason="empty_response_source_fallback", block_count=1)
                log_event(
                    logging.WARNING,
                    "markdown_empty_response_source_fallback",
                    "Recovery для блока снова завершился empty_response; сохраняю исходный текст блока как controlled fallback.",
                    model=model,
                    target_chars=len(target_text_for_leakage),
                    marker_mode=marker_mode,
                    recovery_error=str(recovery_exc),
                )
                return _block_source_fallback_result(
                    target_text_for_leakage,
                    expected_paragraph_ids=expected_paragraph_ids,
                    source_paragraph_chunks=source_paragraph_chunks,
                    marker_mode=marker_mode,
                )
            if _is_retryable_empty_generation_error(recovery_exc) or _is_retryable_marker_validation_error(recovery_exc):
                raise recovery_exc
            raise recovery_exc

    if (
        last_exception is not None
        and _is_non_completed_response_error(last_exception)
        and _can_fallback_to_source_text_after_non_completed_response(target_text)
    ):
        record_model_output_discarded(reason="non_completed_response_source_fallback", block_count=1)
        log_event(
            logging.WARNING,
            "markdown_non_completed_response_source_fallback",
            "Модель повторно вернула non_completed_response; сохраняю исходный текст блока как controlled fallback.",
            model=model,
            target_chars=len(target_text_for_leakage),
            marker_mode=marker_mode,
        )
        return _block_source_fallback_result(
            target_text_for_leakage,
            expected_paragraph_ids=expected_paragraph_ids,
            source_paragraph_chunks=source_paragraph_chunks,
            marker_mode=marker_mode,
        )

    if last_exception is not None:
        raise last_exception

    raise RuntimeError("Не удалось получить ответ модели.")


_THEME_RELATIONSHIP_TYPE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme"
)
_DRAWINGML_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def convert_markdown_to_docx_bytes(
    markdown_text: str,
    *,
    body_font: str | None = None,
    heading_font: str | None = None,
) -> bytes:
    """Convert *markdown_text* to DOCX bytes via Pandoc.

    *body_font* and *heading_font* are optional overrides for the reference
    document. Body-facing styles are updated directly because python-docx writes
    them as explicit ``w:rFonts`` values; heading styles additionally require a
    theme patch because Word gives ``w:asciiTheme=majorHAnsi`` precedence over
    the direct font name. When both are ``None`` (the default) the python-docx
    built-in theme is left unchanged.
    """
    ensure_pandoc_available()
    try:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            markdown_path = temp_path / "result.md"
            docx_path = temp_path / "result.docx"
            reference_docx_path = temp_path / "reference.docx"
            markdown_path.write_text(_preprocess_markdown_for_docx(markdown_text), encoding="utf-8")
            _build_reference_docx(reference_docx_path, body_font=body_font, heading_font=heading_font)
            pypandoc.convert_file(
                str(markdown_path),
                to="docx",
                format="markdown+raw_html+superscript+subscript",
                outputfile=str(docx_path),
                extra_args=[f"--reference-doc={reference_docx_path}"],
            )
            return docx_path.read_bytes()
    except Exception as exc:
        raise RuntimeError(f"Ошибка при сборке DOCX: {exc}") from exc


def _preprocess_markdown_for_docx(markdown_text: str) -> str:
    """Convert extractor-emitted inline HTML into Pandoc-friendly markdown.

    The extraction layer preserves semantic inline signals as HTML-like tags
    such as ``<sup>``, ``<sub>``, ``<u>``, and ``<br/>``. Pandoc's DOCX writer does
    not preserve those shapes reliably through the default markdown path, so we
    translate them into markdown extensions that round-trip into OOXML.
    """
    # Underline runs FIRST. Its escaping must not see the backslashes that
    # _escape_pandoc_script_spaces injects for ^…^/~…~: doubling those would turn
    # the escaped space into a literal backslash and demote the script role.
    # Nested tags survive because ``<sup>``/``<sub>``/``<br/>`` carry no span
    # metacharacters, so the later passes still rewrite them inside the span.
    processed = _INLINE_HTML_UNDERLINE_PATTERN.sub(_render_pandoc_underline_span, markdown_text)
    processed = _INLINE_HTML_SUP_PATTERN.sub(
        lambda match: f"^{_escape_pandoc_script_spaces(match.group(1))}^", processed
    )
    processed = _INLINE_HTML_SUB_PATTERN.sub(
        lambda match: f"~{_escape_pandoc_script_spaces(match.group(1))}~", processed
    )
    processed = _INLINE_HTML_BREAK_PATTERN.sub("\\\n", processed)
    return processed


def _render_pandoc_underline_span(match: "re.Match[str]") -> str:
    """Wrap underlined content into Pandoc's ``[…]{.underline}`` span, escaped.

    The span delimiters are markdown syntax, so unescaped ``[``/``]``/``\\`` in
    the *content* terminate or escape them early and the construct collapses:
    an underlined run reading ``1]`` (DOCX splits runs anywhere, so a trailing
    footnote bracket easily lands in its own underlined run) produced
    ``[1]]{.underline}`` and the reader saw the literal markup in the delivered
    text. Escaping keeps the characters as body text and the underline as a role.
    """
    prefix = _neutralize_pandoc_span_prefix(match.group("prefix"))
    content = _PANDOC_SPAN_SPECIAL_PATTERN.sub(lambda special: f"\\{special.group(0)}", match.group("content"))
    return f"{prefix}[{content}]{{.underline}}"


def _neutralize_pandoc_span_prefix(prefix: str) -> str:
    """Stop the character in front of the span from consuming its opening ``[``.

    ``!`` immediately before ``[`` starts Pandoc's image syntax and an odd run of
    backslashes escapes the bracket outright; either way the span degrades to
    literal ``[…]{.underline}`` in the document. Both are neutralized without
    changing the text the reader sees.
    """
    if prefix.endswith("!"):
        backslashes = prefix[:-1]
        if len(backslashes) % 2 == 1:
            # Already escaped by the preceding backslash: cannot open an image.
            return prefix
        return f"{backslashes}\\!"
    if len(prefix) % 2 == 1:
        return f"{prefix}\\"
    return prefix


def _escape_pandoc_script_spaces(content: str) -> str:
    """Escape spaces so ``^…^`` / ``~…~`` stay real superscript/subscript.

    Pandoc's superscript and subscript may not contain unescaped spaces, so
    ``<sup>note 1</sup>`` would otherwise translate into a literal ``^note 1^``
    that reaches the reader as raw carets. Escaping keeps the vertical-alignment
    role; leaving such content untranslated keeps the text but silently demotes
    it to ordinary body text. The escaped space arrives in OOXML as a
    non-breaking space, which is the accepted cost of preserving the role.
    """
    return content.replace(" ", "\\ ")


def _patch_reference_theme_fonts(
    reference_document: "DocxDocument",
    *,
    body_font: str | None,
    heading_font: str | None,
) -> None:
    """Overwrite the major/minor font slots in the reference document's theme.

    Word resolves ``w:asciiTheme="majorHAnsi"`` (used by all built-in Heading
    styles) and ``w:asciiTheme="minorHAnsi"`` (body/list/caption) by looking
    up the document's embedded ``theme1.xml``.  python-docx's default template
    maps those slots to Calibri (major) and Cambria (minor).

    Patching the theme here means every ``w:asciiTheme`` reference in every
    style automatically picks up the configured font **without** touching
    individual style ``w:rFonts`` elements — this is the OOXML-idiomatic
    approach and the only reliable way to override heading fonts given that
    python-docx's ``Style.font.name`` setter leaves ``w:asciiTheme`` intact.

    Called only when at least one font is configured; both arguments may be
    ``None`` to skip their respective slot.
    """
    try:
        theme_part = reference_document.part.part_related_by(_THEME_RELATIONSHIP_TYPE)
    except KeyError:
        return  # Template has no theme part — nothing to patch.

    root = etree.fromstring(theme_part.blob)

    if heading_font is not None:
        for el in root.findall(f".//{{{_DRAWINGML_NS}}}majorFont/{{{_DRAWINGML_NS}}}latin"):
            el.set("typeface", heading_font)

    if body_font is not None:
        for el in root.findall(f".//{{{_DRAWINGML_NS}}}minorFont/{{{_DRAWINGML_NS}}}latin"):
            el.set("typeface", body_font)

    theme_part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _build_reference_docx(
    reference_docx_path: Path,
    *,
    body_font: str | None = None,
    heading_font: str | None = None,
) -> None:
    reference_document = Document()
    styles = reference_document.styles

    body_baseline = {
        "font_name": body_font,
        "font_size": 11,
        "space_after": 8,
        "line_spacing": 1.15,
    }

    _configure_paragraph_style(styles["Normal"], **body_baseline)

    if "Body Text" in styles:
        _configure_paragraph_style(styles["Body Text"], **body_baseline)

    heading_specs = (
        ("Heading 1", 18, 18, 8),
        ("Heading 2", 16, 16, 7),
        ("Heading 3", 14, 14, 6),
        ("Heading 4", 13, 12, 5),
        ("Heading 5", 12, 10, 4),
        ("Heading 6", 11, 8, 3),
    )
    for style_name, font_size, space_before, space_after in heading_specs:
        if style_name not in styles:
            continue
        _configure_paragraph_style(
            styles[style_name],
            font_name=heading_font,
            font_size=font_size,
            bold=True,
            space_before=space_before,
            space_after=space_after,
            line_spacing=1.1,
            keep_with_next=True,
        )

    if "Caption" in styles:
        _configure_paragraph_style(
            styles["Caption"],
            font_name=body_font,
            font_size=10,
            italic=True,
            space_before=4,
            space_after=10,
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
        )

    if "List Paragraph" in styles:
        _configure_paragraph_style(
            styles["List Paragraph"],
            font_name=body_font,
            font_size=11,
            space_before=0,
            space_after=4,
            line_spacing=1.1,
        )

    if "Table Grid" in styles:
        table_grid_style = cast(Any, styles["Table Grid"])
        if body_font is not None:
            table_grid_style.font.name = body_font
        table_grid_style.font.size = Pt(10)

    _ensure_reference_numbering_definitions(reference_document, body_font=body_font)

    # Patch theme fonts only when explicitly configured. This is required for
    # heading styles, whose built-in w:asciiTheme bindings outrank the direct
    # font name, and keeps theme-bound fallback slots aligned with explicit
    # style overrides.
    if body_font is not None or heading_font is not None:
        _patch_reference_theme_fonts(reference_document, body_font=body_font, heading_font=heading_font)

    reference_document.save(str(reference_docx_path))


def _ensure_reference_numbering_definitions(
    reference_document: "DocxDocument",
    *,
    body_font: str | None = None,
) -> None:
    numbering_part = reference_document.part.numbering_part
    numbering = numbering_part.element
    baseline_specs = (
        {
            "num_fmt": "decimal",
            "level_text_patterns": ("%1.", "%1.%2.", "%1.%2.%3."),
        },
        {
            "num_fmt": "bullet",
            "level_text_patterns": (chr(0x2022), chr(0x25E6), chr(0x25AA)),
        },
    )

    for spec in baseline_specs:
        baseline_abstract_num = _find_reference_baseline_abstract_num(
            numbering,
            num_fmt=spec["num_fmt"],
            level_text_patterns=spec["level_text_patterns"],
            body_font=body_font,
        )
        if baseline_abstract_num is None:
            abstract_num_id = _next_numbering_id(numbering, "w:abstractNum", "abstractNumId")
            _append_multilevel_numbering_definition(
                numbering,
                abstract_num_id=abstract_num_id,
                num_fmt=spec["num_fmt"],
                level_text_patterns=spec["level_text_patterns"],
                body_font=body_font,
            )
            baseline_abstract_num = _find_abstract_num_by_id(numbering, abstract_num_id)
            if baseline_abstract_num is None:
                raise RuntimeError("Не удалось создать baseline numbering definition.")

        abstract_num_id = int(baseline_abstract_num.get(qn("w:abstractNumId")))
        if not _num_instance_exists(numbering, abstract_num_id=abstract_num_id):
            num_id = _next_numbering_id(numbering, "w:num", "numId")
            _append_num_instance(numbering, num_id=num_id, abstract_num_id=abstract_num_id)


def _find_reference_baseline_abstract_num(
    numbering,
    *,
    num_fmt: str,
    level_text_patterns: tuple[str, ...],
    body_font: str | None,
):
    for abstract_num in numbering.xpath('./*[local-name()="abstractNum"]'):
        if _abstract_num_matches_reference_baseline(
            abstract_num,
            num_fmt=num_fmt,
            level_text_patterns=level_text_patterns,
            body_font=body_font,
        ):
            return abstract_num
    return None


def _find_abstract_num_by_id(numbering, abstract_num_id: int):
    matches = numbering.xpath(
        f'./*[local-name()="abstractNum" and @*[local-name()="abstractNumId"]="{abstract_num_id}"]'
    )
    return matches[0] if matches else None


def _iter_num_instances(numbering):
    return numbering.xpath('./*[local-name()="num"]')


def _num_instance_abstract_num_id(num_instance) -> str | None:
    abstract_num_id_values = num_instance.xpath(
        './*[local-name()="abstractNumId"]/@*[local-name()="val"]'
    )
    if not abstract_num_id_values:
        return None
    return str(abstract_num_id_values[0])


def _abstract_num_matches_reference_baseline(
    abstract_num,
    *,
    num_fmt: str,
    level_text_patterns: tuple[str, ...],
    body_font: str | None,
) -> bool:
    levels = abstract_num.xpath('./*[local-name()="lvl"]')
    if len(levels) != len(level_text_patterns):
        return False

    for ilvl, level_text in enumerate(level_text_patterns):
        level_matches = [level for level in levels if level.get(qn("w:ilvl")) == str(ilvl)]
        if len(level_matches) != 1:
            return False
        if not _level_matches_reference_baseline(
            level_matches[0],
            num_fmt=num_fmt,
            level_text=level_text,
            ilvl=ilvl,
            body_font=body_font,
        ):
            return False
    return True


def _level_matches_reference_baseline(
    level,
    *,
    num_fmt: str,
    level_text: str,
    ilvl: int,
    body_font: str | None,
) -> bool:
    num_fmt_values = level.xpath('./*[local-name()="numFmt"]/@*[local-name()="val"]')
    level_text_values = level.xpath('./*[local-name()="lvlText"]/@*[local-name()="val"]')
    left_values = level.xpath('./*[local-name()="pPr"]/*[local-name()="ind"]/@*[local-name()="left"]')
    hanging_values = level.xpath('./*[local-name()="pPr"]/*[local-name()="ind"]/@*[local-name()="hanging"]')
    after_values = level.xpath('./*[local-name()="pPr"]/*[local-name()="spacing"]/@*[local-name()="after"]')
    line_values = level.xpath('./*[local-name()="pPr"]/*[local-name()="spacing"]/@*[local-name()="line"]')
    line_rule_values = level.xpath('./*[local-name()="pPr"]/*[local-name()="spacing"]/@*[local-name()="lineRule"]')
    ascii_fonts = level.xpath('./*[local-name()="rPr"]/*[local-name()="rFonts"]/@*[local-name()="ascii"]')
    hansi_fonts = level.xpath('./*[local-name()="rPr"]/*[local-name()="rFonts"]/@*[local-name()="hAnsi"]')
    cs_fonts = level.xpath('./*[local-name()="rPr"]/*[local-name()="rFonts"]/@*[local-name()="cs"]')

    expected_fonts = [] if body_font is None else [body_font]

    return (
        num_fmt_values == [num_fmt]
        and level_text_values == [level_text]
        and left_values == [str(720 + (ilvl * 360))]
        and hanging_values == ["360"]
        and after_values == ["80"]
        and line_values == ["264"]
        and line_rule_values == ["auto"]
        and ascii_fonts == expected_fonts
        and hansi_fonts == expected_fonts
        and cs_fonts == expected_fonts
    )


def _num_instance_exists(numbering, *, abstract_num_id: int) -> bool:
    expected_abstract_num_id = str(abstract_num_id)
    return any(
        _num_instance_abstract_num_id(num_instance) == expected_abstract_num_id
        for num_instance in _iter_num_instances(numbering)
    )


def _next_numbering_id(numbering, element_name: str, attr_name: str) -> int:
    existing_ids = []
    local_name = element_name.split(":", 1)[1]
    for element in numbering.xpath(f'./*[local-name()="{local_name}"]'):
        value = element.get(qn(f"w:{attr_name}"))
        if value is not None:
            existing_ids.append(int(value))
    return (max(existing_ids) + 1) if existing_ids else 0


def _append_multilevel_numbering_definition(
    numbering,
    *,
    abstract_num_id: int,
    num_fmt: str,
    level_text_patterns: tuple[str, str, str],
    body_font: str | None,
) -> None:
    abstract_num = OxmlElement("w:abstractNum")
    abstract_num.set(qn("w:abstractNumId"), str(abstract_num_id))

    nsid = OxmlElement("w:nsid")
    nsid.set(qn("w:val"), f"{abstract_num_id + 1:08X}")
    abstract_num.append(nsid)

    multi_level_type = OxmlElement("w:multiLevelType")
    multi_level_type.set(qn("w:val"), "multilevel")
    abstract_num.append(multi_level_type)

    template_code = OxmlElement("w:tmpl")
    template_code.set(qn("w:val"), f"{abstract_num_id + 257:08X}")
    abstract_num.append(template_code)

    for ilvl, level_text in enumerate(level_text_patterns):
        abstract_num.append(
            _build_numbering_level(
                ilvl=ilvl,
                num_fmt=num_fmt,
                level_text=level_text,
                body_font=body_font,
            )
        )

    numbering.append(abstract_num)


def _build_numbering_level(*, ilvl: int, num_fmt: str, level_text: str, body_font: str | None):
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), str(ilvl))

    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)

    num_fmt_element = OxmlElement("w:numFmt")
    num_fmt_element.set(qn("w:val"), num_fmt)
    level.append(num_fmt_element)

    level_text_element = OxmlElement("w:lvlText")
    level_text_element.set(qn("w:val"), level_text)
    level.append(level_text_element)

    level_jc = OxmlElement("w:lvlJc")
    level_jc.set(qn("w:val"), "left")
    level.append(level_jc)

    paragraph_properties = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(720 + (ilvl * 360)))
    ind.set(qn("w:hanging"), "360")
    paragraph_properties.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "264")
    spacing.set(qn("w:lineRule"), "auto")
    paragraph_properties.append(spacing)
    level.append(paragraph_properties)

    if body_font is not None:
        run_properties = OxmlElement("w:rPr")
        run_fonts = OxmlElement("w:rFonts")
        run_fonts.set(qn("w:ascii"), body_font)
        run_fonts.set(qn("w:hAnsi"), body_font)
        run_fonts.set(qn("w:cs"), body_font)
        run_properties.append(run_fonts)
        level.append(run_properties)

    return level


def _append_num_instance(numbering, *, num_id: int, abstract_num_id: int) -> None:
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id_element = OxmlElement("w:abstractNumId")
    abstract_num_id_element.set(qn("w:val"), str(abstract_num_id))
    num.append(abstract_num_id_element)
    numbering.append(num)


def _configure_paragraph_style(
    style,
    *,
    font_name: str | None,
    font_size: int,
    bold: bool | None = None,
    italic: bool | None = None,
    space_before: int | None = None,
    space_after: int | None = None,
    line_spacing: float | None = None,
    keep_with_next: bool | None = None,
    alignment=None,
) -> None:
    if font_name is not None:
        style.font.name = font_name
    style.font.size = Pt(font_size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic

    paragraph_format = style.paragraph_format
    if space_before is not None:
        paragraph_format.space_before = Pt(space_before)
    if space_after is not None:
        paragraph_format.space_after = Pt(space_after)
    if line_spacing is not None:
        paragraph_format.line_spacing = line_spacing
    if keep_with_next is not None:
        paragraph_format.keep_with_next = keep_with_next
    if alignment is not None:
        paragraph_format.alignment = alignment


def build_output_filename(filename: str) -> str:
    return f"{Path(filename).stem}_edited.docx"


def build_markdown_filename(filename: str) -> str:
    return f"{Path(filename).stem}_edited.md"
