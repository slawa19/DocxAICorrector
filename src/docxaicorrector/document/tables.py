import re

from docx.table import Table

from docxaicorrector.core.models import ImageAsset, RawTable

_BR_PATTERN = re.compile(r"<br\s*/?>", re.IGNORECASE)
_SUP_PATTERN = re.compile(r"<sup>(.*?)</sup>", re.IGNORECASE | re.DOTALL)
_SUB_PATTERN = re.compile(r"<sub>(.*?)</sub>", re.IGNORECASE | re.DOTALL)


def build_raw_table(
    table: Table,
    image_assets: list[ImageAsset],
    *,
    raw_index: int,
    asset_id: str,
    build_paragraph_text_with_placeholders,
) -> RawTable | None:
    """Emit a KEPT table as a Pandoc-markdown table so the render produces a real ``w:tbl``.

    The extractor historically emitted raw ``<table>`` HTML, which Pandoc's
    ``markdown+raw_html`` reader turns into a dropped ``RawBlock`` (only the cell
    text survives, as loose paragraphs). Emitting a pipe/grid markdown table
    instead makes Pandoc's markdown reader build an actual Word table. Only the
    tabular structure, cell text and inline emphasis are preserved — source
    column widths / colors / border styling / fonts are intentionally dropped
    (minimal-formatting canon).
    """
    markdown_table = render_table_markdown(
        table,
        image_assets,
        build_paragraph_text_with_placeholders=build_paragraph_text_with_placeholders,
    )
    if not markdown_table.strip():
        return None
    return RawTable(raw_index=raw_index, html_text=markdown_table, asset_id=asset_id)


def flatten_table_lines(
    table: Table,
    image_assets: list[ImageAsset],
    *,
    build_paragraph_text_with_placeholders,
) -> list[str]:
    """Flatten a table into a linear stream of body-line strings.

    Cells are emitted in reading order (row-major); ``<br/>`` boundaries and
    multi-paragraph cells split into separate lines; empty cells are skipped.
    Inline emphasis is preserved. Used for scan-origin (OCR) documents where the
    "table" is really a scanned column layout, not authored tabular data.
    """
    lines: list[str] = []
    for row in table.rows:
        for cell in _iter_row_cells(row):
            lines.extend(
                _cell_segments(
                    cell,
                    image_assets,
                    build_paragraph_text_with_placeholders=build_paragraph_text_with_placeholders,
                )
            )
    return lines


def render_table_markdown(
    table: Table,
    image_assets: list[ImageAsset],
    *,
    build_paragraph_text_with_placeholders,
) -> str:
    rows: list[list[list[str]]] = []
    for row in table.rows:
        rendered_row = [
            _cell_segments(
                cell,
                image_assets,
                build_paragraph_text_with_placeholders=build_paragraph_text_with_placeholders,
            )
            for cell in _iter_row_cells(row)
        ]
        rows.append(rendered_row)

    column_count = max((len(row) for row in rows), default=0)
    if column_count == 0:
        return ""
    for row in rows:
        while len(row) < column_count:
            row.append([])

    if not any(any(cell for cell in row) for row in rows):
        return ""

    needs_grid = any(len(cell) > 1 for row in rows for cell in row)
    if needs_grid:
        return _render_grid_table(rows, column_count)
    return _render_pipe_table(rows, column_count)


def _iter_row_cells(row):
    """Yield each distinct cell once (python-docx repeats merged cells)."""
    seen: set[int] = set()
    for cell in row.cells:
        identity = id(cell._tc)
        if identity in seen:
            continue
        seen.add(identity)
        yield cell


def _cell_segments(
    cell,
    image_assets: list[ImageAsset],
    *,
    build_paragraph_text_with_placeholders,
) -> list[str]:
    segments: list[str] = []
    for paragraph in cell.paragraphs:
        text = build_paragraph_text_with_placeholders(paragraph, image_assets).strip()
        if not text:
            continue
        for part in _BR_PATTERN.split(text):
            part = part.strip()
            if part:
                segments.append(part)
    return segments


def _inline_cell(text: str) -> str:
    """Convert inline HTML markers to Pandoc markdown and neutralise table syntax."""
    text = _SUP_PATTERN.sub(lambda match: f"^{match.group(1)}^", text)
    text = _SUB_PATTERN.sub(lambda match: f"~{match.group(1)}~", text)
    return text.replace("|", "\\|")


def _has_header_row(rows: list[list[list[str]]]) -> bool:
    return len(rows) > 1 and all(cell for cell in rows[0])


def _render_pipe_table(rows: list[list[list[str]]], column_count: int) -> str:
    def render_row(cells: list[list[str]]) -> str:
        rendered = [_inline_cell(cell[0]) if cell else "" for cell in cells]
        return "| " + " | ".join(rendered) + " |"

    if _has_header_row(rows):
        header_row = rows[0]
        body_rows = rows[1:]
    else:
        header_row = [[] for _ in range(column_count)]
        body_rows = rows

    lines = [render_row(header_row)]
    lines.append("|" + "|".join(" --- " for _ in range(column_count)) + "|")
    for row in body_rows:
        lines.append(render_row(row))
    return "\n".join(lines)


def _render_grid_table(rows: list[list[list[str]]], column_count: int) -> str:
    grid = [[_grid_cell_lines(cell) for cell in row] for row in rows]

    widths = [3] * column_count
    for grid_row in grid:
        for column_index, cell_lines in enumerate(grid_row):
            for line in cell_lines:
                widths[column_index] = max(widths[column_index], len(line))

    def separator(fill: str) -> str:
        return "+" + "+".join(fill * (width + 2) for width in widths) + "+"

    def render_block(grid_row: list[list[str]]) -> list[str]:
        height = max((len(cell_lines) for cell_lines in grid_row), default=1)
        block: list[str] = []
        for line_index in range(height):
            parts: list[str] = []
            for column_index in range(column_count):
                cell_lines = grid_row[column_index]
                content = cell_lines[line_index] if line_index < len(cell_lines) else ""
                parts.append(" " + content.ljust(widths[column_index]) + " ")
            block.append("|" + "|".join(parts) + "|")
        return block

    output = [separator("-")]
    if _has_header_row(rows):
        output.extend(render_block(grid[0]))
        output.append(separator("="))
        body = grid[1:]
    else:
        body = grid

    for grid_row in body:
        output.extend(render_block(grid_row))
        output.append(separator("-"))
    return "\n".join(output)


def _grid_cell_lines(cell_segments: list[str]) -> list[str]:
    """Lay out a cell's segments as physical grid lines.

    Multiple segments become separate paragraphs inside the cell (blank line
    between them) so line boundaries survive without relying on fragile
    trailing-backslash hard breaks that padding would defeat.
    """
    segments = [_inline_cell(segment) for segment in cell_segments]
    if not segments:
        return [""]
    lines: list[str] = []
    for index, segment in enumerate(segments):
        if index > 0:
            lines.append("")
        lines.append(segment)
    return lines
