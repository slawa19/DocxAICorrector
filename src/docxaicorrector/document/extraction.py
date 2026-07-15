import re
import zipfile
from collections import Counter
from dataclasses import replace
from io import BytesIO
from math import isclose
from pathlib import Path
from collections.abc import Mapping
from typing import cast

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from docxaicorrector.core.constants import (
    MAX_DOCX_ARCHIVE_SIZE_BYTES,
    MAX_DOCX_COMPRESSION_RATIO,
    MAX_DOCX_ENTRY_COUNT,
    MAX_DOCX_UNCOMPRESSED_SIZE_BYTES,
)
from docxaicorrector.document.boundaries import (
    evaluate_paragraph_boundary as _evaluate_paragraph_boundary_impl,
    normalize_paragraph_boundaries as _normalize_paragraph_boundaries_impl,
    resolve_paragraph_boundary_normalization_settings as _resolve_paragraph_boundary_normalization_settings_impl,
    summarize_boundary_normalization_metrics,
    write_paragraph_boundary_report_artifact as _write_paragraph_boundary_report_artifact_impl,
)
from docxaicorrector.document.boundary_review import (
    build_ai_review_candidates as _build_ai_review_candidates_impl,
    coerce_int_config_value as _coerce_int_config_value_impl,
    request_ai_review_recommendations as _request_ai_review_recommendations_impl,
    resolve_paragraph_boundary_ai_review_settings as _resolve_paragraph_boundary_ai_review_settings_impl,
    run_paragraph_boundary_ai_review as _run_paragraph_boundary_ai_review_impl,
    write_paragraph_boundary_ai_review_artifact as _write_paragraph_boundary_ai_review_artifact_impl,
)
from docxaicorrector.document.relations import (
    apply_relation_side_effects,
    build_paragraph_relations,
    write_relation_normalization_report_artifact as _write_relation_normalization_report_artifact_impl,
)
from docxaicorrector.document.layout_cleanup import (
    clean_paragraph_layout_artifacts,
    write_layout_cleanup_report_artifact as _write_layout_cleanup_report_artifact_impl,
)
from docxaicorrector.document.structure_repair import repair_pdf_derived_structure
from docxaicorrector.document.roles import (
    detect_explicit_list_kind,
    extract_explicit_heading_level,
    find_child_element,
    get_xml_attribute,
    has_heading_text_signal,
    infer_heuristic_heading_level,
    infer_role_confidence,
    is_caption_style,
    is_likely_caption_text,
    is_probable_heading,
    normalize_front_matter_display_title,
    paragraph_is_effectively_bold,
    paragraph_is_effectively_italic,
    promote_short_standalone_headings,
    reclassify_adjacent_captions,
    resolve_effective_paragraph_font_size,
    resolve_paragraph_alignment,
    resolve_paragraph_outline_level,
    xml_local_name,
)
from docxaicorrector.document.shared_xml import (
    build_source_xml_fingerprint,
    extract_num_pr_level,
    extract_run_element_images,
    resolve_num_pr_details,
    resolve_paragraph_num_pr,
)
from docxaicorrector.document.tables import (
    build_raw_table as _build_raw_table_impl,
    flatten_table_lines as _flatten_table_lines_impl,
)
from docxaicorrector.document.provenance import classify_document_scan_origin
from docxaicorrector.core.models import (
    PARAGRAPH_BOUNDARY_AI_REVIEW_MODE_VALUES,
    PARAGRAPH_BOUNDARY_NORMALIZATION_MODE_VALUES,
    ImageAsset,
    ParagraphBoundaryDecision,
    ParagraphBoundaryNormalizationReport,
    ParagraphRelation,
    ParagraphUnit,
    RawBlock,
    RawParagraph,
    RawTable,
    RelationNormalizationReport,
    LayoutArtifactCleanupReport,
    StructureRepairReport,
    normalize_heuristic_structural_role_hint,
)
from docxaicorrector.processing.processing_runtime import read_uploaded_file_bytes, resolve_uploaded_filename
from docxaicorrector.runtime.artifact_retention import (
    PARAGRAPH_BOUNDARY_AI_REVIEW_MAX_AGE_SECONDS,
    PARAGRAPH_BOUNDARY_AI_REVIEW_MAX_COUNT,
    PARAGRAPH_BOUNDARY_REPORTS_MAX_AGE_SECONDS,
    PARAGRAPH_BOUNDARY_REPORTS_MAX_COUNT,
    LAYOUT_CLEANUP_REPORTS_MAX_AGE_SECONDS,
    LAYOUT_CLEANUP_REPORTS_MAX_COUNT,
    RELATION_NORMALIZATION_REPORTS_MAX_AGE_SECONDS,
    RELATION_NORMALIZATION_REPORTS_MAX_COUNT,
)


IMAGE_PLACEHOLDER_PATTERN = re.compile(r"\[\[DOCX_IMAGE_img_\d+\]\]")
RELATIONSHIP_NAMESPACE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
ORDERED_LIST_FORMATS = {
    "aiueo",
    "cardinalText",
    "chicago",
    "decimal",
    "decimalEnclosedCircle",
    "decimalEnclosedFullstop",
    "decimalEnclosedParen",
    "decimalFullWidth",
    "decimalFullWidth2",
    "decimalHalfWidth",
    "ganada",
    "hebrew1",
    "hebrew2",
    "hindiConsonants",
    "hindiCounting",
    "hindiNumbers",
    "hindiVowels",
    "ideographDigital",
    "ideographEnclosedCircle",
    "ideographLegalTraditional",
    "ideographTraditional",
    "iroha",
    "japaneseCounting",
    "japaneseDigitalTenThousand",
    "japaneseLegal",
    "koreanCounting",
    "koreanDigital",
    "koreanDigital2",
    "koreanLegal",
    "lowerLetter",
    "lowerRoman",
    "numberInDash",
    "ordinal",
    "ordinalText",
    "russianLower",
    "russianUpper",
    "taiwaneseCounting",
    "taiwaneseCountingThousand",
    "taiwaneseDigital",
    "thaiCounting",
    "thaiLetters",
    "thaiNumbers",
    "upperLetter",
    "upperRoman",
    "vietnameseCounting",
}
UNORDERED_LIST_FORMATS = {"bullet", "none"}
PARAGRAPH_BOUNDARY_REPORTS_DIR = Path(".run") / "paragraph_boundary_reports"
RELATION_NORMALIZATION_REPORTS_DIR = Path(".run") / "relation_normalization_reports"
LAYOUT_CLEANUP_REPORTS_DIR = Path(".run") / "layout_cleanup_reports"
PARAGRAPH_BOUNDARY_AI_REVIEW_DIR = Path(".run") / "paragraph_boundary_ai_review"
_TYPOGRAPHIC_BULLET_CHARS = {"\u2014", "\u2013"}
_INLINE_BREAK_PATTERN = re.compile(r"<br\s*/?>", re.IGNORECASE)
_TOC_HEADER_LINE_VALUES = {"contents", "содержание"}
_TOC_CANDIDATE_WORD_PATTERN = re.compile(r"\w+(?:[-']\w+)*", re.UNICODE)
_STAGE0_TOC_HEADER_VALUES = {"contents", "table of contents", "содержание"}
_STAGE0_TOC_SUFFIX_PATTERN = re.compile(r"\.{2,}\s*\d+\s*$")
_STAGE0_ISOLATED_MARKER_PATTERN = re.compile(r"^(?:\s*[•●\-*]\s*|\s*\d+[\.)]\s*)$")
_STAGE0_SCRIPTURE_REFERENCE_PATTERN = re.compile(r"\b(?:[A-Za-zА-Яа-яЁё]+)\s+\d+:\d+(?:-\d+)?\b")
_STAGE0_SPACING_BEFORE_PATTERN = re.compile(r"<(?:\w+:)?spacing\b[^>]*\b(?:\w+:)?before=\"(?P<before>\d+)\"")
_STAGE0_SPACING_BEFORE_AUTOSPACING_PATTERN = re.compile(
    r"<(?:\w+:)?spacing\b[^>]*\b(?:\w+:)?beforeAutospacing=\"(?P<flag>[^\"]+)\""
)


def extract_paragraph_units_from_docx(uploaded_file) -> list[ParagraphUnit]:
    paragraphs, _ = extract_document_content_from_docx(uploaded_file)
    return paragraphs


def extract_inline_images(uploaded_file) -> list[ImageAsset]:
    _, image_assets = extract_document_content_from_docx(uploaded_file)
    return image_assets


def extract_document_content_from_docx(uploaded_file) -> tuple[list[ParagraphUnit], list[ImageAsset]]:
    paragraphs, image_assets, _, _, _, _, _ = extract_document_content_with_normalization_reports(uploaded_file)
    return paragraphs, image_assets


def extract_document_content_with_normalization_reports(
    uploaded_file,
    *,
    app_config: Mapping[str, object] | None = None,
) -> tuple[
    list[ParagraphUnit],
    list[ImageAsset],
    ParagraphBoundaryNormalizationReport,
    list[ParagraphRelation],
    RelationNormalizationReport,
    LayoutArtifactCleanupReport,
    StructureRepairReport,
]:
    source_bytes = _read_uploaded_docx_bytes(uploaded_file)
    validate_docx_source_bytes(source_bytes)
    document = Document(BytesIO(source_bytes))
    scan_origin = classify_document_scan_origin(source_bytes)
    raw_blocks, image_assets = _build_raw_document_blocks(document, is_scan_origin=scan_origin.is_scan_origin)
    normalization_mode, save_boundary_debug_artifacts = _resolve_paragraph_boundary_normalization_settings()
    normalized_blocks, boundary_report = _normalize_paragraph_boundaries(raw_blocks, mode=normalization_mode)
    structure_recovery_enabled, structure_recovery_mode = _resolve_structure_recovery_runtime(app_config=app_config)
    paragraphs = _build_logical_paragraph_units(
        normalized_blocks,
        structure_recovery_enabled=structure_recovery_enabled,
        structure_recovery_mode=structure_recovery_mode,
    )
    paragraphs = _normalize_inline_break_paragraphs(
        paragraphs,
        signal_only=structure_recovery_enabled and structure_recovery_mode == "ai_first",
    )
    promote_short_standalone_headings(
        paragraphs,
        structure_recovery_enabled=structure_recovery_enabled,
        structure_recovery_mode=structure_recovery_mode,
    )
    normalize_front_matter_display_title(
        paragraphs,
        structure_recovery_enabled=structure_recovery_enabled,
        structure_recovery_mode=structure_recovery_mode,
    )
    (
        cleanup_enabled,
        cleanup_min_repeat_count,
        cleanup_max_repeated_text_chars,
        cleanup_save_debug_artifacts,
        cleanup_mode,
    ) = _resolve_layout_artifact_cleanup_settings(app_config=app_config)
    paragraphs, cleanup_report = clean_paragraph_layout_artifacts(
        paragraphs,
        enabled=cleanup_enabled,
        min_repeat_count=cleanup_min_repeat_count,
        max_repeated_text_chars=cleanup_max_repeated_text_chars,
        cleanup_mode=cleanup_mode,
        structure_recovery_enabled=structure_recovery_enabled,
        structure_recovery_mode=structure_recovery_mode,
    )
    paragraphs, structure_repair_report = repair_pdf_derived_structure(
        paragraphs,
        app_config=app_config,
        structure_recovery_enabled=structure_recovery_enabled,
        structure_recovery_mode=structure_recovery_mode,
    )
    _reassign_paragraph_identities(paragraphs)
    (
        relation_enabled,
        relation_profile,
        enabled_relation_kinds,
        save_relation_debug_artifacts,
    ) = _resolve_relation_normalization_settings()
    (
        ai_review_enabled,
        ai_review_mode,
        ai_review_candidate_limit,
        ai_review_timeout_seconds,
        ai_review_max_tokens_per_candidate,
        ai_review_model,
    ) = _resolve_paragraph_boundary_ai_review_settings()
    try:
        relations, relation_report = build_paragraph_relations(
            paragraphs,
            enabled_relation_kinds=enabled_relation_kinds if relation_enabled else (),
            structure_phase="pre_ai_diagnostic",
        )
    except TypeError as exc:
        if "structure_phase" not in str(exc):
            raise
        relations, relation_report = build_paragraph_relations(
            paragraphs,
            enabled_relation_kinds=enabled_relation_kinds if relation_enabled else (),
        )
    apply_relation_side_effects(paragraphs, relations)
    reclassify_adjacent_captions(
        paragraphs,
        structure_recovery_enabled=structure_recovery_enabled,
        structure_recovery_mode=structure_recovery_mode,
    )
    _annotate_stage0_structure_signals(paragraphs)

    if ai_review_enabled and ai_review_mode != "off":
        _run_paragraph_boundary_ai_review(
            source_name=resolve_uploaded_filename(uploaded_file),
            source_bytes=source_bytes,
            mode=ai_review_mode,
            model=ai_review_model,
            raw_blocks=raw_blocks,
            paragraphs=paragraphs,
            boundary_report=boundary_report,
            relation_report=relation_report,
            candidate_limit=ai_review_candidate_limit,
            timeout_seconds=ai_review_timeout_seconds,
            max_tokens_per_candidate=ai_review_max_tokens_per_candidate,
        )

    if save_boundary_debug_artifacts:
        _write_paragraph_boundary_report_artifact(
            source_name=resolve_uploaded_filename(uploaded_file),
            source_bytes=source_bytes,
            mode=normalization_mode,
            report=boundary_report,
        )
    if relation_enabled and save_relation_debug_artifacts:
        _write_relation_normalization_report_artifact(
            source_name=resolve_uploaded_filename(uploaded_file),
            source_bytes=source_bytes,
            profile=relation_profile,
            enabled_relation_kinds=enabled_relation_kinds,
            report=relation_report,
        )
    cleanup_report_path = None
    if cleanup_save_debug_artifacts:
        cleanup_report_path = _write_layout_cleanup_report_artifact(
            source_name=resolve_uploaded_filename(uploaded_file),
            source_bytes=source_bytes,
            report=cleanup_report,
        )
    cleanup_report.artifact_path = cleanup_report_path
    return paragraphs, image_assets, boundary_report, relations, relation_report, cleanup_report, structure_repair_report


def extract_document_content_with_boundary_report(
    uploaded_file,
) -> tuple[list[ParagraphUnit], list[ImageAsset], ParagraphBoundaryNormalizationReport]:
    paragraphs, image_assets, boundary_report, _, _, _, _ = extract_document_content_with_normalization_reports(uploaded_file)
    return paragraphs, image_assets, boundary_report


def _build_paragraph_text_with_placeholders(
    paragraph,
    image_assets: list[ImageAsset],
    *,
    include_image_placeholders: bool = True,
    allow_run_markdown: bool = True,
) -> str:
    parts: list[str] = []
    for child in paragraph._element:
        local_name = xml_local_name(child.tag)
        if local_name == "r":
            parts.append(
                _render_run_element(
                    child,
                    paragraph.part,
                    image_assets,
                    allow_hyperlink_markdown=allow_run_markdown,
                    include_image_placeholders=include_image_placeholders,
                )
            )
            continue
        if local_name == "hyperlink":
            parts.append(
                _render_hyperlink_element(
                    child,
                    paragraph,
                    image_assets,
                    allow_hyperlink_markdown=allow_run_markdown,
                    include_image_placeholders=include_image_placeholders,
                )
            )
    return "".join(parts)


def _build_raw_document_blocks(
    document,
    *,
    is_scan_origin: bool = False,
) -> tuple[list[RawBlock], list[ImageAsset]]:
    raw_blocks: list[RawBlock] = []
    image_assets: list[ImageAsset] = []
    table_count = 0

    for block_kind, block in _iter_document_block_items(document):
        if block_kind == "paragraph":
            for raw_block in _build_raw_paragraph_blocks(cast(Paragraph, block), image_assets, raw_index=len(raw_blocks)):
                raw_blocks.append(raw_block)
            continue
        table_count += 1
        if is_scan_origin:
            # Scan-origin (OCR) documents: the "table" is a scanned column layout,
            # not authored tabular data — flatten it into linear body paragraphs.
            raw_blocks.extend(_build_flattened_table_blocks(cast(Table, block), image_assets, start_raw_index=len(raw_blocks)))
            continue
        raw_block = _build_raw_table(
            cast(Table, block),
            image_assets,
            raw_index=len(raw_blocks),
            asset_id=f"table_{table_count:03d}",
        )
        if raw_block is not None:
            raw_blocks.append(raw_block)

    return raw_blocks, image_assets


def _build_flattened_table_blocks(
    table: Table,
    image_assets: list[ImageAsset],
    *,
    start_raw_index: int,
) -> list[RawParagraph]:
    blocks: list[RawParagraph] = []
    for line in _flatten_table_lines_impl(
        table,
        image_assets,
        build_paragraph_text_with_placeholders=_build_paragraph_text_with_placeholders,
    ):
        raw_index = start_raw_index + len(blocks)
        blocks.append(
            RawParagraph(
                raw_index=raw_index,
                text=line,
                style_name="",
                role_hint="body",
                origin_raw_indexes=(raw_index,),
                origin_raw_texts=(line,),
                layout_origin="table_flattened",
                boundary_source="raw",
                boundary_confidence="high",
            )
        )
    return blocks


def _build_raw_paragraph_blocks(paragraph, image_assets: list[ImageAsset], *, raw_index: int) -> list[RawParagraph]:
    raw_blocks: list[RawParagraph] = []
    has_textboxes = _paragraph_has_textbox_content(paragraph)
    direct_text = _build_paragraph_text_with_placeholders(paragraph, image_assets, include_image_placeholders=not has_textboxes)

    if direct_text.strip():
        raw_block = _build_raw_paragraph(
            paragraph,
            image_assets,
            raw_index=raw_index,
            text_override=direct_text,
            layout_origin="paragraph",
        )
        if raw_block is not None:
            raw_blocks.append(raw_block)

    for textbox_paragraph in _iter_textbox_paragraphs(paragraph):
        raw_block = _build_raw_paragraph(
            textbox_paragraph,
            image_assets,
            raw_index=raw_index + len(raw_blocks),
            allow_run_markdown=False,
            layout_origin="textbox",
        )
        if raw_block is not None:
            if raw_blocks and raw_blocks[-1].text == raw_block.text:
                continue
            raw_blocks.append(raw_block)

    return raw_blocks


def _build_raw_paragraph(
    paragraph,
    image_assets: list[ImageAsset],
    *,
    raw_index: int,
    text_override: str | None = None,
    allow_run_markdown: bool = True,
    layout_origin: str = "paragraph",
) -> RawParagraph | None:
    text = (
        text_override
        if text_override is not None
        else _build_paragraph_text_with_placeholders(paragraph, image_assets, allow_run_markdown=allow_run_markdown)
    ).strip()
    if not text:
        return None

    style_name = paragraph.style.name if paragraph.style and paragraph.style.name else ""
    normalized_style = style_name.strip().lower()
    explicit_heading_level = extract_explicit_heading_level(paragraph, style_name)
    heading_level = explicit_heading_level
    heading_source = "explicit" if explicit_heading_level is not None else None
    if heading_level is None and not is_caption_style(normalized_style):
        if is_probable_heading(paragraph, text, normalized_style) or _is_markdown_strong_heading_candidate(
            text,
            normalized_style,
        ):
            heading_level = infer_heuristic_heading_level(text)
            heading_source = "heuristic"
    role = classify_paragraph_role(text, style_name, heading_level=heading_level)
    list_metadata = _extract_paragraph_list_metadata(paragraph, text, style_name, role)
    if (
        role != "list"
        and list_metadata["list_kind"] is not None
        and not _should_preserve_heading_role_against_list_metadata(
            role=role,
            heading_level=heading_level,
            text=text,
        )
    ):
        role = "list"
    if role == "list" and list_metadata.get("_is_typographic_emdash_bullet"):
        role = "body"
        list_metadata = _empty_list_metadata()
    if role == "body":
        compact_toc_text = _build_compact_toc_run_cluster_text(paragraph)
        if compact_toc_text is not None:
            text = compact_toc_text
    asset_id = _extract_paragraph_asset_id(text, role=role)
    role_confidence = infer_role_confidence(
        role=role,
        text=text,
        normalized_style=normalized_style,
        explicit_heading_level=explicit_heading_level,
        heading_source=heading_source,
    )
    return RawParagraph(
        raw_index=raw_index,
        text=text,
        style_name=style_name,
        paragraph_properties_xml=_extract_paragraph_properties_xml(paragraph),
        paragraph_alignment=resolve_paragraph_alignment(paragraph),
        is_bold=paragraph_is_effectively_bold(paragraph),
        is_italic=paragraph_is_effectively_italic(paragraph),
        font_size_pt=resolve_effective_paragraph_font_size(paragraph),
        explicit_heading_level=explicit_heading_level,
        heading_level=heading_level,
        heading_source=heading_source,
        list_kind=cast(str | None, list_metadata["list_kind"]),
        list_level=cast(int, list_metadata["list_level"]),
        list_numbering_format=cast(str | None, list_metadata["list_numbering_format"]),
        list_num_id=cast(str | None, list_metadata["list_num_id"]),
        list_abstract_num_id=cast(str | None, list_metadata["list_abstract_num_id"]),
        list_num_xml=cast(str | None, list_metadata["list_num_xml"]),
        list_abstract_num_xml=cast(str | None, list_metadata["list_abstract_num_xml"]),
        role_hint=role,
        source_xml_fingerprint=build_source_xml_fingerprint(paragraph),
        origin_raw_indexes=(raw_index,),
        origin_raw_texts=(text,),
        layout_origin=layout_origin,
        boundary_source="raw",
        boundary_confidence="explicit" if role_confidence == "explicit" else "high",
    )


def _should_preserve_heading_role_against_list_metadata(
    *,
    role: str,
    heading_level: int | None,
    text: str,
) -> bool:
    return role == "heading" and heading_level is not None and has_heading_text_signal(text)


def _is_markdown_strong_heading_candidate(text: str, normalized_style: str) -> bool:
    stripped_text = text.strip()
    normalized_text = stripped_text.replace("**", "").replace("*", "").strip()
    word_count = len(normalized_text.split())
    if normalized_text.endswith(".") and word_count > 4:
        return False
    return (
        stripped_text.startswith("**")
        and stripped_text.endswith("**")
        and not is_caption_style(normalized_style)
        and 0 < len(normalized_text) <= 140
        and word_count <= 18
        and has_heading_text_signal(normalized_text)
    )


def _build_raw_table(table: Table, image_assets: list[ImageAsset], *, raw_index: int, asset_id: str) -> RawTable | None:
    return _build_raw_table_impl(
        table,
        image_assets,
        raw_index=raw_index,
        asset_id=asset_id,
        build_paragraph_text_with_placeholders=_build_paragraph_text_with_placeholders,
    )


def _build_logical_paragraph_units(
    raw_blocks: list[RawBlock],
    *,
    structure_recovery_enabled: bool = False,
    structure_recovery_mode: str = "legacy",
) -> list[ParagraphUnit]:
    paragraphs: list[ParagraphUnit] = []
    signal_only = structure_recovery_enabled and structure_recovery_mode == "ai_first"
    for block in raw_blocks:
        if isinstance(block, RawParagraph):
            role = block.role_hint
            structural_role = block.role_hint
            heading_level = block.heading_level
            heading_source = block.heading_source
            heuristic_role_hint = None
            heuristic_heading_level_hint = None
            if signal_only and block.role_hint == "heading" and block.heading_source == "heuristic":
                role = "body"
                structural_role = "body"
                heuristic_role_hint = "heading"
                heuristic_heading_level_hint = block.heading_level or 2
                heading_level = None
                heading_source = None
            paragraph = ParagraphUnit(
                text=block.text,
                role=role,
                asset_id=_extract_paragraph_asset_id(block.text, role=role),
                paragraph_properties_xml=block.paragraph_properties_xml,
                paragraph_alignment=block.paragraph_alignment,
                heading_level=heading_level,
                heading_source=heading_source,
                list_kind=block.list_kind,
                list_level=block.list_level,
                list_numbering_format=block.list_numbering_format,
                list_num_id=block.list_num_id,
                list_abstract_num_id=block.list_abstract_num_id,
                list_num_xml=block.list_num_xml,
                list_abstract_num_xml=block.list_abstract_num_xml,
                structural_role=structural_role,
                role_confidence=infer_role_confidence(
                    role=role,
                    text=block.text,
                    normalized_style=block.style_name.strip().lower(),
                    explicit_heading_level=block.explicit_heading_level,
                    heading_source=heading_source,
                ),
                heuristic_role_hint=heuristic_role_hint,
                heuristic_heading_level_hint=heuristic_heading_level_hint,
                style_name=block.style_name,
                is_bold=block.is_bold,
                is_italic=block.is_italic,
                font_size_pt=block.font_size_pt,
                origin_raw_indexes=list(block.origin_raw_indexes or (block.raw_index,)),
                origin_raw_texts=list(block.origin_raw_texts or (block.text,)),
                layout_origin=block.layout_origin,
                boundary_source=block.boundary_source,
                boundary_confidence=block.boundary_confidence,
                boundary_rationale=block.boundary_rationale,
            )
        else:
            paragraph = ParagraphUnit(
                text=block.html_text,
                role="table",
                asset_id=block.asset_id,
                structural_role="table",
                role_confidence="explicit",
                origin_raw_indexes=[block.raw_index],
                origin_raw_texts=[block.html_text],
            )
        _assign_paragraph_identity(paragraph, len(paragraphs))
        paragraphs.append(paragraph)
    return paragraphs


def _normalize_inline_break_paragraphs(paragraphs: list[ParagraphUnit], *, signal_only: bool = False) -> list[ParagraphUnit]:
    normalized: list[ParagraphUnit] = []
    for paragraph in paragraphs:
        if paragraph.role in {"image", "table"} or not _INLINE_BREAK_PATTERN.search(paragraph.text):
            normalized.append(paragraph)
            continue

        lines = _split_inline_break_lines(paragraph.text)
        if not lines:
            continue
        if len(lines) < 2:
            normalized.append(_copy_paragraph_unit(paragraph, text=_join_inline_break_lines(lines)))
            continue

        if _should_expand_inline_break_paragraph(paragraph, lines):
            normalized.extend(_expand_inline_break_paragraph(paragraph, lines, signal_only=signal_only))
            continue

        normalized.append(_copy_paragraph_unit(paragraph, text=_join_inline_break_lines(lines)))

    _annotate_toc_region_candidates(normalized, signal_only=signal_only)
    for index, paragraph in enumerate(normalized):
        _assign_paragraph_identity(paragraph, index)
    return normalized


def _split_inline_break_lines(text: str) -> list[str]:
    return [part.strip() for part in _INLINE_BREAK_PATTERN.split(text) if part.strip()]


def _join_inline_break_lines(lines: list[str]) -> str:
    return " ".join(line.strip() for line in lines if line.strip())


def _copy_paragraph_unit(paragraph: ParagraphUnit, *, text: str) -> ParagraphUnit:
    return replace(
        paragraph,
        text=text,
        paragraph_id="",
        source_index=paragraph.source_index,
        logical_index=-1,
        origin_raw_indexes=list(paragraph.origin_raw_indexes),
        origin_raw_texts=list(paragraph.origin_raw_texts),
    )


def _annotate_stage0_structure_signals(paragraphs: list[ParagraphUnit]) -> None:
    style_cluster_ids = _build_stage0_style_cluster_ids(paragraphs)
    font_size_z_scores = _build_stage0_font_size_z_scores(paragraphs)
    last_index = max(len(paragraphs) - 1, 1)

    for position, paragraph in enumerate(paragraphs):
        text = str(getattr(paragraph, "text", "") or "").strip()
        hint_texts = [str(getattr(hint, "text", "") or "").strip() for hint in (paragraph.heuristic_embedded_structure_hints or [])]
        hint_structural_roles = [
            str(getattr(hint, "structural_role", "body") or "body")
            for hint in (paragraph.heuristic_embedded_structure_hints or [])
        ]
        paragraph.style_cluster_id = style_cluster_ids[position]
        paragraph.font_size_z_score = font_size_z_scores[position]
        paragraph.position_fraction = round(position / last_index, 3)
        paragraph.page_number = _extract_stage0_page_number_hint(text, paragraph)
        paragraph.vertical_gap_before_pt = _extract_stage0_vertical_gap_before_pt(paragraph)
        paragraph.is_isolated_marker = _is_stage0_isolated_marker_text(text) or any(
            _is_stage0_isolated_marker_text(hint_text) for hint_text in hint_texts
        )
        paragraph.toc_pattern_hint = _is_stage0_toc_pattern_hint(text, paragraph) or any(
            role in {"toc_header", "toc_entry"} for role in hint_structural_roles
        )
        paragraph.scripture_reference_hint = _is_stage0_scripture_reference_text(text) or any(
            _is_stage0_scripture_reference_text(hint_text) for hint_text in hint_texts
        )
        paragraph.boundary_normalization_applied = str(getattr(paragraph, "boundary_source", "raw") or "raw") != "raw"


def _build_stage0_style_cluster_ids(paragraphs: list[ParagraphUnit]) -> list[int | None]:
    normalized_styles = [str(getattr(paragraph, "style_name", "") or "").strip().lower() for paragraph in paragraphs]
    nonempty_styles = [style for style in normalized_styles if style]
    if not nonempty_styles:
        return [None for _ in paragraphs]
    default_style = Counter(nonempty_styles).most_common(1)[0][0]
    cluster_map: dict[str, int] = {}
    next_cluster_id = 1
    cluster_ids: list[int | None] = []
    for style in normalized_styles:
        if not style or style == default_style:
            cluster_ids.append(None)
            continue
        if style not in cluster_map:
            cluster_map[style] = next_cluster_id
            next_cluster_id += 1
        cluster_ids.append(cluster_map[style])
    return cluster_ids


def _build_stage0_font_size_z_scores(paragraphs: list[ParagraphUnit]) -> list[float | None]:
    sizes = [getattr(paragraph, "font_size_pt", None) for paragraph in paragraphs]
    numeric_sizes = [float(size) for size in sizes if isinstance(size, (int, float))]
    if not numeric_sizes:
        return [None for _ in paragraphs]
    mean = sum(numeric_sizes) / len(numeric_sizes)
    variance = sum((size - mean) ** 2 for size in numeric_sizes) / len(numeric_sizes)
    stddev = variance ** 0.5
    z_scores: list[float | None] = []
    for size in sizes:
        if not isinstance(size, (int, float)):
            z_scores.append(None)
            continue
        if isclose(stddev, 0.0):
            z_scores.append(0.0)
            continue
        z_scores.append(round((float(size) - mean) / stddev, 1))
    return z_scores


def _extract_stage0_page_number_hint(text: str, paragraph: ParagraphUnit) -> int | None:
    if not bool(getattr(paragraph, "is_likely_page_number", False)):
        return None
    stripped = text.strip()
    if stripped.isdigit():
        return int(stripped)
    match = re.search(r"\b(\d{1,4})\b", stripped)
    if match is None:
        return None
    return int(match.group(1))


def _extract_stage0_vertical_gap_before_pt(paragraph: ParagraphUnit) -> float | None:
    direct_value = getattr(paragraph, "vertical_gap_before_pt", None)
    if isinstance(direct_value, int | float):
        return _round_stage0_vertical_gap_before_pt(float(direct_value))

    paragraph_properties_xml = str(getattr(paragraph, "paragraph_properties_xml", "") or "").strip()
    if not paragraph_properties_xml:
        return None
    autospacing_match = _STAGE0_SPACING_BEFORE_AUTOSPACING_PATTERN.search(paragraph_properties_xml)
    if autospacing_match and str(autospacing_match.group("flag") or "").strip().lower() in {"1", "true", "on"}:
        return None
    spacing_match = _STAGE0_SPACING_BEFORE_PATTERN.search(paragraph_properties_xml)
    if spacing_match is None:
        return None
    try:
        before_twips = int(spacing_match.group("before"))
    except (TypeError, ValueError):
        return None
    return _round_stage0_vertical_gap_before_pt(before_twips / 20.0)


def _round_stage0_vertical_gap_before_pt(value: float) -> float:
    return round(value * 2.0) / 2.0


def _is_stage0_isolated_marker_text(text: str) -> bool:
    return bool(_STAGE0_ISOLATED_MARKER_PATTERN.fullmatch(str(text or "").strip()))


def _is_stage0_toc_pattern_hint(text: str, paragraph: ParagraphUnit) -> bool:
    if normalize_heuristic_structural_role_hint(getattr(paragraph, "heuristic_structural_role_hint", None)) in {"toc_header", "toc_entry"}:
        return True
    if getattr(paragraph, "structural_role", None) in {"toc_header", "toc_entry"}:
        return True
    normalized = str(text or "").strip().casefold()
    if normalized in _STAGE0_TOC_HEADER_VALUES:
        return True
    return bool(_STAGE0_TOC_SUFFIX_PATTERN.search(str(text or "").strip()))


def _is_stage0_scripture_reference_text(text: str) -> bool:
    return bool(_STAGE0_SCRIPTURE_REFERENCE_PATTERN.search(str(text or "").strip()))


def _extract_paragraph_properties_xml(paragraph) -> str | None:
    paragraph_properties = find_child_element(paragraph._element, "pPr")
    if paragraph_properties is None:
        return None
    return paragraph_properties.xml


def _build_compact_toc_run_cluster_text(paragraph) -> str | None:
    segments = _extract_compact_run_clusters(paragraph)
    if not _is_compact_toc_run_cluster(segments):
        return None
    return "<br/>".join(segments)


def _extract_compact_run_clusters(paragraph) -> list[str]:
    segments: list[str] = []
    current_parts: list[str] = []

    for child in paragraph._element:
        if xml_local_name(child.tag) != "r":
            continue
        raw_text = _extract_run_text(child)
        if not raw_text:
            continue
        if "<br/>" in raw_text or "\t" in raw_text:
            return []
        formatted_text = _apply_run_markdown(raw_text, child)
        if not raw_text.strip():
            if current_parts:
                segment = "".join(current_parts).strip()
                if segment:
                    segments.append(segment)
                current_parts = []
            continue
        current_parts.append(formatted_text)

    if current_parts:
        segment = "".join(current_parts).strip()
        if segment:
            segments.append(segment)

    return segments


def _is_compact_toc_run_cluster(segments: list[str]) -> bool:
    if len(segments) < 2:
        return False

    normalized_segments = [segment.strip() for segment in segments if segment.strip()]
    if len(normalized_segments) < 2:
        return False
    if not all(_is_toc_candidate_text(segment) for segment in normalized_segments):
        return False

    word_counts = [len(_TOC_CANDIDATE_WORD_PATTERN.findall(segment)) for segment in normalized_segments]
    total_words = sum(word_counts)
    if len(normalized_segments) == 2:
        if total_words > 20:
            return False
        if min(word_counts) < 3:
            return False
        if not (any(count >= 4 for count in word_counts) or any(has_heading_text_signal(segment) for segment in normalized_segments)):
            return False
        return True

    if total_words > 14:
        return False
    return all(count <= 5 for count in word_counts)


def _should_expand_inline_break_paragraph(paragraph: ParagraphUnit, lines: list[str]) -> bool:
    if paragraph.role not in {"body", "heading", "list"}:
        return False
    if len(lines) < 2:
        return False
    if _is_toc_header_line(lines[0]):
        return sum(1 for line in lines[1:] if _is_toc_candidate_text(line)) >= 2
    return len(lines) >= 2 and all(_is_toc_candidate_text(line) for line in lines)


def _is_toc_header_line(text: str) -> bool:
    return text.strip().lower() in _TOC_HEADER_LINE_VALUES


def _is_toc_candidate_text(text: str) -> bool:
    stripped = text.strip()
    if not stripped:
        return False
    if IMAGE_PLACEHOLDER_PATTERN.search(stripped):
        return False
    if stripped.startswith("<table"):
        return False
    if len(stripped) > 160:
        return False
    word_count = len(_TOC_CANDIDATE_WORD_PATTERN.findall(stripped))
    if word_count == 0 or word_count > 16:
        return False
    if stripped.endswith((".", ";")):
        return False
    return True


def _expand_inline_break_paragraph(paragraph: ParagraphUnit, lines: list[str], *, signal_only: bool) -> list[ParagraphUnit]:
    expanded: list[ParagraphUnit] = []
    header_cluster = _is_toc_header_line(lines[0]) and len(lines) >= 3
    for index, line in enumerate(lines):
        clone = _copy_paragraph_unit(paragraph, text=line)
        if header_cluster and index == 0:
            _apply_or_hint_stage0_toc_role(clone, structural_role="toc_header", signal_only=signal_only)
        elif header_cluster or _is_toc_candidate_text(line):
            _apply_or_hint_stage0_toc_role(clone, structural_role="toc_entry", signal_only=signal_only)
        expanded.append(clone)
    return expanded


def _annotate_toc_region_candidates(paragraphs: list[ParagraphUnit], *, signal_only: bool) -> None:
    index = 0
    while index < len(paragraphs):
        paragraph = paragraphs[index]
        if not _is_toc_header_line(paragraph.text):
            index += 1
            continue

        look_ahead = index + 1
        while look_ahead < len(paragraphs) and _is_toc_candidate_paragraph(paragraphs[look_ahead]):
            look_ahead += 1

        if look_ahead - index >= 3:
            _apply_or_hint_stage0_toc_role(paragraph, structural_role="toc_header", signal_only=signal_only)
            for toc_index in range(index + 1, look_ahead):
                toc_paragraph = paragraphs[toc_index]
                _apply_or_hint_stage0_toc_role(toc_paragraph, structural_role="toc_entry", signal_only=signal_only)
            index = look_ahead
            continue

        index += 1


def _is_toc_candidate_paragraph(paragraph: ParagraphUnit) -> bool:
    if paragraph.role in {"image", "table"}:
        return False
    if paragraph.attached_to_asset_id is not None:
        return False
    return _is_toc_candidate_text(paragraph.text)


def _apply_or_hint_stage0_toc_role(paragraph: ParagraphUnit, *, structural_role: str, signal_only: bool) -> None:
    paragraph.heuristic_structural_role_hint = normalize_heuristic_structural_role_hint(structural_role)
    if signal_only:
        return
    if paragraph.role == "heading" and paragraph.heading_source != "explicit":
        paragraph.role = "body"
        paragraph.heading_level = None
        paragraph.heading_source = None
    paragraph.structural_role = normalize_heuristic_structural_role_hint(structural_role) or "body"


def _resolve_paragraph_boundary_normalization_settings() -> tuple[str, bool]:
    return _resolve_paragraph_boundary_normalization_settings_impl(
        allowed_modes=PARAGRAPH_BOUNDARY_NORMALIZATION_MODE_VALUES,
    )


def _resolve_relation_normalization_settings() -> tuple[bool, str, tuple[str, ...], bool]:
    from docxaicorrector.document.relations import _resolve_relation_normalization_settings as _resolve_relation_normalization_settings_impl

    return _resolve_relation_normalization_settings_impl()


def _resolve_layout_artifact_cleanup_settings(*, app_config: Mapping[str, object] | None = None) -> tuple[bool, int, int, bool, str]:
    from docxaicorrector.core.config import load_app_config

    if app_config is None:
        app_config = load_app_config()
    return (
        bool(app_config.get("layout_artifact_cleanup_enabled", True)),
        _coerce_int_config_value_impl(app_config.get("layout_artifact_cleanup_min_repeat_count", 3), 3),
        _coerce_int_config_value_impl(app_config.get("layout_artifact_cleanup_max_repeated_text_chars", 80), 80),
        bool(app_config.get("layout_artifact_cleanup_save_debug_artifacts", True)),
        str(app_config.get("layout_artifact_cleanup_mode", "flag") or "flag").strip().lower() or "flag",
    )


def _resolve_structure_recovery_runtime(*, app_config: Mapping[str, object] | None) -> tuple[bool, str]:
    if app_config is None:
        return False, "legacy"
    structure_recovery_enabled = bool(app_config.get("structure_recovery_enabled", False))
    structure_recovery_mode = str(app_config.get("structure_recovery_mode", "ai_first") or "ai_first").strip().lower()
    if not structure_recovery_enabled:
        return False, "legacy"
    if structure_recovery_mode not in {"legacy", "ai_first"}:
        return structure_recovery_enabled, "ai_first"
    return structure_recovery_enabled, structure_recovery_mode


def _resolve_paragraph_boundary_ai_review_settings() -> tuple[bool, str, int, int, int, str]:
    return _resolve_paragraph_boundary_ai_review_settings_impl(
        allowed_modes=PARAGRAPH_BOUNDARY_AI_REVIEW_MODE_VALUES,
    )


def _coerce_int_config_value(value: object, default: int) -> int:
    return _coerce_int_config_value_impl(value, default)


def _build_ai_review_candidates(
    *,
    raw_blocks: list[RawBlock],
    paragraphs: list[ParagraphUnit],
    boundary_report: ParagraphBoundaryNormalizationReport,
    relation_report: RelationNormalizationReport,
    candidate_limit: int,
) -> list[dict[str, object]]:
    return _build_ai_review_candidates_impl(
        raw_blocks=raw_blocks,
        paragraphs=paragraphs,
        boundary_report=boundary_report,
        relation_report=relation_report,
        candidate_limit=candidate_limit,
    )


def _request_ai_review_recommendations(
    *,
    model: str,
    candidates: list[dict[str, object]],
    timeout_seconds: int,
    max_tokens_per_candidate: int,
) -> dict[str, dict[str, object]]:
    return _request_ai_review_recommendations_impl(
        model=model,
        candidates=candidates,
        timeout_seconds=timeout_seconds,
        max_tokens_per_candidate=max_tokens_per_candidate,
    )


def _write_paragraph_boundary_ai_review_artifact(
    *,
    source_name: str,
    source_bytes: bytes,
    mode: str,
    decisions: list[dict[str, object]],
    error_code: str | None = None,
) -> str | None:
    return _write_paragraph_boundary_ai_review_artifact_impl(
        source_name=source_name,
        source_bytes=source_bytes,
        mode=mode,
        decisions=decisions,
        target_dir=PARAGRAPH_BOUNDARY_AI_REVIEW_DIR,
        max_age_seconds=PARAGRAPH_BOUNDARY_AI_REVIEW_MAX_AGE_SECONDS,
        max_count=PARAGRAPH_BOUNDARY_AI_REVIEW_MAX_COUNT,
        error_code=error_code,
    )


def _run_paragraph_boundary_ai_review(
    *,
    source_name: str,
    source_bytes: bytes,
    mode: str,
    model: str,
    raw_blocks: list[RawBlock],
    paragraphs: list[ParagraphUnit],
    boundary_report: ParagraphBoundaryNormalizationReport,
    relation_report: RelationNormalizationReport,
    candidate_limit: int,
    timeout_seconds: int,
    max_tokens_per_candidate: int,
) -> str | None:
    return _run_paragraph_boundary_ai_review_impl(
        source_name=source_name,
        source_bytes=source_bytes,
        mode=mode,
        model=model,
        raw_blocks=raw_blocks,
        paragraphs=paragraphs,
        boundary_report=boundary_report,
        relation_report=relation_report,
        candidate_limit=candidate_limit,
        timeout_seconds=timeout_seconds,
        max_tokens_per_candidate=max_tokens_per_candidate,
        target_dir=PARAGRAPH_BOUNDARY_AI_REVIEW_DIR,
        max_age_seconds=PARAGRAPH_BOUNDARY_AI_REVIEW_MAX_AGE_SECONDS,
        max_count=PARAGRAPH_BOUNDARY_AI_REVIEW_MAX_COUNT,
        request_ai_review_recommendations_impl=_request_ai_review_recommendations,
    )


def _normalize_paragraph_boundaries(
    raw_blocks: list[RawBlock],
    *,
    mode: str,
) -> tuple[list[RawBlock], ParagraphBoundaryNormalizationReport]:
    return _normalize_paragraph_boundaries_impl(
        raw_blocks,
        mode=mode,
        detect_explicit_list_kind=detect_explicit_list_kind,
        has_heading_text_signal=has_heading_text_signal,
    )


def _evaluate_paragraph_boundary(left: RawParagraph, right: RawParagraph) -> ParagraphBoundaryDecision:
    return _evaluate_paragraph_boundary_impl(
        left,
        right,
        detect_explicit_list_kind=detect_explicit_list_kind,
        has_heading_text_signal=has_heading_text_signal,
    )


def _write_paragraph_boundary_report_artifact(
    *,
    source_name: str,
    source_bytes: bytes,
    mode: str,
    report: ParagraphBoundaryNormalizationReport,
) -> str | None:
    return _write_paragraph_boundary_report_artifact_impl(
        source_name=source_name,
        source_bytes=source_bytes,
        mode=mode,
        report=report,
        target_dir=PARAGRAPH_BOUNDARY_REPORTS_DIR,
        max_age_seconds=PARAGRAPH_BOUNDARY_REPORTS_MAX_AGE_SECONDS,
        max_count=PARAGRAPH_BOUNDARY_REPORTS_MAX_COUNT,
    )


def _write_relation_normalization_report_artifact(
    *,
    source_name: str,
    source_bytes: bytes,
    profile: str,
    enabled_relation_kinds: tuple[str, ...],
    report: RelationNormalizationReport,
) -> str | None:
    return _write_relation_normalization_report_artifact_impl(
        source_name=source_name,
        source_bytes=source_bytes,
        profile=profile,
        enabled_relation_kinds=enabled_relation_kinds,
        report=report,
        target_dir=RELATION_NORMALIZATION_REPORTS_DIR,
        max_age_seconds=RELATION_NORMALIZATION_REPORTS_MAX_AGE_SECONDS,
        max_count=RELATION_NORMALIZATION_REPORTS_MAX_COUNT,
    )


def _write_layout_cleanup_report_artifact(
    *,
    source_name: str,
    source_bytes: bytes,
    report: LayoutArtifactCleanupReport,
) -> str | None:
    return _write_layout_cleanup_report_artifact_impl(
        source_name=source_name,
        source_bytes=source_bytes,
        report=report,
        target_dir=LAYOUT_CLEANUP_REPORTS_DIR,
        max_age_seconds=LAYOUT_CLEANUP_REPORTS_MAX_AGE_SECONDS,
        max_count=LAYOUT_CLEANUP_REPORTS_MAX_COUNT,
    )


def _assign_paragraph_identity(paragraph: ParagraphUnit, logical_index: int) -> None:
    if int(getattr(paragraph, "source_index", -1)) < 0:
        paragraph.source_index = logical_index
    # logical_index is assigned only after the final extraction topology is known.
    # It is the dense Stage 1/2/3 coordinate for the current paragraph list, while
    # source_index and origin_raw_indexes retain provenance across earlier splits.
    paragraph.logical_index = logical_index
    paragraph.paragraph_id = f"p{logical_index:04d}"
    if not paragraph.structural_role or paragraph.structural_role == "body":
        paragraph.structural_role = paragraph.role
    if not paragraph.origin_raw_indexes:
        paragraph.origin_raw_indexes = [int(getattr(paragraph, "source_index", logical_index))]
    if not paragraph.origin_raw_texts:
        paragraph.origin_raw_texts = [paragraph.text]


def _reassign_paragraph_identities(paragraphs: list[ParagraphUnit]) -> None:
    for logical_index, paragraph in enumerate(paragraphs):
        _assign_paragraph_identity(paragraph, logical_index)


def _iter_document_block_items(document):
    for child in document.element.body.iterchildren():
        local_name = xml_local_name(child.tag)
        if local_name == "p":
            yield "paragraph", Paragraph(child, document)
        elif local_name == "tbl":
            yield "table", Table(child, document)


def _read_uploaded_docx_bytes(uploaded_file) -> bytes:
    try:
        source_bytes = read_uploaded_file_bytes(uploaded_file)
    except ValueError as exc:
        raise ValueError("Не удалось прочитать содержимое DOCX-файла.") from exc
    if zipfile.is_zipfile(BytesIO(source_bytes)):
        return source_bytes
    source_name = resolve_uploaded_filename(uploaded_file)
    raise ValueError(
        "Ожидался уже нормализованный DOCX-архив, но получен ненормализованный входной файл: "
        f"{source_name}"
    )


def _paragraph_has_textbox_content(paragraph) -> bool:
    return any(True for _ in _iter_textbox_content_elements(paragraph._element))


def _iter_textbox_content_elements(element):
    for descendant in element.iter():
        if descendant is element:
            continue
        if xml_local_name(descendant.tag) == "txbxContent":
            yield descendant


def _iter_textbox_paragraphs(paragraph):
    for textbox_content in _iter_textbox_content_elements(paragraph._element):
        for child in textbox_content:
            if xml_local_name(child.tag) == "p":
                yield Paragraph(child, paragraph.part)


def _render_hyperlink_element(
    hyperlink_element,
    paragraph,
    image_assets: list[ImageAsset],
    *,
    include_image_placeholders: bool = True,
    allow_hyperlink_markdown: bool = True,
) -> str:
    text_parts: list[str] = []
    for child in hyperlink_element:
        if xml_local_name(child.tag) != "r":
            continue
        text_parts.append(
            _render_run_element(
                child,
                paragraph.part,
                image_assets,
                allow_hyperlink_markdown=False,
                include_image_placeholders=include_image_placeholders,
            )
        )

    text = "".join(text_parts)
    if not text.strip():
        return text

    relationship_id = hyperlink_element.get(f"{{{RELATIONSHIP_NAMESPACE}}}id")
    if not allow_hyperlink_markdown or not relationship_id:
        return text

    relationship = paragraph.part.rels.get(relationship_id)
    url = getattr(relationship, "target_ref", None)
    if not url:
        return text
    return f"[{text}]({url})"


def _render_run_element(
    run_element,
    part,
    image_assets: list[ImageAsset],
    *,
    allow_hyperlink_markdown: bool = True,
    include_image_placeholders: bool = True,
) -> str:
    text = _extract_run_text(run_element)
    formatted_text = _apply_run_markdown(text, run_element) if allow_hyperlink_markdown else text
    image_placeholders = _extract_run_image_placeholders(run_element, part, image_assets) if include_image_placeholders else []
    return formatted_text + "".join(image_placeholders)


def _extract_run_text(run_element) -> str:
    text_parts: list[str] = []
    for child in run_element:
        local_name = xml_local_name(child.tag)
        if local_name == "t":
            # Only <w:t> carries visible run text. <w:instrText> holds field
            # instruction codes (HYPERLINK "…", PAGEREF, TOC \o …, MERGEFORMAT)
            # and <w:delText> holds tracked-change deletions — neither is visible
            # content, so including them leaks field codes / removed text into the
            # translated body.
            text_parts.append(child.text or "")
            continue
        if local_name == "tab":
            text_parts.append("\t")
            continue
        if local_name in {"br", "cr"}:
            text_parts.append("<br/>")
    return "".join(text_parts)


def _apply_run_markdown(text: str, run_element) -> str:
    if not text or not text.strip():
        return text

    run_properties = find_child_element(run_element, "rPr")
    if run_properties is None:
        return text

    is_bold = _run_toggle_property_is_on(run_properties, "b")
    is_italic = _run_toggle_property_is_on(run_properties, "i")
    is_underline = _run_toggle_property_is_on(run_properties, "u")
    vertical_align = _extract_vertical_align(run_properties)

    formatted = text
    if is_bold and is_italic:
        formatted = f"***{formatted}***"
    elif is_bold:
        formatted = f"**{formatted}**"
    elif is_italic:
        formatted = f"*{formatted}*"

    if is_underline:
        formatted = f"<u>{formatted}</u>"
    if vertical_align == "superscript":
        formatted = f"<sup>{formatted}</sup>"
    elif vertical_align == "subscript":
        formatted = f"<sub>{formatted}</sub>"
    return formatted


_OOXML_TOGGLE_OFF_VALUES = {"0", "false", "off", "none"}


def _run_toggle_property_is_on(run_properties, local_name: str) -> bool:
    element = find_child_element(run_properties, local_name)
    if element is None:
        return False
    value = get_xml_attribute(element, "val")
    if value is None:
        # An OOXML toggle property with no w:val defaults to enabled.
        return True
    return value.strip().lower() not in _OOXML_TOGGLE_OFF_VALUES


def _extract_vertical_align(run_properties) -> str | None:
    vertical_align = find_child_element(run_properties, "vertAlign")
    return get_xml_attribute(vertical_align, "val") if vertical_align is not None else None


def _extract_run_element_images(run_element, part) -> list[tuple[bytes, str | None, int | None, int | None, dict[str, object]]]:
    return extract_run_element_images(
        run_element,
        part,
        relationship_namespace=RELATIONSHIP_NAMESPACE,
    )


def _extract_run_image_placeholders(run_element, part, image_assets: list[ImageAsset]) -> list[str]:
    placeholders: list[str] = []
    for image_blob, mime_type, width_emu, height_emu, source_forensics in _extract_run_element_images(run_element, part):
        image_index = len(image_assets) + 1
        placeholder = f"[[DOCX_IMAGE_img_{image_index:03d}]]"
        image_assets.append(
            ImageAsset(
                image_id=f"img_{image_index:03d}",
                placeholder=placeholder,
                original_bytes=image_blob,
                mime_type=mime_type,
                position_index=image_index - 1,
                width_emu=width_emu,
                height_emu=height_emu,
                source_forensics=source_forensics,
            )
        )
        placeholders.append(placeholder)
    return placeholders


def _extract_paragraph_asset_id(text: str, *, role: str) -> str | None:
    if role != "image":
        return None
    placeholders = IMAGE_PLACEHOLDER_PATTERN.findall(text)
    if len(placeholders) != 1:
        return None
    placeholder = placeholders[0]
    match = re.match(r"\[\[DOCX_IMAGE_(img_\d+)\]\]", placeholder)
    if match is None:
        return None
    return match.group(1)


def _empty_list_metadata() -> dict[str, object]:
    return {
        "list_kind": None,
        "list_level": 0,
        "list_numbering_format": None,
        "list_num_id": None,
        "list_abstract_num_id": None,
        "list_num_xml": None,
        "list_abstract_num_xml": None,
    }


def _is_typographic_emdash_bullet(numbering_details: dict[str, str | None]) -> bool:
    return (
        numbering_details.get("num_format") == "bullet"
        and (numbering_details.get("lvl_text") or "") in _TYPOGRAPHIC_BULLET_CHARS
    )


def _extract_paragraph_list_metadata(paragraph, text: str, style_name: str, role: str) -> dict[str, object]:
    metadata: dict[str, object] = _empty_list_metadata()

    explicit_kind = detect_explicit_list_kind(text)
    style_level = _extract_style_list_level(style_name)
    num_pr = _resolve_paragraph_num_pr(paragraph)

    if role != "list" and explicit_kind is None and num_pr is None:
        return metadata

    if explicit_kind is not None:
        metadata["list_kind"] = explicit_kind
        if num_pr is not None:
            numbering_details = _resolve_num_pr_details(paragraph, num_pr)
            if _is_typographic_emdash_bullet(numbering_details):
                return _empty_list_metadata()
            numbering_format = numbering_details["num_format"]
            metadata["list_level"] = _extract_num_pr_level(num_pr)
            metadata["list_numbering_format"] = numbering_format
            metadata["list_num_id"] = numbering_details["num_id"]
            metadata["list_abstract_num_id"] = numbering_details["abstract_num_id"]
            metadata["list_num_xml"] = numbering_details["num_xml"]
            metadata["list_abstract_num_xml"] = numbering_details["abstract_num_xml"]
            if numbering_format in ORDERED_LIST_FORMATS:
                metadata["list_kind"] = "ordered"
            elif numbering_format in UNORDERED_LIST_FORMATS:
                metadata["list_kind"] = "unordered"
        return metadata

    if num_pr is not None:
        list_level = max(_extract_num_pr_level(num_pr), style_level)
        numbering_details = _resolve_num_pr_details(paragraph, num_pr)
        if _is_typographic_emdash_bullet(numbering_details):
            metadata["_is_typographic_emdash_bullet"] = True
            return metadata
        numbering_format = numbering_details["num_format"]
        metadata["list_level"] = list_level
        metadata["list_numbering_format"] = numbering_format
        metadata["list_num_id"] = numbering_details["num_id"]
        metadata["list_abstract_num_id"] = numbering_details["abstract_num_id"]
        metadata["list_num_xml"] = numbering_details["num_xml"]
        metadata["list_abstract_num_xml"] = numbering_details["abstract_num_xml"]
        if numbering_format in ORDERED_LIST_FORMATS:
            metadata["list_kind"] = "ordered"
            return metadata
        if numbering_format in UNORDERED_LIST_FORMATS:
            metadata["list_kind"] = "unordered"
            return metadata
        if role != "list":
            return metadata

    if role != "list":
        return metadata

    normalized_style = style_name.strip().lower()
    if any(token in normalized_style for token in ("number", "num", "нумер", "числ")):
        metadata["list_kind"] = "ordered"
        metadata["list_level"] = style_level
        return metadata
    if any(token in normalized_style for token in ("bullet", "bulleted", "маркир", "маркер")):
        metadata["list_kind"] = "unordered"
        metadata["list_level"] = style_level
        return metadata
    metadata["list_kind"] = "unordered"
    metadata["list_level"] = style_level
    return metadata


def _extract_style_list_level(style_name: str) -> int:
    match = re.search(r"(\d+)\s*$", style_name.strip())
    if match is None:
        return 0
    try:
        return max(0, int(match.group(1)) - 1)
    except ValueError:
        return 0


def _resolve_paragraph_num_pr(paragraph):
    return resolve_paragraph_num_pr(
        paragraph,
        find_child_element=find_child_element,
    )


def _extract_num_pr_level(num_pr) -> int:
    return extract_num_pr_level(
        num_pr,
        find_child_element=find_child_element,
        get_xml_attribute=get_xml_attribute,
    )


def _resolve_num_pr_details(paragraph, num_pr) -> dict[str, str | None]:
    return resolve_num_pr_details(
        paragraph,
        num_pr,
        xml_local_name=xml_local_name,
        find_child_element=find_child_element,
        get_xml_attribute=get_xml_attribute,
    )


def _validate_docx_archive(source_bytes: bytes) -> None:
    if len(source_bytes) > MAX_DOCX_ARCHIVE_SIZE_BYTES:
        raise RuntimeError("DOCX-файл превышает допустимый размер архива.")

    try:
        with zipfile.ZipFile(BytesIO(source_bytes)) as archive:
            entries = archive.infolist()
    except zipfile.BadZipFile as exc:
        raise RuntimeError("Передан поврежденный или неподдерживаемый DOCX-архив.") from exc

    if not entries:
        raise RuntimeError("Передан пустой DOCX-архив.")
    if len(entries) > MAX_DOCX_ENTRY_COUNT:
        raise RuntimeError("DOCX-архив содержит слишком много файлов и отклонен из соображений безопасности.")

    total_uncompressed_size = sum(max(0, entry.file_size) for entry in entries)
    total_compressed_size = sum(max(0, entry.compress_size) for entry in entries)
    if total_uncompressed_size > MAX_DOCX_UNCOMPRESSED_SIZE_BYTES:
        raise RuntimeError("DOCX-архив слишком велик после распаковки и отклонен из соображений безопасности.")
    if total_compressed_size > 0 and (total_uncompressed_size / total_compressed_size) > MAX_DOCX_COMPRESSION_RATIO:
        raise RuntimeError("DOCX-архив имеет подозрительно высокий коэффициент сжатия и отклонен из соображений безопасности.")

    for entry in entries:
        entry_name = entry.filename
        parts = entry_name.replace("\\", "/").split("/")
        if any(part == ".." for part in parts):
            raise RuntimeError("DOCX-архив содержит подозрительные пути и отклонён из соображений безопасности.")
        if entry_name.startswith("/"):
            raise RuntimeError("DOCX-архив содержит абсолютные пути и отклонён из соображений безопасности.")

    filenames = {entry.filename for entry in entries}
    if "[Content_Types].xml" not in filenames:
        raise RuntimeError("Передан невалидный DOCX-архив: отсутствует [Content_Types].xml.")


def validate_docx_source_bytes(source_bytes: bytes) -> None:
    _validate_docx_archive(source_bytes)


def classify_paragraph_role(text: str, style_name: str, *, heading_level: int | None = None) -> str:
    from docxaicorrector.document.roles import classify_paragraph_role as _classify_paragraph_role

    return _classify_paragraph_role(text, style_name, heading_level=heading_level)


def has_heading_text_signal(text: str) -> bool:
    from docxaicorrector.document.roles import has_heading_text_signal as _has_heading_text_signal

    return _has_heading_text_signal(text)


def build_document_text(paragraphs: list[ParagraphUnit]) -> str:
    return "\n\n".join(paragraph.rendered_text for paragraph in paragraphs).strip()


def inspect_placeholder_integrity(markdown_text: str, image_assets: list[ImageAsset]) -> dict[str, str]:
    status_map: dict[str, str] = {}
    expected_placeholders = {asset.placeholder for asset in image_assets}
    for asset in image_assets:
        occurrence_count = markdown_text.count(asset.placeholder)
        if occurrence_count == 1:
            status_map[asset.image_id] = "ok"
        elif occurrence_count == 0:
            status_map[asset.image_id] = "lost"
        else:
            status_map[asset.image_id] = "duplicated"
    for unexpected_placeholder in sorted(set(IMAGE_PLACEHOLDER_PATTERN.findall(markdown_text)) - expected_placeholders):
        status_map[f"unexpected:{unexpected_placeholder}"] = "unexpected"
    return status_map
